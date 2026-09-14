"""Carga pequena, determinística e idempotente para o banco isolado de testes.

O modo padrão é somente leitura. Para gravar, use ``--apply`` e configure
``TEST_DATABASE_URL`` apontando para um banco cujo nome contenha ``test``.
"""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
import hashlib
import json
import os
from pathlib import Path
import re
import sys
import unicodedata

from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.database.config import load_postgres_config  # noqa: E402
from app.database.database import Database  # noqa: E402
from app.core.operator_sectors import operator_route_resource_label  # noqa: E402


DEFAULT_SOURCE = Path(
    os.environ.get(
        "GESTOR_SEED_SOURCE",
        PROJECT_ROOT / "dados" / "Relatorio_de_Operacao_OF_com_OPs_Ficticias_CORRIGIDO.xlsx",
    )
)
DEFAULT_STATUS_SOURCE = PROJECT_ROOT / "data" / "pcfactory_status_recursos.csv"
DEFAULT_RESOURCE_SOURCE = PROJECT_ROOT / "docs" / "pcfactory_recursos_sem_pessoas.csv"
OP_PATTERN = re.compile(r"[A-Z]{6}\d{5}")
SECTORS = ("Corte", "Dobra", "Usinagem", "Serra", "Solda", "Pintura")
SEED_LOCK_ID = 874_210_308
CUT_TEST_TASK_CODE = "TSTCORTE-001"


@dataclass(frozen=True)
class SourceOperation:
    filial: str
    tipo: str
    roteiro: str
    produto_codigo: str
    numero_operacao: str
    codigo_recurso: str
    descricao_operacao: str
    produto_descricao: str
    tempo_medio_segundos: float | None
    codigo_op: str
    tipo_setor: str | None
    origem_aba: str

    @property
    def key(self):
        return self.codigo_op, self.numero_operacao, self.codigo_recurso


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header_key(value):
    text = unicodedata.normalize("NFKD", _text(value)).casefold()
    text = "".join(character for character in text if not unicodedata.combining(character))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _header_indexes(headers):
    keys = [_header_key(value) for value in headers]

    def first(predicate, required=True):
        index = next((position for position, key in enumerate(keys) if predicate(key)), None)
        if required and index is None:
            raise ValueError(f"Cabeçalho obrigatório ausente: {headers}")
        return index

    return {
        "filial": first(lambda key: key == "filial", required=False),
        "tipo": first(lambda key: key == "tipo", required=False),
        "roteiro": first(lambda key: "roteiro" in key and "cod" in key, required=False),
        "produto": first(lambda key: key == "produto"),
        "numero_operacao": first(lambda key: "num" in key and "opera" in key),
        "recurso": first(lambda key: "recurso" in key),
        "descricao_operacao": first(lambda key: "opera" in key and ("desc" in key or "descri" in key)),
        "produto_descricao": first(lambda key: "produto" in key and ("desc" in key or "descri" in key)),
        "tempo": first(lambda key: "tempo" in key and "medio" in key, required=False),
        "op": first(lambda key: key.startswith("op ") and "fict" in key),
    }


def _seconds(value):
    if value in (None, ""):
        return None
    if isinstance(value, timedelta):
        return round(value.total_seconds(), 3)
    if isinstance(value, time):
        return float(value.hour * 3600 + value.minute * 60 + value.second)
    if isinstance(value, (int, float)):
        number = float(value)
        return round(number * 86400 if 0 < number <= 1 else number, 3)
    match = re.fullmatch(r"(\d{1,3}):(\d{2})(?::(\d{2}))?", str(value).strip())
    if match:
        hours, minutes, seconds = (int(item or 0) for item in match.groups())
        return float(hours * 3600 + minutes * 60 + seconds)
    return None


def classify_sector(resource):
    code = _text(resource).upper()
    if "LASER" in code or "PLASMA" in code:
        return "Corte"
    if "DOBRA" in code:
        return "Dobra"
    if any(token in code for token in ("CNC", "TORNO", "FRESA", "USIN", "MCV", "ROSQ", "MANDRI")):
        return "Usinagem"
    if "SERRA" in code:
        return "Serra"
    if "SOLDA" in code:
        return "Solda"
    if "PINT" in code:
        return "Pintura"
    return None


def _flag(value):
    return str(value or "").strip().casefold() in {"1", "true", "sim", "yes"}


def read_status_catalog(path):
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = []
        for raw in csv.DictReader(stream, delimiter=";"):
            code = _text(raw.get("codigo"))
            name = _text(raw.get("nome"))
            if not code or not name:
                continue
            rows.append(
                {
                    "codigo": code,
                    "nome": name,
                    "grupo_id": int(raw["grupo_id"]) if _text(raw.get("grupo_id")) else None,
                    "grupo_codigo": _text(raw.get("grupo_codigo")),
                    "grupo_nome": _text(raw.get("grupo_nome")),
                    "habilitado": _flag(raw.get("habilitado")),
                    "classificacao": (
                        int(raw["classificacao"])
                        if _text(raw.get("classificacao"))
                        else None
                    ),
                    "setup": _flag(raw.get("setup")),
                    "retrabalho": _flag(raw.get("retrabalho")),
                    "retorno_automatico": _flag(raw.get("retorno_automatico")),
                    "parada_geral": _flag(raw.get("parada_geral")),
                    "oculto": _flag(raw.get("oculto")),
                    "requer_detalhe": _flag(raw.get("requer_detalhe")),
                    "requer_comentario": _flag(raw.get("requer_comentario")),
                    "acao_padrao": (
                        int(raw["acao_padrao"])
                        if _text(raw.get("acao_padrao"))
                        else None
                    ),
                }
            )
    if len({row["codigo"] for row in rows}) != len(rows):
        raise ValueError("Catálogo de status contém códigos duplicados.")
    return rows


def read_resource_catalog(path):
    with path.open("r", encoding="utf-8-sig", newline="") as stream:
        rows = []
        for raw in csv.DictReader(stream, delimiter=";"):
            code = _text(raw.get("codigo")).upper()
            name = _text(raw.get("nome"))
            if not code or not name:
                continue
            rows.append(
                {
                    "codigo": code,
                    "nome": name,
                    "tipo_setor": classify_sector(f"{code} {name}"),
                    "fonte": _text(raw.get("fonte")) or "PC Factory D0021",
                }
            )
    if len({row["codigo"] for row in rows}) != len(rows):
        raise ValueError("Catálogo de recursos contém códigos duplicados.")
    return rows


def read_source(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    operations = {}
    total_rows = 0
    valid_rows = 0
    try:
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            rows = sheet.iter_rows(values_only=True)
            headers = tuple(next(rows, None) or ())
            try:
                indexes = _header_indexes(headers)
            except ValueError:
                continue
            for values in rows:
                total_rows += 1
                values = tuple(values or ())

                def at(name):
                    position = indexes[name]
                    return values[position] if position is not None and position < len(values) else None

                product = _text(at("produto")).upper()
                operation_number = _text(at("numero_operacao"))
                resource = _text(at("recurso")).upper()
                op = _text(at("op")).upper()
                if not product or not operation_number or not resource or not OP_PATTERN.fullmatch(op):
                    continue
                valid_rows += 1
                item = SourceOperation(
                    filial=_text(at("filial")),
                    tipo=_text(at("tipo")),
                    roteiro=_text(at("roteiro")),
                    produto_codigo=product,
                    numero_operacao=operation_number,
                    codigo_recurso=resource,
                    descricao_operacao=_text(at("descricao_operacao")) or resource,
                    produto_descricao=_text(at("produto_descricao")) or product,
                    tempo_medio_segundos=_seconds(at("tempo")),
                    codigo_op=op,
                    tipo_setor=classify_sector(resource),
                    origem_aba=sheet_name,
                )
                operations.setdefault(item.key, item)
    finally:
        workbook.close()
    return list(operations.values()), total_rows, valid_rows


def select_sample(operations, per_sector=3):
    candidates = defaultdict(list)
    seen = defaultdict(set)
    for row in sorted(
        operations,
        key=lambda item: (item.codigo_op, int(item.numero_operacao) if item.numero_operacao.isdigit() else 999999, item.codigo_recurso),
    ):
        if row.tipo_setor and row.codigo_op not in seen[row.tipo_setor]:
            candidates[row.tipo_setor].append(row)
            seen[row.tipo_setor].add(row.codigo_op)
    missing = [sector for sector in SECTORS if len(candidates[sector]) < per_sector]
    if missing:
        raise ValueError(f"Planilha sem amostra suficiente para: {', '.join(missing)}")
    scenarios = {sector: candidates[sector][:per_sector] for sector in SECTORS}
    selected_ops = {row.codigo_op for rows in scenarios.values() for row in rows}
    selected_routes = [
        row for row in operations if row.codigo_op in selected_ops and row.tipo_setor in SECTORS
    ]
    return scenarios, sorted(selected_routes, key=lambda item: item.key)


def source_sha256(path):
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def quantity_for(op):
    number = int(op[-5:])
    return (40, 60, 80, 100)[number % 4]


def build_summary(source, operations, total_rows, valid_rows, scenarios, routes):
    selected_ops = sorted({row.codigo_op for row in routes})
    return {
        "source": str(source),
        "source_sha256": source_sha256(source),
        "sheets_and_raw_rows": {"rows": total_rows},
        "valid_rows": valid_rows,
        "unique_operations": len(operations),
        "selected_ops": selected_ops,
        "selected_route_rows": len(routes),
        "scenarios": {
            sector: [
                {
                    "op": row.codigo_op,
                    "produto": row.produto_codigo,
                    "operacao": row.numero_operacao,
                    "recurso": row.codigo_recurso,
                }
                for row in rows
            ]
            for sector, rows in scenarios.items()
        },
    }


def _upsert_pcfactory_catalogs(cursor, statuses, resources, routes, now):
    resource_by_code = {row["codigo"]: dict(row) for row in resources}
    for route in routes:
        resource_by_code.setdefault(
            route.codigo_recurso,
            {
                "codigo": route.codigo_recurso,
                "nome": route.codigo_recurso,
                "tipo_setor": route.tipo_setor,
                "fonte": "planilha_teste_roteiro",
            },
        )
    for row in resource_by_code.values():
        cursor.execute(
            """
            INSERT INTO catalogo_recursos_pcfactory (
                codigo, nome, tipo_setor, habilitado, fonte, sincronizado_em
            ) VALUES (%s, %s, %s, TRUE, %s, %s)
            ON CONFLICT (codigo) DO UPDATE SET
                nome = EXCLUDED.nome,
                tipo_setor = COALESCE(EXCLUDED.tipo_setor, catalogo_recursos_pcfactory.tipo_setor),
                habilitado = TRUE,
                fonte = EXCLUDED.fonte,
                sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (
                row["codigo"], row["nome"], row.get("tipo_setor"),
                row.get("fonte") or "PC Factory D0021", now,
            ),
        )
    for row in statuses:
        cursor.execute(
            """
            INSERT INTO catalogo_status_recursos (
                codigo, nome, grupo_id, grupo_codigo, grupo_nome, habilitado,
                classificacao, setup, retrabalho, retorno_automatico,
                parada_geral, oculto, requer_detalhe, requer_comentario,
                acao_padrao, fonte, sincronizado_em
            ) VALUES (
                %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                %s, %s, %s, %s, %s, 'TBLResourceStatus.xls', %s
            )
            ON CONFLICT (codigo) DO UPDATE SET
                nome = EXCLUDED.nome,
                grupo_id = EXCLUDED.grupo_id,
                grupo_codigo = EXCLUDED.grupo_codigo,
                grupo_nome = EXCLUDED.grupo_nome,
                habilitado = EXCLUDED.habilitado,
                classificacao = EXCLUDED.classificacao,
                setup = EXCLUDED.setup,
                retrabalho = EXCLUDED.retrabalho,
                retorno_automatico = EXCLUDED.retorno_automatico,
                parada_geral = EXCLUDED.parada_geral,
                oculto = EXCLUDED.oculto,
                requer_detalhe = EXCLUDED.requer_detalhe,
                requer_comentario = EXCLUDED.requer_comentario,
                acao_padrao = EXCLUDED.acao_padrao,
                fonte = EXCLUDED.fonte,
                sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (
                row["codigo"], row["nome"], row["grupo_id"],
                row["grupo_codigo"], row["grupo_nome"], row["habilitado"],
                row["classificacao"], row["setup"], row["retrabalho"],
                row["retorno_automatico"], row["parada_geral"], row["oculto"],
                row["requer_detalhe"], row["requer_comentario"],
                row["acao_padrao"], now,
            ),
        )
    return len(resource_by_code), len(statuses)


def _upsert_catalog(cursor, routes, now):
    primary = {}
    for row in routes:
        primary.setdefault(row.codigo_op, row)
    for op, row in primary.items():
        cursor.execute(
            """
            INSERT INTO catalogo_pcp_ops (
                codigo_op, produto_codigo, produto_descricao, quantidade, unidade,
                data_emissao, status_pcp, filial, local_estoque, roteiro, recurso,
                ativo, sincronizado_em
            ) VALUES (%s, %s, %s, %s, 'UN', %s, 'LIBERADA', %s, 'TESTE', %s, %s, TRUE, %s)
            ON CONFLICT (codigo_op) DO UPDATE SET
                produto_codigo = EXCLUDED.produto_codigo,
                produto_descricao = EXCLUDED.produto_descricao,
                quantidade = EXCLUDED.quantidade,
                unidade = EXCLUDED.unidade,
                data_emissao = EXCLUDED.data_emissao,
                status_pcp = EXCLUDED.status_pcp,
                filial = EXCLUDED.filial,
                local_estoque = EXCLUDED.local_estoque,
                roteiro = EXCLUDED.roteiro,
                recurso = EXCLUDED.recurso,
                ativo = TRUE,
                sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (
                op, row.produto_codigo, row.produto_descricao, quantity_for(op), now.date(),
                row.filial, row.roteiro, row.codigo_recurso, now,
            ),
        )
    for index, row in enumerate(routes, start=1):
        order = int(row.numero_operacao) if row.numero_operacao.isdigit() else index
        cursor.execute(
            """
            INSERT INTO catalogo_operacoes_op (
                codigo_op, produto_codigo, produto_descricao, numero_operacao,
                codigo_recurso, descricao_operacao, tipo_setor, filial, tipo,
                roteiro, tempo_medio_segundos, ordem, fonte, ativo, sincronizado_em
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                      'planilha_teste', TRUE, %s)
            ON CONFLICT (codigo_op, numero_operacao, codigo_recurso) DO UPDATE SET
                produto_codigo = EXCLUDED.produto_codigo,
                produto_descricao = EXCLUDED.produto_descricao,
                descricao_operacao = EXCLUDED.descricao_operacao,
                tipo_setor = EXCLUDED.tipo_setor,
                filial = EXCLUDED.filial,
                tipo = EXCLUDED.tipo,
                roteiro = EXCLUDED.roteiro,
                tempo_medio_segundos = EXCLUDED.tempo_medio_segundos,
                ordem = EXCLUDED.ordem,
                fonte = EXCLUDED.fonte,
                ativo = TRUE,
                sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (
                row.codigo_op, row.produto_codigo, row.produto_descricao,
                row.numero_operacao, row.codigo_recurso, row.descricao_operacao,
                row.tipo_setor, row.filial, row.tipo, row.roteiro,
                row.tempo_medio_segundos, order, now,
            ),
        )
    return primary


def _upsert_tasks(cursor, primary, now, grouped_ops=()):
    grouped_ops = {str(op).strip().upper() for op in grouped_ops}
    legacy_task_codes = [f"TSTUI-{op[-5:]}" for op in grouped_ops]
    if legacy_task_codes:
        cursor.execute(
            "DELETE FROM tarefas WHERE codigo_tarefa = ANY(%s)",
            (legacy_task_codes,),
        )
    task_ids = {}
    for index, (op, row) in enumerate(sorted(primary.items()), start=1):
        if op in grouped_ops:
            continue
        task_code = f"TSTUI-{op[-5:]}"
        status = ("Destacando", "Finalizado", "Despachado")[index % 3]
        cursor.execute(
            """
            INSERT INTO tarefas (
                codigo_tarefa, material, espessura, status, data_inicio_destaque,
                data_finalizacao, data_despacho, observacoes
            ) VALUES (%s, 'Aço carbono - teste', 3.00, %s, %s, %s, %s,
                      'SEED FICTÍCIO ISOLADO')
            ON CONFLICT (codigo_tarefa) DO UPDATE SET
                material = EXCLUDED.material,
                espessura = EXCLUDED.espessura,
                status = EXCLUDED.status,
                data_inicio_destaque = EXCLUDED.data_inicio_destaque,
                data_finalizacao = EXCLUDED.data_finalizacao,
                data_despacho = EXCLUDED.data_despacho,
                observacoes = EXCLUDED.observacoes
            RETURNING id
            """,
            (
                task_code,
                status,
                now - timedelta(days=3),
                now - timedelta(days=2) if status in {"Finalizado", "Despachado"} else None,
                now - timedelta(days=1) if status == "Despachado" else None,
            ),
        )
        task_id = int(cursor.fetchone()["id"])
        task_ids[op] = task_id
        destination = f"Aguardando {row.tipo_setor or 'Dobra'}"
        cursor.execute(
            """
            INSERT INTO op_por_tarefa (
                tarefa_id, codigo_op, id_peca, setor_destino_original,
                setor_destino_atual, quantidade_original, quantidade_atual, editado
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, FALSE)
            ON CONFLICT (tarefa_id, codigo_op) DO UPDATE SET
                id_peca = EXCLUDED.id_peca,
                setor_destino_original = EXCLUDED.setor_destino_original,
                setor_destino_atual = EXCLUDED.setor_destino_atual,
                quantidade_original = EXCLUDED.quantidade_original,
                quantidade_atual = EXCLUDED.quantidade_atual,
                editado = FALSE
            """,
            (
                task_id, op, row.produto_codigo, destination, destination,
                quantity_for(op), quantity_for(op),
            ),
        )
    return task_ids


def _appointment_scenarios():
    return {
        "Dobra": (("1303", "Em processo"), ("2204", "Aguardando"), ("Gasparini", "Finalizado")),
        "Usinagem": (("Eurostec", "Em processo"), ("Fresadora FTV31", "Setup"), ("Romi D 1000", "Aguardando")),
        "Serra": (("SFG-330", "Parada"), ("S4220", "Em processo"), ("SFHA-10", "Finalizado")),
        "Solda": (("Estação 1", "Retrabalho"), ("Estação 2", "Em processo"), ("Estação 3", "Aguardando")),
        "Pintura": (("Pintura", "Em processo"), ("Pintura", "Aguardando"), ("Pintura", "Finalizado")),
    }


def _seed_appointments(cursor, scenarios, task_ids, selected_ops, now, excluded_ops=()):
    excluded_ops = {str(op).strip().upper() for op in excluded_ops}
    cursor.execute(
        "DELETE FROM historico WHERE op = ANY(%s) AND tipo = 'Apontamento Operador'",
        (selected_ops,),
    )
    cursor.execute("DELETE FROM apontamentos_operacionais WHERE op = ANY(%s)", (selected_ops,))
    cursor.execute(
        """
        SELECT id, codigo_op, numero_operacao, codigo_recurso
        FROM catalogo_operacoes_op
        WHERE codigo_op = ANY(%s)
        """,
        (selected_ops,),
    )
    operation_ids = {
        (row["codigo_op"], row["numero_operacao"], row["codigo_recurso"]): row["id"]
        for row in cursor.fetchall()
    }
    total = 0
    for sector, definitions in _appointment_scenarios().items():
        eligible_rows = [
            row for row in scenarios[sector]
            if row.codigo_op not in excluded_ops
        ]
        for index, (row, (machine, status)) in enumerate(zip(eligible_rows, definitions)):
            quantity = quantity_for(row.codigo_op)
            entry = now - timedelta(hours=8 + index * 3)
            started = None if status == "Aguardando" else entry + timedelta(minutes=25)
            finished = started + timedelta(minutes=70) if status == "Finalizado" else None
            scrap = min(5, quantity) if status == "Finalizado" and index % 2 == 0 else 0
            good = quantity - scrap if status == "Finalizado" else 0
            reason = {
                "Parada": "0029 - Aguardando ponte",
                "Setup": "1005 - Set-Up",
                "Retrabalho": "0040 - AGUARDANDO RETRABALHO PRODUÇÃO",
            }.get(status)
            reason_code = {
                "Parada": "0029",
                "Setup": "1005",
                "Retrabalho": "0040",
            }.get(status)
            operator = f"SEED {sector.upper()}"
            cursor.execute(
                """
                INSERT INTO apontamentos_operacionais (
                    op, peca, tarefa_id, tipo_setor, maquina, status, quantidade,
                    operador_fila, data_entrada, operador_inicio, data_inicio,
                    operador_fim, data_fim, setor_destino, catalogo_operacao_id,
                    numero_operacao, codigo_recurso, descricao_operacao,
                    produto_codigo, produto_descricao, quantidade_boa,
                    quantidade_refugo, motivo_parada, comentario,
                    codigo_status_recurso
                ) VALUES (
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s,
                    %s, %s, %s, %s, %s, %s, %s, %s, %s, 'Cenário controlado de teste', %s
                ) RETURNING id
                """,
                (
                    row.codigo_op, row.produto_codigo, task_ids[row.codigo_op], sector,
                    machine, status, quantity, operator, entry,
                    operator if started else None, started,
                    operator if finished else None, finished,
                    "Almoxarifado" if finished else None,
                    operation_ids[row.key], row.numero_operacao, row.codigo_recurso,
                    row.descricao_operacao, row.produto_codigo, row.produto_descricao,
                    good, scrap, reason, reason_code,
                ),
            )
            appointment_id = int(cursor.fetchone()["id"])
            events = [("fila", entry, None, 0, 0)]
            if started:
                events.append(("producao", started, None, 0, 0))
            if status in {"Parada", "Setup", "Retrabalho"}:
                events.append((status.lower(), started + timedelta(minutes=35), reason, 0, 0))
            if finished:
                events.append(("finalizado", finished, None, good, scrap))
            for state, event_time, event_reason, event_good, event_scrap in events:
                cursor.execute(
                    """
                    INSERT INTO eventos_apontamento_operador (
                        apontamento_id, estado, motivo, comentario, quantidade_boa,
                        quantidade_refugo, operador, data_hora, codigo_status_recurso
                    ) VALUES (%s, %s, %s, 'Cenário controlado de teste', %s, %s, %s, %s, %s)
                    """,
                    (
                        appointment_id, state, event_reason, event_good, event_scrap,
                        operator, event_time, reason_code if state in {"parada", "setup", "retrabalho"} else None,
                    ),
                )
                if state != "fila":
                    cursor.execute(
                        """
                        INSERT INTO historico (
                            op, tipo, setor, motivo, quantidade, operador,
                            data_hora, peca, tarefa_id
                        ) VALUES (%s, 'Apontamento Operador', %s, %s, %s, %s, %s, %s, %s)
                        """,
                        (
                            row.codigo_op, machine, event_reason or f"Estado: {state}",
                            event_good + event_scrap if state == "finalizado" else quantity,
                            operator, event_time, row.produto_codigo, task_ids[row.codigo_op],
                        ),
                    )
            total += 1
    return total


def _route_order(row):
    try:
        return int(row.numero_operacao)
    except (TypeError, ValueError):
        return 999_999


def destination_after_cut(routes, resource_names=None):
    """Retorna o posto/setor real depois do Corte, sem alias ``Aguardando``."""

    resource_names = dict(resource_names or {})
    found_cut = False
    for row in sorted(routes, key=lambda item: (_route_order(item), item.codigo_recurso)):
        if row.tipo_setor == "Corte":
            found_cut = True
            continue
        if found_cut and row.tipo_setor:
            return operator_route_resource_label(
                row.tipo_setor,
                row.codigo_recurso,
                resource_names.get(row.codigo_recurso),
            )
    return "Almoxarifado"


def _seed_cut(cursor, rows, routes, resources, now):
    rows = list(rows)
    resource_names = {
        str(resource.get("codigo") or "").strip(): str(resource.get("nome") or "").strip()
        for resource in (resources or ())
    }
    route_by_op = defaultdict(list)
    for route in routes:
        route_by_op[route.codigo_op].append(route)

    cursor.execute("DELETE FROM apontamentos_corte WHERE plano_hash LIKE 'seed:%'")
    cursor.execute("DELETE FROM catalogo_sigmanest_planos_corte WHERE plano_hash LIKE 'seed:%'")
    cursor.execute(
        "DELETE FROM catalogo_sigmanest_tarefas WHERE codigo_tarefa LIKE 'TSTCUT-%'"
    )
    cursor.execute(
        "DELETE FROM catalogo_sigmanest_programas WHERE codigo_tarefa = %s",
        (CUT_TEST_TASK_CODE,),
    )
    cursor.execute(
        "DELETE FROM catalogo_sigmanest_ops WHERE codigo_tarefa = %s",
        (CUT_TEST_TASK_CODE,),
    )

    cursor.execute(
        """
        INSERT INTO catalogo_sigmanest_tarefas (
            codigo_tarefa, material, espessura, ativo, sincronizado_em
        ) VALUES (%s, 'Aço carbono - teste', 3.00, TRUE, %s)
        ON CONFLICT (codigo_tarefa) DO UPDATE SET
            material = EXCLUDED.material,
            espessura = EXCLUDED.espessura,
            ativo = TRUE,
            sincronizado_em = EXCLUDED.sincronizado_em
        """,
        (CUT_TEST_TASK_CODE, now),
    )
    cursor.execute(
        """
        INSERT INTO tarefas (
            codigo_tarefa, material, espessura, status, data_inicio_destaque,
            data_finalizacao, data_despacho, observacoes
        ) VALUES (%s, 'Aço carbono - teste', 3.00, NULL, NULL, NULL, NULL,
                  'SEED FICTÍCIO ISOLADO; DESTAQUE NÃO PERTENCE AO ROTEIRO')
        ON CONFLICT (codigo_tarefa) DO UPDATE SET
            material = EXCLUDED.material,
            espessura = EXCLUDED.espessura,
            status = NULL,
            data_inicio_destaque = NULL,
            data_finalizacao = NULL,
            data_despacho = NULL,
            observacoes = EXCLUDED.observacoes
        RETURNING id
        """,
        (CUT_TEST_TASK_CODE,),
    )
    task_id = int(cursor.fetchone()["id"])
    cursor.execute("DELETE FROM historico WHERE tarefa_id = %s", (task_id,))
    cursor.execute("DELETE FROM op_por_tarefa WHERE tarefa_id = %s", (task_id,))

    total_plans = 0
    machine_sequences = defaultdict(int)
    for row in rows:
        task = CUT_TEST_TASK_CODE
        program = f"TST-{row.codigo_op[-5:]}"
        machine = (
            "MESSER_XPR_300"
            if "PLASMA" in row.codigo_recurso.upper()
            else "AMADA_ENSIS"
        )
        machine_sequences[machine] += 1
        sequence = machine_sequences[machine]
        op_routes = route_by_op[row.codigo_op]
        route_sectors = {route.tipo_setor for route in op_routes}
        destination = destination_after_cut(op_routes, resource_names)
        cursor.execute(
            """
            INSERT INTO catalogo_sigmanest_programas (
                codigo_tarefa, programa, ativo, sincronizado_em
            ) VALUES (%s, %s, TRUE, %s)
            ON CONFLICT (codigo_tarefa, programa) DO UPDATE SET
                ativo = TRUE, sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (task, program, now),
        )
        line_hash = hashlib.sha256(f"{task}|{row.codigo_op}|{row.produto_codigo}".encode()).hexdigest()
        cursor.execute(
            """
            INSERT INTO catalogo_sigmanest_ops (
                linha_hash, codigo_tarefa, codigo_op, id_peca, setor_destino,
                quantidade, dobra, usinagem, solda, chanfro, ativo, sincronizado_em
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, NULL, TRUE, %s)
            ON CONFLICT (linha_hash) DO UPDATE SET
                codigo_tarefa = EXCLUDED.codigo_tarefa, codigo_op = EXCLUDED.codigo_op,
                id_peca = EXCLUDED.id_peca, setor_destino = EXCLUDED.setor_destino,
                quantidade = EXCLUDED.quantidade, ativo = TRUE,
                sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (
                line_hash, task, row.codigo_op, row.produto_codigo, destination,
                quantity_for(row.codigo_op),
                "Sim" if "Dobra" in route_sectors else None,
                "Sim" if "Usinagem" in route_sectors else None,
                "Sim" if "Solda" in route_sectors else None,
                now,
            ),
        )
        cursor.execute(
            """
            INSERT INTO op_por_tarefa (
                tarefa_id, codigo_op, id_peca, setor_destino_original,
                setor_destino_atual, quantidade_original, quantidade_atual, editado
            ) VALUES (%s, %s, %s, %s, %s, %s, %s, FALSE)
            ON CONFLICT (tarefa_id, codigo_op) DO UPDATE SET
                id_peca = EXCLUDED.id_peca,
                setor_destino_original = EXCLUDED.setor_destino_original,
                setor_destino_atual = EXCLUDED.setor_destino_atual,
                quantidade_original = EXCLUDED.quantidade_original,
                quantidade_atual = EXCLUDED.quantidade_atual,
                editado = FALSE
            """,
            (
                task_id, row.codigo_op, row.produto_codigo, destination, destination,
                quantity_for(row.codigo_op), quantity_for(row.codigo_op),
            ),
        )
        plan_hash = f"seed:{task}:{machine}:{row.codigo_op}"
        cursor.execute(
            """
            INSERT INTO catalogo_sigmanest_planos_corte (
                plano_hash, codigo_tarefa, programa, nome_chapa, sequencia_nesting,
                area_usada, fracao_sucata, quantidade_processo, maquina_sigmanest,
                tempo_previsto_segundos, tempo_previsto_formatado, data_programa,
                status_programa, ativo, sincronizado_em
            ) VALUES (%s, %s, %s, %s, %s, 12.50, 0.08, %s, %s, 3600,
                      '1h 00min', %s, 'LIBERADO', TRUE, %s)
            ON CONFLICT (plano_hash) DO UPDATE SET
                codigo_tarefa = EXCLUDED.codigo_tarefa,
                programa = EXCLUDED.programa,
                nome_chapa = EXCLUDED.nome_chapa,
                sequencia_nesting = EXCLUDED.sequencia_nesting,
                quantidade_processo = EXCLUDED.quantidade_processo,
                maquina_sigmanest = EXCLUDED.maquina_sigmanest,
                data_programa = EXCLUDED.data_programa,
                status_programa = EXCLUDED.status_programa,
                ativo = TRUE,
                sincronizado_em = EXCLUDED.sincronizado_em
            """,
            (
                plan_hash, task, program, row.produto_descricao, sequence,
                quantity_for(row.codigo_op), machine, now.date(), now,
            ),
        )
        total_plans += 1
    return total_plans, task_id


def _seed_users(cursor, db, password, now):
    test_users = (
        ("admin", "admin"),
        ("destaque", "operador_destaque"),
        ("dobra", "operador_dobra"),
        ("usinagem", "operador_usinagem"),
        ("serra", "operador_serra"),
        ("corte", "operador_corte"),
        ("pintura", "operador_pintura"),
        ("solda", "operador_solda"),
    )
    legacy_names = (
        "Teste Admin", "Teste Destaque", "Teste Dobra", "Teste Usinagem", "Teste Serra",
        "Teste Corte", "Teste Pintura", "Teste Solda",
    )
    for name, level in test_users:
        cursor.execute(
            """
            INSERT INTO usuarios (nome, senha_hash, nivel, ativo, data_criacao)
            VALUES (%s, %s, %s, TRUE, %s)
            ON CONFLICT (nome) DO UPDATE SET
                senha_hash = EXCLUDED.senha_hash,
                nivel = EXCLUDED.nivel,
                ativo = TRUE
            """,
            (name, db._hash_senha(password), level, now),
        )
    cursor.execute("DELETE FROM usuarios WHERE nome = ANY(%s)", (list(legacy_names),))
    return len(test_users)


def _seed_operator_badges(cursor):
    cursor.execute(
        """
        INSERT INTO operadores_apontamento (cracha, nome, ativo, fonte)
        VALUES ('1', 'Iago', TRUE, 'teste_ficticio')
        ON CONFLICT (cracha) DO UPDATE SET
            nome = EXCLUDED.nome,
            ativo = TRUE,
            fonte = EXCLUDED.fonte
        """
    )
    return 1


def apply_seed(summary, scenarios, routes, statuses, resources, with_users=False):
    config = load_postgres_config(testing=True)
    target = config.safe_target
    database_name = str(target.get("dbname") or "").lower()
    if "test" not in database_name:
        raise RuntimeError("Carga recusada: o nome do banco TEST_DATABASE_URL deve conter 'test'.")
    password = os.getenv("SEED_TEST_PASSWORD", "")
    if with_users and not password:
        raise RuntimeError("Defina SEED_TEST_PASSWORD para criar os usuários opcionais.")

    db = Database(config=config)
    now = datetime.now().replace(microsecond=0)
    selected_ops = sorted({row.codigo_op for row in routes})
    try:
        with db.connection() as connection, connection.cursor() as cursor:
            cursor.execute("SELECT pg_advisory_xact_lock(%s)", (SEED_LOCK_ID,))
            resource_count, status_count = _upsert_pcfactory_catalogs(
                cursor, statuses, resources, routes, now
            )
            badge_count = _seed_operator_badges(cursor)
            primary = _upsert_catalog(cursor, routes, now)
            cut_ops = {row.codigo_op for row in scenarios["Corte"]}
            task_ids = _upsert_tasks(cursor, primary, now, grouped_ops=cut_ops)
            plan_count, cut_task_id = _seed_cut(
                cursor, scenarios["Corte"], routes, resources, now
            )
            task_ids.update({op: cut_task_id for op in cut_ops})
            appointment_count = _seed_appointments(
                cursor, scenarios, task_ids, selected_ops, now,
                excluded_ops=cut_ops,
            )
            user_count = _seed_users(cursor, db, password, now) if with_users else 0
            cursor.execute(
                """
                INSERT INTO eventos_sistema (
                    tipo, origem, referencia, mensagem, operador, detalhes, data_hora
                ) VALUES ('seed_dados_ficticios', 'scripts/seed_dados_ficticios.py',
                          %s, 'Carga fictícia de teste aplicada.', 'SEED', %s, %s)
                """,
                (
                    summary["source_sha256"],
                    json.dumps({"ops": len(selected_ops), "apontamentos": appointment_count}),
                    now,
                ),
            )
            cursor.execute(
                """
                SELECT
                    (SELECT COUNT(*) FROM catalogo_pcp_ops WHERE codigo_op = ANY(%s)) AS catalogo_ops,
                    (SELECT COUNT(*) FROM catalogo_operacoes_op WHERE codigo_op = ANY(%s)) AS roteiro_operacoes,
                    (SELECT COUNT(*) FROM apontamentos_operacionais WHERE op = ANY(%s)) AS apontamentos,
                    (SELECT COUNT(*) FROM eventos_apontamento_operador evento
                     JOIN apontamentos_operacionais apontamento ON apontamento.id = evento.apontamento_id
                     WHERE apontamento.op = ANY(%s)) AS eventos,
                    (SELECT COUNT(*) FROM catalogo_sigmanest_planos_corte WHERE plano_hash LIKE 'seed:%%') AS planos_corte,
                    (SELECT COUNT(*) FROM catalogo_recursos_pcfactory) AS recursos_pcfactory,
                    (SELECT COUNT(*) FROM catalogo_status_recursos) AS status_recursos,
                    (SELECT COUNT(*) FROM operadores_apontamento WHERE ativo = TRUE) AS operadores_cracha,
                    (SELECT COUNT(*) FROM catalogo_status_recursos
                     WHERE habilitado = TRUE AND oculto = FALSE
                       AND COALESCE(grupo_codigo, '') <> '0001'
                       AND setup = FALSE AND retrabalho = FALSE) AS motivos_parada
                """,
                (selected_ops, selected_ops, selected_ops, selected_ops),
            )
            counts = dict(cursor.fetchone())
        return {
            **summary,
            "mode": "apply",
            "target": target,
            "users_created": user_count,
            "catalogs_loaded": {
                "resources": resource_count,
                "statuses": status_count,
                "operator_badges": badge_count,
            },
            "database_counts": counts,
        }
    finally:
        db.close()


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=DEFAULT_SOURCE)
    parser.add_argument("--status-source", type=Path, default=DEFAULT_STATUS_SOURCE)
    parser.add_argument("--resource-source", type=Path, default=DEFAULT_RESOURCE_SOURCE)
    parser.add_argument("--apply", action="store_true", help="Grava somente em TEST_DATABASE_URL.")
    parser.add_argument(
        "--with-users",
        action="store_true",
        help="Cria usuários setoriais usando SEED_TEST_PASSWORD.",
    )
    parser.add_argument("--report", type=Path, help="Salva também o resumo em JSON.")
    return parser.parse_args()


def main():
    args = parse_args()
    source = args.source.resolve()
    if not source.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {source}")
    status_source = args.status_source.resolve()
    resource_source = args.resource_source.resolve()
    if not status_source.is_file():
        raise FileNotFoundError(f"Catálogo de status não encontrado: {status_source}")
    if not resource_source.is_file():
        raise FileNotFoundError(f"Catálogo de recursos não encontrado: {resource_source}")
    operations, total_rows, valid_rows = read_source(source)
    statuses = read_status_catalog(status_source)
    resources = read_resource_catalog(resource_source)
    scenarios, routes = select_sample(operations)
    summary = build_summary(source, operations, total_rows, valid_rows, scenarios, routes)
    summary["pcfactory_catalogs"] = {
        "status_source": str(status_source),
        "statuses": len(statuses),
        "resource_source": str(resource_source),
        "resources": len(resources),
    }
    result = apply_seed(
        summary, scenarios, routes, statuses, resources, args.with_users
    ) if args.apply else {
        **summary,
        "mode": "dry-run",
        "message": "Nenhuma gravação foi realizada.",
    }
    serialized = json.dumps(result, ensure_ascii=False, indent=2, default=str)
    print(serialized)
    if args.report:
        args.report.parent.mkdir(parents=True, exist_ok=True)
        args.report.write_text(serialized + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()
