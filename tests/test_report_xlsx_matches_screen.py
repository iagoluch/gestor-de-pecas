"""X2: o .xlsx do relatório mostra os mesmos totais que a tela (JSON) do mesmo filtro.

A rota JSON e a rota .xlsx chamam o mesmo `facade.relatorio`; aqui um facade
falso devolve números distintivos e o teste lê o JSON e a planilha gerada.
Só se compara o que existe dos dois lados.
"""

import unittest
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from backend.api.dependencies.facade import get_frontend_facade
from backend.api.main import create_app
from tests.test_web_api import ApiFakeDatabase, _settings

PRODUCAO = {
    "availability": "disponivel", "good": 120, "scrap": 7, "rework": 3,
    "ops": 4, "resources": 2, "reason": "",
}
FTT = {"value": 95.5, "unit": "%", "availability": "disponivel"}
SETORES = [
    {"setor": "Corte", "producao_boa": 80, "refugo": 5, "retrabalho": 3, "tempo_produtivo_segundos": 3600},
    {"setor": "Dobra", "producao_boa": 40, "refugo": 2, "retrabalho": 0, "tempo_produtivo_segundos": 1800},
]
FILTRO = "?date_from=2026-10-01&date_to=2026-10-20"


class _Facade:
    def relatorio(self, report_type, filters):
        return {
            "production": PRODUCAO,
            "quality": {"ftt": FTT, "availability": "disponivel", "by_product": []},
            "sectors": SETORES,
            "orders": {"items": []},
        }


class ReportXlsxMatchesScreenTests(unittest.TestCase):
    def setUp(self):
        db = ApiFakeDatabase()
        db.criar_usuario("Gestor Web", "senha-segura", "gestor")
        self.app = create_app(settings=_settings(), database_factory=lambda: db)
        self.app.dependency_overrides[get_frontend_facade] = lambda: _Facade()
        self.client = TestClient(self.app)
        login = self.client.post("/api/v1/auth/login", json={"username": "Gestor Web", "password": "senha-segura"})
        self.assertEqual(login.status_code, 200, login.text)

    def tearDown(self):
        self.client.close()

    def _card(self, ws, label):
        # Cartão: rótulo em maiúsculas; o valor fica na linha de baixo, mesma coluna.
        for row in ws.iter_rows():
            for cell in row:
                if cell.value == label:
                    return ws.cell(row=cell.row + 1, column=cell.column).value
        self.fail(f"cartão {label!r} não encontrado na Visão Geral")

    def test_producao_xlsx_bate_com_a_tela(self):
        tela = self.client.get(f"/api/v1/reports/producao{FILTRO}")
        self.assertEqual(tela.status_code, 200, tela.text)
        tela = tela.json()
        xlsx = self.client.get(f"/api/v1/reports/producao/export.xlsx{FILTRO}")
        self.assertEqual(xlsx.status_code, 200, xlsx.text)
        wb = load_workbook(BytesIO(xlsx.content), data_only=False)
        try:
            visao = wb["Visão Geral"]
            p = tela["production"]
            for label, esperado in (
                ("PRODUÇÃO BOA", p["good"]),
                ("REFUGO", p["scrap"]),
                ("RETRABALHO", p["rework"]),
                ("OPS EXECUTADAS", p["ops"]),
                ("RECURSOS UTILIZADOS", p["resources"]),
            ):
                self.assertEqual(self._card(visao, label), esperado, label)
            # Percentual é gravado como fração (95,5% -> 0,955).
            self.assertAlmostEqual(self._card(visao, "FTT") * 100, tela["quality"]["ftt"]["value"], places=6)

            # Linha "Total" da aba Setores = soma dos setores que a tela mostra.
            ws = wb["Setores"]
            cabecalho = next(
                r for r in ws.iter_rows() if any(c.value == "Produção boa" for c in r)
            )
            colunas = {c.value: c.column for c in cabecalho if c.value}
            total = next(r for r in ws.iter_rows(min_row=cabecalho[0].row + 1) if r[0].value == "Total")
            for titulo, chave in (("Produção boa", "producao_boa"), ("Refugo", "refugo"), ("Retrabalho", "retrabalho")):
                soma = sum(s[chave] for s in tela["sectors"])
                self.assertEqual(total[colunas[titulo] - 1].value, soma, titulo)
            # E os totais da tela de setores fecham com os cartões.
            self.assertEqual(sum(s["producao_boa"] for s in tela["sectors"]), p["good"])
        finally:
            wb.close()


if __name__ == "__main__":
    unittest.main()
