"""Dois formatos de planilha: "tecnico" (completo) e "simples" (1 página para a diretoria).

O simples é só uma escolha do que já existe no técnico: mesma fonte, mesmos
números, nenhum recálculo. Estes testes garantem o formato (1 aba, 1 página,
texto de negócio, percentual legível) e que o técnico não perdeu nada.
"""

from __future__ import annotations

from datetime import datetime
from functools import partial
from io import BytesIO
import re
import tempfile
import unittest

from openpyxl import load_workbook

from backend.api.report_workbook import build_report_workbook
from mes.contracts import AnalyticsFilter, ReportRequest
from mes.services.industrial_reports import IndustrialReportService
from tests.fakes import FakeDatabase
from tests.test_intelligence_reports import NOW, CanonicalReportFacade
from tests.test_report_workbook_clean import (
    EXPECTED_EXECUTIVE_SHEETS,
    SECTION_TYPES,
    SNAKE_RE,
    UUID_RE,
    WORD_RE,
    RichCanonicalFacade,
    executive_payloads,
)


TECH_RE = re.compile(r"\b(backend|can[oô]nic[oa]s?|payload|json|api|null|None)\b", re.I)
PERCENT_FORMATS = {"0%", "0.00%", "0.000%"}
EXECUTIVE_TYPES = ("gerencial", "producao", "perdas", "indicadores", "dados_analiticos")
FILTROS = AnalyticsFilter(inicio=datetime(2026, 8, 1), fim=NOW, setor="Usinagem")
MAX_LINHAS_RESUMO = 60


class ModoWorkbookMixin:
    def _abrir(self, content: bytes):
        return load_workbook(BytesIO(content))

    def assert_resumo(self, content: bytes, label: str, *, com_grafico: bool = True) -> None:
        workbook = self._abrir(content)
        try:
            self.assertEqual(workbook.sheetnames, ["Resumo"], label)
            ws = workbook["Resumo"]
            self.assertEqual(ws.sheet_state, "visible", label)
            self.assertLessEqual(ws.max_row, MAX_LINHAS_RESUMO, f"{label}: não cabe numa página")
            self.assertEqual(ws.page_setup.orientation, "landscape", label)
            self.assertEqual(int(ws.page_setup.fitToWidth), 1, label)
            self.assertEqual(int(ws.page_setup.fitToHeight), 1, label)
            self.assertTrue(ws.sheet_properties.pageSetUpPr.fitToPage, label)
            self.assertEqual(ws.oddFooter.right.text, "Página &P de &N", label)
            self.assertIn("Período:", ws["A3"].value, label)
            self.assertIn("Emitido em:", ws["A3"].value, label)
            tipos = [type(grafico).__name__ for grafico in ws._charts]
            # Ranking em barras e, quando os valores somam um todo, a pizza ao lado.
            self.assertLessEqual(len(tipos), 2, f"{label}: gráficos demais")
            if com_grafico:
                self.assertIn("BarChart", tipos, f"{label}: ranking principal sem gráfico")
            if "PieChart" in tipos:
                self.assertEqual(len(tipos), 2, label)
                self.assertTrue(ws.conditional_formatting, f"{label}: pizza sem barras na tabela")
            self._assert_celulas(ws, label)
        finally:
            workbook.close()

    def _assert_celulas(self, ws, label: str) -> None:
        for row in ws.iter_rows():
            for cell in row:
                local = f"{label} {ws.title}!{cell.coordinate}: {cell.value!r}"
                if cell.data_type == "f":
                    self.fail(f"{local}: fórmula")
                if "%" in str(cell.number_format) and cell.value is not None:
                    self.assertIn(cell.number_format, PERCENT_FORMATS, local)
                    self.assertIsInstance(cell.value, (int, float), local)
                if not isinstance(cell.value, str):
                    continue
                texto = cell.value
                self.assertIsNone(TECH_RE.search(texto), local)
                self.assertIsNone(UUID_RE.search(texto), local)
                self.assertIsNone(SNAKE_RE.search(texto), local)
                self.assertIsNone(WORD_RE.search(texto), local)
                self.assertNotIn("{", texto, local)
                # Percentual escrito em frase: no máximo duas casas ("90,45%").
                self.assertIsNone(re.search(r"\d[.,]\d{3,}\s*%", texto), local)


class ModoValidacaoTests(unittest.TestCase):
    def test_modo_invalido_levanta_value_error(self):
        with self.assertRaises(ValueError):
            build_report_workbook("gerencial", {}, FILTROS, generated_at=NOW, modo="executivo")

    def test_padrao_e_tecnico(self):
        payload = executive_payloads()["gerencial"]
        padrao = load_workbook(BytesIO(build_report_workbook("gerencial", payload, FILTROS, generated_at=NOW)))
        tecnico = load_workbook(BytesIO(
            build_report_workbook("gerencial", payload, FILTROS, generated_at=NOW, modo="tecnico")
        ))
        try:
            self.assertEqual(padrao.sheetnames, tecnico.sheetnames)
        finally:
            padrao.close()
            tecnico.close()


class ExecutivoModosTests(ModoWorkbookMixin, unittest.TestCase):
    def test_simples_tem_uma_aba_resumo_de_uma_pagina(self):
        for report_type, payload in executive_payloads().items():
            with self.subTest(report_type=report_type):
                content = build_report_workbook(report_type, payload, FILTROS, generated_at=NOW, modo="simples")
                self.assert_resumo(content, report_type)

    def test_simples_com_payload_vazio_mostra_traco_e_nunca_zero(self):
        for report_type in (*EXECUTIVE_TYPES, "desconhecido"):
            with self.subTest(report_type=report_type):
                content = build_report_workbook(report_type, {}, FILTROS, generated_at=NOW, modo="simples")
                self.assert_resumo(content, report_type, com_grafico=False)
                workbook = self._abrir(content)
                try:
                    numeros = [
                        cell.value for row in workbook["Resumo"].iter_rows() for cell in row
                        if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
                    ]
                    self.assertEqual(numeros, [], f"{report_type}: dado ausente virou número")
                finally:
                    workbook.close()

    def test_simples_nao_tem_aba_de_qualidade_dos_dados(self):
        for report_type, payload in executive_payloads().items():
            with self.subTest(report_type=report_type):
                workbook = self._abrir(
                    build_report_workbook(report_type, payload, FILTROS, generated_at=NOW, modo="simples")
                )
                try:
                    self.assertNotIn("Qualidade dos Dados", workbook.sheetnames)
                finally:
                    workbook.close()

    def test_simples_repete_os_numeros_do_tecnico(self):
        payload = executive_payloads()["gerencial"]
        tecnico = self._abrir(build_report_workbook("gerencial", payload, FILTROS, generated_at=NOW))
        simples = self._abrir(build_report_workbook("gerencial", payload, FILTROS, generated_at=NOW, modo="simples"))
        try:
            numeros = lambda ws: {
                cell.value for row in ws.iter_rows() for cell in row
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            }
            # Todo número do resumo já existe no painel técnico.
            self.assertTrue(numeros(simples["Resumo"]) <= numeros(tecnico["Visão Geral"]) | numeros(tecnico["Setores"]))
        finally:
            tecnico.close()
            simples.close()

    def test_tecnico_mantem_todas_as_abas_de_antes(self):
        for report_type, payload in executive_payloads().items():
            with self.subTest(report_type=report_type):
                workbook = self._abrir(build_report_workbook(report_type, payload, FILTROS, generated_at=NOW))
                try:
                    abas = set(workbook.sheetnames)
                    esperado = set(EXPECTED_EXECUTIVE_SHEETS.get(report_type, ["Visão Geral"]))
                    self.assertTrue(esperado <= abas, f"{report_type}: faltou {esperado - abas}")
                    self.assertTrue(all(ws.sheet_state == "visible" for ws in workbook.worksheets))
                finally:
                    workbook.close()

    def test_tecnico_nao_corta_listas_longas(self):
        payload = executive_payloads()["gerencial"]
        overview = dict(payload["overview"])
        insights = dict(overview.get("insights") or {})
        insights["losses"] = [
            {"cause": f"Motivo {indice}", "impact_value": float(100 + indice)} for indice in range(30)
        ]
        overview["insights"] = insights
        workbook = self._abrir(build_report_workbook(
            "gerencial", {**payload, "overview": overview}, FILTROS, generated_at=NOW,
        ))
        try:
            textos = {cell.value for row in workbook["Perdas"].iter_rows() for cell in row if isinstance(cell.value, str)}
            self.assertTrue({f"Motivo {indice}" for indice in range(30)} <= textos)
        finally:
            workbook.close()


class SecoesModosTests(ModoWorkbookMixin, unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.temp.cleanup()

    def _generate(self, facade, report_type, modo):
        service = IndustrialReportService(
            facade, FakeDatabase(), artifact_dir=self.temp.name,
            workbook_builder=partial(build_report_workbook, modo=modo), now_func=lambda: NOW,
        )
        request = ReportRequest(
            report_type=report_type, inicio=datetime(2026, 8, 1), fim=NOW,
            op="PCMIXQ01001" if report_type == "rastreabilidade" else None,
        )
        artifact = service.generate(request, created_by=7, source="manual")
        _record, path = service.get(artifact["id"], user_id=7)
        return path.read_bytes()

    def test_simples_de_cada_secao_e_um_resumo(self):
        for facade in (RichCanonicalFacade(), CanonicalReportFacade()):
            for report_type in SECTION_TYPES:
                with self.subTest(facade=type(facade).__name__, report_type=report_type):
                    self.assert_resumo(self._generate(facade, report_type, "simples"), report_type)

    def test_tecnico_das_secoes_continua_com_visao_geral_primeiro(self):
        for report_type in SECTION_TYPES:
            with self.subTest(report_type=report_type):
                workbook = self._abrir(self._generate(RichCanonicalFacade(), report_type, "tecnico"))
                try:
                    self.assertEqual(workbook.sheetnames[0], "Visão Geral")
                finally:
                    workbook.close()


if __name__ == "__main__":
    unittest.main()
