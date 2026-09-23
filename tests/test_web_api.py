import base64
import json
import tempfile
import unittest
from contextlib import contextmanager
from dataclasses import replace
from pathlib import Path
from datetime import date, datetime, time
from io import BytesIO
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.database.schema import SCHEMA_VERSION
from backend.api.config import WebSettings
from backend.api.main import create_app
from mes.services.operator_flow import OperatorFlowService
from tests.fakes import FakeDatabase
from tests.helpers import liberar_primeira_peca


class _Cursor:
    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, traceback):
        return False

    def execute(self, _query, _params=None):
        return None

    def fetchone(self):
        return {"ok": 1}


class _Connection:
    def cursor(self):
        return _Cursor()


class ApiFakeDatabase(FakeDatabase):
    @contextmanager
    def connection(self):
        yield _Connection()


def _settings():
    return WebSettings(
        environment="test",
        session_secret="test-secret-that-is-long-enough-for-session-signatures-123456",
        allowed_hosts=("testserver",),
        allowed_origins=("http://testserver",),
        session_ttl_seconds=3600,
    )


class WebApiTests(unittest.TestCase):
    def setUp(self):
        self.db = ApiFakeDatabase()
        self.db.criar_usuario("Gestor Web", "senha-segura", "gestor")
        self.db.criar_usuario("Andon Web", "senha-andon", "andon")
        self.db.criar_usuario("Operador Web", "senha-operador", "operador_dobra")
        self.db.criar_usuario("Operador Corte Web", "senha-corte", "operador_corte")
        self.db.criar_usuario("Operador Destaque Web", "senha-destaque", "operador_destaque")
        # Wave 6F — dois operadores no mesmo perfil de estação (a estação vem do
        # login), que é o cenário real de disputa pelo posto.
        self.db.criar_usuario("Soldador Web A", "senha-solda-a", "estacao1aco")
        self.db.criar_usuario("Soldador Web B", "senha-solda-b", "estacao1aco")
        task_id = self.db.inserir_tarefa("T-WEB", material="AÇO 304", espessura=3)
        self.db.inserir_op_na_tarefa(task_id, "OP-WEB", "PECA-WEB", "Dobra", 2)
        self.db.catalog_operations.append({
            "id": 101,
            "codigo_op": "OP-WEB",
            "numero_operacao": "20",
            "codigo_recurso": "DOBRA3",
            "descricao_operacao": "DOBRA",
            "tipo_setor": "Dobra",
            "produto_codigo": "PECA-WEB",
            "produto_descricao": "Peça de teste Web",
            "quantidade": 2,
            "ordem": 1,
            "ativo": True,
        })
        for suffix, operation_id in (("A", 201), ("B", 202)):
            solda_task_id = self.db.inserir_tarefa(f"T-SOLDA-{suffix}")
            self.db.inserir_op_na_tarefa(solda_task_id, f"OP-SOLDA-{suffix}", f"PECA-SOLDA-{suffix}", "Solda", 1)
            self.db.catalog_operations.append({
                "id": operation_id,
                "codigo_op": f"OP-SOLDA-{suffix}",
                "numero_operacao": "30",
                "codigo_recurso": "SOLDA4",
                "descricao_operacao": "SOLDAGEM",
                "tipo_setor": "Solda Aço",
                "produto_codigo": f"PECA-SOLDA-{suffix}",
                "produto_descricao": "Peça de solda Web",
                "quantidade": 1,
                "ordem": 1,
                "ativo": True,
            })
        self.db.cut_plans.append({
            "plano_hash": "hash-web-corte-1",
            "codigo_tarefa": "T-CORTE-WEB",
            "programa": "P-WEB-01",
            "nome_chapa": "CHAPA-WEB",
            "sequencia_nesting": 1,
            "material": "AÇO 304",
            "espessura": 3.0,
            "maquina_sigmanest": "AMADA_ENSIS",
            "quantidade_processo": 4,
            "tempo_previsto_segundos": 120,
            "data_programa": date.today(),
            "area_usada": 1000,
            "fracao_sucata": 0.1,
            "ativo": True,
        })
        self.app = create_app(settings=_settings(), database_factory=lambda: self.db)
        self.client = TestClient(self.app)

    def tearDown(self):
        self.client.close()

    def login_manager(self):
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": "Gestor Web", "password": "senha-segura"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response

    def login_operator(self, username="Operador Web", password="senha-operador"):
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": username, "password": password},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response

    def login_andon(self):
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": "Andon Web", "password": "senha-andon"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        return response

    def csrf(self):
        return {"X-CSRF-Token": self.client.cookies.get("gestor_csrf")}

    def _concluir_plano_laser_da_tarefa(self, codigo_tarefa):
        """Corta no Laser a única chapa da tarefa, liberando o Destaque."""

        plano_hash = f"hash-laser-{codigo_tarefa}"
        self.db.cut_plans.append({
            "plano_hash": plano_hash,
            "codigo_tarefa": codigo_tarefa,
            "programa": "P-LASER-01",
            "nome_chapa": "CHAPA-LASER",
            "sequencia_nesting": 1,
            "sigmanest_repeat_id": 1,
            "material": "AÇO 304",
            "espessura": 3.0,
            "maquina_sigmanest": "AMADA_ENSIS",
            "quantidade_processo": 2,
            "tempo_previsto_segundos": 120,
            "data_programa": date.today(),
            "ativo": True,
        })
        started = self.db.iniciar_apontamento_corte(
            plano_hash, "Laser Ensis 3015", "Operador Corte Web", "2026-01-01"
        )
        self.assertIsNotNone(started)
        self.assertIsNotNone(
            self.db.finalizar_apontamento_corte(started["id"], "Operador Corte Web")
        )
        return plano_hash

    def test_configuracao_de_simulacao_exige_relogio_deterministico(self):
        base = {
            "GESTOR_WEB_ENV": "test",
            "GESTOR_WEB_SESSION_SECRET": "test-secret-that-is-long-enough-for-session-signatures-123456",
            "GESTOR_SIMULATION_MODE": "true",
        }
        with self.assertRaisesRegex(RuntimeError, "GESTOR_SIMULATION_NOW"):
            WebSettings.from_env(base)
        configured = WebSettings.from_env({
            **base,
            "GESTOR_SIMULATION_NOW": "2026-08-24T08:32:00",
        })
        self.assertTrue(configured.simulation_mode)
        self.assertEqual(configured.simulation_reference_time.isoformat(), "2026-08-24T08:32:00")

    def test_filtro_web_na_simulacao_nao_ultrapassa_relogio_de_referencia(self):
        self.db.safe_target = {"dbname": "gestor_pecas_test_simulacao"}
        settings = replace(
            _settings(),
            simulation_mode=True,
            simulation_reference_time=datetime(2026, 8, 24, 8, 32),
        )
        app = create_app(settings=settings, database_factory=lambda: self.db)
        with TestClient(app) as client:
            login = client.post(
                "/api/v1/auth/login",
                json={"username": "Gestor Web", "password": "senha-segura"},
            )
            self.assertEqual(login.status_code, 200, login.text)
            response = client.get(
                "/api/v1/management/overview",
                params={
                    "inicio": "2026-08-01T00:00:00",
                    "fim": "2026-08-24T23:59:59",
                },
            )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.json()["periodo"]["fim"], "2026-08-24T08:32:00")

    def test_andon_filtra_estado_aberto_pelo_relogio_de_referencia(self):
        reference = datetime(2026, 8, 24, 8, 32)
        received = []
        self.db.safe_target = {"dbname": "gestor_pecas_test_simulacao"}
        self.db.listar_estados_recurso_atuais = lambda **kwargs: received.append(kwargs) or []
        settings = replace(
            _settings(),
            simulation_mode=True,
            simulation_reference_time=reference,
        )
        app = create_app(settings=settings, database_factory=lambda: self.db)
        with TestClient(app) as client:
            login = client.post(
                "/api/v1/auth/login",
                json={"username": "Gestor Web", "password": "senha-segura"},
            )
            self.assertEqual(login.status_code, 200, login.text)
            response = client.get("/api/v1/andon")

        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(received)
        self.assertEqual(received[0]["reference_time"], reference)

    def test_capabilities_expoem_backend_only_sem_banco(self):
        response = self.client.get("/api/v1/system/capabilities")
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["calculation_policy"], "backend_only")
        self.assertFalse(payload["ui_policy"]["management_correction_or_edit_tab"])
        self.assertTrue(payload["api"]["operator_web_enabled"])
        self.assertEqual(payload["active_data_source"], "postgresql_test_only")
        self.assertEqual(
            payload["corporate_integration"]["provider"],
            "totvs_production_order_push_v1",
        )
        self.assertFalse(payload["corporate_integration"]["configured"])
        self.assertFalse(payload["corporate_integration"]["planning_read_enabled"])
        self.assertFalse(payload["corporate_integration"]["execution_write_enabled"])

    def test_health_nao_expoe_dsn(self):
        response = self.client.get("/api/v1/system/health")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["schema_version"], SCHEMA_VERSION)
        serialized = json.dumps(response.json()).casefold()
        self.assertNotIn("password", serialized)
        self.assertNotIn("database_url", serialized)

    def test_respostas_levam_headers_de_seguranca(self):
        response = self.client.get("/api/v1/system/health")
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")
        self.assertIn("frame-ancestors 'none'", response.headers["Content-Security-Policy"])
        self.assertIn("camera=()", response.headers["Permissions-Policy"])
        # Implantação sem HTTPS declarado (cookie_secure=False) não recebe HSTS.
        self.assertNotIn("Strict-Transport-Security", response.headers)

    def test_hsts_so_quando_implantacao_declara_https(self):
        settings = replace(_settings(), cookie_secure=True)
        client = TestClient(create_app(settings=settings, database_factory=lambda: self.db))
        self.addCleanup(client.close)
        response = client.get("/api/v1/system/health")
        hsts = response.headers["Strict-Transport-Security"]
        self.assertIn("max-age=", hsts)
        self.assertNotIn("includeSubDomains", hsts)

    def test_login_usa_cookie_http_only_e_nao_retorna_senha(self):
        response = self.login_manager()
        self.assertIn("httponly", response.headers["set-cookie"].casefold())
        self.assertNotIn("senha-segura", response.text)
        self.assertEqual(response.json()["role"], "gestor")
        session = self.client.get("/api/v1/auth/session")
        self.assertEqual(session.status_code, 200)
        self.assertTrue(session.json()["management_access"])

    def test_credencial_invalida_tem_erro_padronizado(self):
        response = self.client.post(
            "/api/v1/auth/login",
            json={"username": "Gestor Web", "password": "incorreta"},
        )
        self.assertEqual(response.status_code, 401)
        self.assertEqual(response.json()["code"], "invalid_credentials")
        self.assertIn("request_id", response.json())
        self.assertNotIn("traceback", response.text.casefold())

    def test_operador_nao_acessa_area_gerencial(self):
        self.client.post(
            "/api/v1/auth/login",
            json={"username": "Operador Web", "password": "senha-operador"},
        )
        response = self.client.get("/api/v1/management/overview")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["code"], "management_access_denied")

    def test_sessao_operador_expoe_somente_setor_e_recursos_do_perfil(self):
        response = self.login_operator()
        payload = response.json()
        self.assertEqual(payload["operator_sector"], "Dobra")
        self.assertEqual(payload["operator_resources"], ["1303", "2204", "Gasparini"])
        self.assertFalse(payload["management_access"])
        context = self.client.get("/api/v1/operator/context")
        self.assertEqual(context.status_code, 200, context.text)
        self.assertEqual(context.json()["workflow"], "workbench")

    def test_operador_nao_pode_consultar_ou_apontar_recurso_de_outro_setor(self):
        self.login_operator()
        denied = self.client.get("/api/v1/operator/workbench", params={"resource": "Estação 1"})
        self.assertEqual(denied.status_code, 403)
        command = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={
                "action": "Início",
                "resource": "Estação 1",
                "op": "OP-WEB",
                "operation_id": 101,
            },
        )
        self.assertEqual(command.status_code, 403)
        self.assertEqual(command.json()["code"], "operator_resource_denied")

    def test_fluxo_web_operador_reutiliza_servico_transacional_e_csrf(self):
        self.login_operator()
        operations = self.client.get("/api/v1/operator/operations/OP-WEB")
        self.assertEqual(operations.status_code, 200, operations.text)
        self.assertEqual(operations.json()["items"][0]["id"], 101)
        payload = {"action": "Início", "resource": "1303", "op": "OP-WEB", "operation_id": 101}
        no_csrf = self.client.post("/api/v1/operator/actions", json=payload)
        self.assertEqual(no_csrf.status_code, 403)

        # Wave 6B — o Iniciar é livre; o portão Setup/Qualidade é do Finalizar.
        primeira_peca = {"resource": "1303", "op": "OP-WEB", "operation_id": 101}
        self.db.salvar_template_qualidade(
            "PECA-WEB", [{"sequencia": 1, "descricao": "Altura", "padrao": "12,0 +/- 0,2"}]
        )
        started = self.client.post(
            "/api/v1/operator/actions", headers=self.csrf(), json=payload
        )
        self.assertEqual(started.status_code, 200, started.text)

        # Antes da conferência, o Finalizar é recusado e a mensagem manda o
        # operador apontar o Setup — o popup não abre por aqui.
        sem_conferencia = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Finalizado", "good": 2, "badges": ["1"]},
        )
        self.assertEqual(sem_conferencia.status_code, 409, sem_conferencia.text)
        self.assertEqual(
            sem_conferencia.json()["code"], "primeira_peca_gate_obrigatorio"
        )
        self.assertIn("Aponte o Setup", sem_conferencia.json()["message"])

        estado = self.client.get("/api/v1/operator/first-piece", params=primeira_peca)
        self.assertEqual(estado.status_code, 200, estado.text)
        self.assertTrue(estado.json()["gate_estruturado"])
        cotas = estado.json()["checklist"]["template"]["cotas"]
        self.assertEqual([item["sequencia"] for item in cotas], [1])
        self.assertEqual(cotas[0]["limite_inferior"], 11.8)
        self.assertEqual(cotas[0]["limite_superior"], 12.2)

        # O popup não confirma Setup: quem grava esse fato é o botão Setup.
        sem_setup = self.client.post(
            "/api/v1/operator/first-piece",
            headers=self.csrf(),
            json={
                **primeira_peca,
                "action": "checklist",
                "measures": [{"sequencia": 1, "medida": "12,1"}],
            },
        )
        self.assertEqual(sem_setup.status_code, 409, sem_setup.text)
        self.assertEqual(sem_setup.json()["code"], "primeira_peca_setup_pendente")

        setup_apontado = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Setup"},
        )
        self.assertEqual(setup_apontado.status_code, 200, setup_apontado.text)
        self.assertEqual(setup_apontado.json()["data"]["status"], "Setup")
        self.assertEqual(
            self.client.post(
                "/api/v1/operator/actions", headers=self.csrf(), json=payload
            ).status_code,
            200,
        )

        checklist = self.client.post(
            "/api/v1/operator/first-piece",
            headers=self.csrf(),
            json={
                **primeira_peca,
                "action": "checklist",
                "measures": [{"sequencia": 1, "medida": "12,1"}],
            },
        )
        self.assertEqual(checklist.status_code, 200, checklist.text)
        self.assertTrue(checklist.json()["data"]["liberado"])
        self.assertIn("Lote liberado", checklist.json()["message"])

        setup = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Setup"},
        )
        self.assertEqual(setup.status_code, 200, setup.text)
        resumed = self.client.post("/api/v1/operator/actions", headers=self.csrf(), json=payload)
        self.assertEqual(resumed.status_code, 200, resumed.text)

        stopped = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Parada", "stop_reason_code": "0029"},
        )
        self.assertEqual(stopped.status_code, 200, stopped.text)
        resumed = self.client.post("/api/v1/operator/actions", headers=self.csrf(), json=payload)
        self.assertEqual(resumed.status_code, 200, resumed.text)
        # Wave 4: OP-WEB prevê 2 peças e boas + refugo têm esse teto. Uma peça
        # boa é parcial e devolve a OP à fila com o saldo acumulado.
        partial = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Finalizado", "good": 1, "scrap": 0, "badges": ["1"]},
        )
        self.assertEqual(partial.status_code, 200, partial.text)
        self.assertEqual(partial.json()["code"], "finalizacao_parcial")
        self.assertEqual(self.db.appointments[0]["quantidade_boa"], 1)
        self.assertEqual(self.db.appointments[0]["status"], "Aguardando")

        reopened = self.client.post("/api/v1/operator/actions", headers=self.csrf(), json=payload)
        self.assertEqual(reopened.status_code, 200, reopened.text)
        # O refugo fecha o saldo sem virar peça boa em nenhum ponto do caminho.
        finished = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Finalizado", "good": 0, "scrap": 1, "badges": ["1"]},
        )
        self.assertEqual(finished.status_code, 200, finished.text)
        self.assertEqual(self.db.appointments[0]["quantidade_boa"], 1)
        self.assertEqual(self.db.appointments[0]["quantidade_refugo"], 1)
        self.assertEqual(self.db.appointments[0]["status"], "Finalizado")

        excedente = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**payload, "action": "Finalizado", "good": 1, "scrap": 0, "badges": ["1"]},
        )
        self.assertEqual(excedente.status_code, 409)

    def test_setup_exige_inicio_e_parada_pode_ser_sem_op(self):
        self.login_operator()
        setup = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={"action": "Setup", "resource": "1303", "op": "OP-WEB", "operation_id": 101},
        )
        self.assertEqual(setup.status_code, 409, setup.text)
        self.assertEqual(setup.json()["code"], "setup_exige_inicio")

        stopped = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={"action": "Parada", "resource": "2204", "stop_reason_code": "0029"},
        )
        self.assertEqual(stopped.status_code, 200, stopped.text)
        self.assertEqual(stopped.json()["code"], "parada_recurso_sem_op")
        self.assertIsNone(stopped.json()["data"].get("op"))

        # A parada sem OP precisa ter retomada sem OP: sem ela o recurso ficava
        # travado, porque nenhuma outra ação é aceita sem informar uma OP.
        workbench = self.client.get("/api/v1/operator/workbench?resource=2204")
        self.assertEqual(workbench.status_code, 200, workbench.text)
        self.assertEqual(workbench.json()["resource_state"]["categoria"], "parada")

        resumed = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={"action": "Retomar", "resource": "2204"},
        )
        self.assertEqual(resumed.status_code, 200, resumed.text)
        self.assertEqual(resumed.json()["code"], "retomada_recurso_sem_op")
        self.assertEqual(
            self.client.get("/api/v1/operator/workbench?resource=2204").json()["resource_state"]["categoria"],
            "fila",
        )

    def test_backend_rejeita_transicao_invalida_sem_depender_do_frontend(self):
        self.login_operator()
        base = {
            "resource": "1303",
            "op": "OP-WEB",
            "operation_id": 101,
        }
        # Wave 6B — o portão Setup/Qualidade vem antes de produzir em Dobra.
        liberar_primeira_peca(
            self.db,
            "Operador Web",
            op="OP-WEB",
            setor="Dobra",
            recurso="1303",
            operacao=self.db.catalog_operations[0],
        )
        started = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Início"},
        )
        self.assertEqual(started.status_code, 200, started.text)
        stopped = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Parada", "stop_reason_code": "0029"},
        )
        self.assertEqual(stopped.status_code, 200, stopped.text)

        invalid = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Setup"},
        )
        self.assertEqual(invalid.status_code, 409, invalid.text)
        self.assertEqual(invalid.json()["code"], "transicao_invalida")

    def test_api_preserva_retrabalho_apos_parada_e_setup(self):
        self.login_operator()
        base = {
            "resource": "1303",
            "op": "OP-WEB",
            "operation_id": 101,
        }
        rework = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Retrabalho"},
        )
        self.assertEqual(rework.status_code, 200, rework.text)

        stopped = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Parada", "stop_reason_code": "0029"},
        )
        self.assertEqual(stopped.status_code, 200, stopped.text)
        resumed = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Retomar"},
        )
        self.assertEqual(resumed.status_code, 200, resumed.text)
        self.assertEqual(resumed.json()["data"]["status"], "Retrabalho")

        setup = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Setup"},
        )
        self.assertEqual(setup.status_code, 200, setup.text)
        returned = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Retornar"},
        )
        self.assertEqual(returned.status_code, 200, returned.text)
        self.assertEqual(returned.json()["data"]["status"], "Retrabalho")

    def test_api_rejeita_retrabalho_diretamente_para_producao(self):
        self.login_operator()
        base = {
            "resource": "1303",
            "op": "OP-WEB",
            "operation_id": 101,
        }
        rework = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Retrabalho"},
        )
        self.assertEqual(rework.status_code, 200, rework.text)

        invalid = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={**base, "action": "Início"},
        )

        self.assertEqual(invalid.status_code, 409, invalid.text)
        self.assertEqual(invalid.json()["code"], "transicao_invalida")
        self.assertEqual(self.db.appointments[0]["status"], "Retrabalho")

    def test_solda_rejeita_setup_tambem_na_api(self):
        self.login_operator("Soldador Web A", "senha-solda-a")
        response = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={"action": "Setup", "resource": "Estação 1", "op": "OP-SOLDA-A", "operation_id": 201},
        )
        self.assertEqual(response.status_code, 409, response.text)
        self.assertEqual(response.json()["code"], "setup_indisponivel_setor")

    def test_fluxo_web_corte_preserva_fila_automatica_e_nesting(self):
        self.login_operator("Operador Corte Web", "senha-corte")
        queue = self.client.get("/api/v1/cutting/queue", params={"resource": "Laser Ensis 3015"})
        self.assertEqual(queue.status_code, 200, queue.text)
        self.assertEqual(queue.json()["items"][0]["status"], "Aguardando")
        self.assertEqual(queue.json()["items"][0]["nestings"][0]["tempo_previsto_segundos"], 120)
        filtered = self.client.get(
            "/api/v1/cutting/queue",
            params={"resource": "Laser Ensis 3015", "search": "hash-web-corte-1"},
        )
        self.assertEqual(filtered.status_code, 200, filtered.text)
        self.assertEqual(len(filtered.json()["items"]), 1)
        missing = self.client.get(
            "/api/v1/cutting/queue",
            params={"resource": "Laser Ensis 3015", "search": "plano-inexistente"},
        )
        self.assertEqual(missing.json()["items"], [])
        started = self.client.post(
            "/api/v1/cutting/actions",
            headers=self.csrf(),
            json={"action": "Início", "resource": "Laser Ensis 3015", "plan_hash": "hash-web-corte-1"},
        )
        self.assertEqual(started.status_code, 200, started.text)
        appointment_id = started.json()["data"]["apontamento_ids_em_processo"][0]
        stopped = self.client.post(
            "/api/v1/cutting/actions",
            headers=self.csrf(),
            json={"action": "Parada", "resource": "Laser Ensis 3015", "stop_reason_code": "0029"},
        )
        self.assertEqual(stopped.status_code, 200, stopped.text)
        resumed = self.client.post(
            "/api/v1/cutting/actions",
            headers=self.csrf(),
            json={"action": "Retomada", "resource": "Laser Ensis 3015"},
        )
        self.assertEqual(resumed.status_code, 200, resumed.text)
        finished = self.client.post(
            "/api/v1/cutting/actions",
            headers=self.csrf(),
            json={"action": "Finalizado", "resource": "Laser Ensis 3015", "appointment_id": appointment_id},
        )
        self.assertEqual(finished.status_code, 200, finished.text)
        self.assertEqual(self.db.cut_appointments[0]["status"], "Finalizado")

    def test_fluxo_web_destaque_exige_cracha_no_fim(self):
        # O Destaque só é liberado por plano de Laser concluído no Corte: sem
        # essa chapa a tarefa nem entra na fila do posto.
        self._concluir_plano_laser_da_tarefa("T-WEB")
        self.login_operator("Operador Destaque Web", "senha-destaque")
        task = self.client.get("/api/v1/highlight/tasks/T-WEB")
        self.assertEqual(task.status_code, 200, task.text)
        started = self.client.post(
            "/api/v1/highlight/actions",
            headers=self.csrf(),
            json={"action": "Início", "task_code": "T-WEB"},
        )
        self.assertEqual(started.status_code, 200, started.text)
        self.assertEqual(started.json()["current"]["timing"]["current_mode"], "execucao")
        invalid = self.client.post(
            "/api/v1/highlight/actions",
            headers=self.csrf(),
            json={"action": "Fim", "task_code": "T-WEB", "badge": "inexistente"},
        )
        self.assertEqual(invalid.status_code, 409)
        finished = self.client.post(
            "/api/v1/highlight/actions",
            headers=self.csrf(),
            json={"action": "Fim", "task_code": "T-WEB", "badge": "1"},
        )
        self.assertEqual(finished.status_code, 200, finished.text)
        self.assertEqual(self.db.buscar_tarefa_por_codigo("T-WEB")["status"], "Finalizado")

    def test_destaque_retoma_parada_registrada_sem_tarefa(self):
        """A parada sem tarefa é do posto — e precisa de retomada sem tarefa.

        Sem ela o Destaque ficava parado para sempre: toda ação exigia uma
        tarefa que a parada física não possui.
        """

        self.login_operator("Operador Destaque Web", "senha-destaque")
        stopped = self.client.post(
            "/api/v1/highlight/actions",
            headers=self.csrf(),
            json={"action": "Parada", "stop_reason_code": "0029"},
        )
        self.assertEqual(stopped.status_code, 200, stopped.text)
        self.assertEqual(stopped.json()["code"], "parada_recurso_sem_op")
        self.assertEqual(
            self.client.get("/api/v1/highlight/queue").json()["resource_state"]["categoria"],
            "parada",
        )

        # A tarefa parada continua sendo retomada pelo Início do destaque.
        escopo_invalido = self.client.post(
            "/api/v1/highlight/actions",
            headers=self.csrf(),
            json={"action": "Retomar", "task_code": "T-WEB"},
        )
        self.assertEqual(escopo_invalido.status_code, 409, escopo_invalido.text)
        self.assertEqual(escopo_invalido.json()["code"], "highlight_resume_task_scoped")

        resumed = self.client.post(
            "/api/v1/highlight/actions",
            headers=self.csrf(),
            json={"action": "Retomar"},
        )
        self.assertEqual(resumed.status_code, 200, resumed.text)
        self.assertEqual(resumed.json()["code"], "retomada_recurso_sem_op")
        self.assertEqual(
            self.client.get("/api/v1/highlight/queue").json()["resource_state"]["categoria"],
            "fila",
        )

    def test_solda_bloqueia_estacao_ocupada_tambem_no_servidor(self):
        self.login_operator("Soldador Web A", "senha-solda-a")
        started = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={"action": "Início", "resource": "Estação 1", "op": "OP-SOLDA-A", "operation_id": 201},
        )
        self.assertEqual(started.status_code, 200, started.text)

        self.login_operator("Soldador Web B", "senha-solda-b")
        stations = self.client.get("/api/v1/operator/stations")
        station = next(item for item in stations.json()["items"] if item["resource"] == "Estação 1")
        self.assertEqual(station["status"], "Ocupada")
        self.assertEqual(station["operator"], "Soldador Web A")
        denied = self.client.post(
            "/api/v1/operator/actions",
            headers=self.csrf(),
            json={"action": "Início", "resource": "Estação 1", "op": "OP-SOLDA-B", "operation_id": 202},
        )
        self.assertEqual(denied.status_code, 409, denied.text)
        self.assertEqual(denied.json()["code"], "operator_resource_occupied")

    def test_rotas_gerenciais_consumem_facade_e_preservam_ausencia_real_de_dados(self):
        self.login_manager()
        overview = self.client.get("/api/v1/management/overview")
        self.assertEqual(overview.status_code, 200, overview.text)
        self.assertIsNone(overview.json()["kpis"]["oee"]["value"])
        self.assertEqual(
            overview.json()["kpis"]["oee"]["availability"],
            "sem_registros",
        )
        analytics = self.client.get("/api/v1/analytics/oee")
        self.assertEqual(analytics.status_code, 200)
        self.assertIsNone(analytics.json()["value"])
        self.assertEqual(
            analytics.json()["evolution"]["reason"],
            "Dados insuficientes para exibir a evolução do OEE no período selecionado.",
        )

    def test_api_expoe_dashboard_de_excecoes_e_explicacao_de_kpi(self):
        self.login_manager()
        insights = self.client.get("/api/v1/management/insights")
        self.assertEqual(insights.status_code, 200, insights.text)
        payload = insights.json()
        self.assertEqual(
            set(payload["kpi_explanations"]),
            {"oee", "availability", "performance", "ftt"},
        )
        self.assertEqual(payload["policies"]["calculation"], "backend_only")

        explanation = self.client.get("/api/v1/management/kpis/oee/explanation")
        self.assertEqual(explanation.status_code, 200, explanation.text)
        self.assertEqual(explanation.json()["key"], "oee")
        self.assertIn("components", explanation.json())
        self.assertIn("evidence", explanation.json())

        unknown = self.client.get("/api/v1/management/kpis/inventado/explanation")
        self.assertEqual(unknown.status_code, 404, unknown.text)
        self.assertEqual(unknown.json()["code"], "unknown_management_kpi")

    def test_operador_nao_acessa_insights_gerenciais(self):
        self.login_operator()
        response = self.client.get("/api/v1/management/insights")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["code"], "management_access_denied")

    def test_filtros_invalidos_e_paginacao_sao_servidor(self):
        self.login_manager()
        invalid = self.client.get(
            "/api/v1/orders",
            params={"inicio": "2026-08-20T10:00:00", "fim": "2026-08-20T09:00:00"},
        )
        self.assertEqual(invalid.status_code, 400)
        self.assertEqual(invalid.json()["code"], "invalid_period")
        page = self.client.get("/api/v1/orders", params={"page": 1, "page_size": 10})
        self.assertEqual(page.status_code, 200)
        self.assertEqual(page.json()["page"]["page_size"], 10)

    def test_data_implausivel_rejeitada_sem_500(self):
        self.login_manager()
        response = self.client.get(
            "/api/v1/audit/appointments",
            params={"fim": "0263-10-17T17:14:48"},
        )
        self.assertEqual(response.status_code, 422, response.text)
        self.assertEqual(response.json()["code"], "invalid_date")

    def test_logout_exige_csrf_e_expira_cookies(self):
        self.login_manager()
        denied = self.client.post("/api/v1/auth/logout")
        self.assertEqual(denied.status_code, 403)
        csrf = self.client.cookies.get("gestor_csrf")
        response = self.client.post("/api/v1/auth/logout", headers={"X-CSRF-Token": csrf})
        self.assertEqual(response.status_code, 204)
        self.assertEqual(self.client.get("/api/v1/auth/session").status_code, 401)

    def test_relatorio_csv_e_utf8(self):
        self.login_manager()
        response = self.client.get("/api/v1/reports/producao/export.csv")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue(response.content.startswith(b"\xef\xbb\xbf"))
        self.assertIn("attachment", response.headers["content-disposition"])

    def test_relatorio_excel_formatado(self):
        self.login_manager()
        response = self.client.get("/api/v1/reports/producao/export.xlsx")
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(
            response.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", response.headers["content-disposition"])

        workbook = load_workbook(BytesIO(response.content), data_only=False)
        try:
            # Relatório MES em camadas: painel executivo, análise e, por último,
            # a aba técnica oculta que preserva a rastreabilidade.
            self.assertEqual(workbook.sheetnames[0], "Visão Geral")
            self.assertEqual(workbook.sheetnames[-1], "Dados Técnicos")
            self.assertEqual(workbook["Dados Técnicos"].sheet_state, "hidden")
            summary = workbook["Visão Geral"]
            self.assertEqual(summary["A1"].value, "GESTOR DE PEÇAS")
            self.assertEqual(summary["A2"].value, "Relatório de Produção")
            self.assertIn("Período:", summary["A3"].value)
            self.assertFalse(summary.sheet_view.showGridLines)
            self.assertEqual(summary.page_setup.orientation, "landscape")
            # Sem registros no período a aba mostra o estado explícito em vez de
            # uma tabela vazia ou de zeros inventados.
            self.assertEqual(
                workbook["Setores"]["A7"].value,
                "Nenhum setor com produção registrada no período filtrado.",
            )
            self.assertEqual(summary["A10"].value, "—")
            # Nas tabelas o cabeçalho fica congelado e filtrável.
            tecnica = workbook["Dados Técnicos"]
            self.assertEqual(tecnica.freeze_panes, "A8")
            self.assertTrue(tecnica.auto_filter.ref)
        finally:
            workbook.close()

    def test_relatorio_gerado_tem_download_autorizado_e_idor_negado(self):
        self.db.criar_usuario("Outro Gestor", "senha-outro", "gestor")
        with tempfile.TemporaryDirectory() as artifact_dir:
            settings = replace(_settings(), report_artifact_dir=artifact_dir)
            app = create_app(settings=settings, database_factory=lambda: self.db)
            with TestClient(app) as client:
                login = client.post(
                    "/api/v1/auth/login",
                    json={"username": "Gestor Web", "password": "senha-segura"},
                )
                csrf = client.cookies.get("gestor_csrf")
                denied_csrf = client.post(
                    "/api/v1/reports/generate",
                    json={"report_type": "completo", "period_kind": "hoje"},
                )
                self.assertEqual(denied_csrf.status_code, 403)
                generated = client.post(
                    "/api/v1/reports/generate",
                    headers={"X-CSRF-Token": csrf},
                    json={"report_type": "completo", "period_kind": "hoje"},
                )
                self.assertEqual(generated.status_code, 201, generated.text)
                artifact = generated.json()
                self.assertNotIn("storage_path", artifact)
                download = client.get(artifact["download_url"])
                self.assertEqual(download.status_code, 200, download.text)
                self.assertEqual(
                    download.headers["content-type"],
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                )
                self.assertTrue(download.content.startswith(b"PK"))

                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Outro Gestor", "password": "senha-outro"},
                )
                forbidden = client.get(artifact["download_url"])
                self.assertEqual(forbidden.status_code, 404)

    def test_agendamento_exige_horario_configurado_e_eh_do_proprietario(self):
        with tempfile.TemporaryDirectory() as artifact_dir:
            settings = replace(_settings(), report_artifact_dir=artifact_dir)
            app = create_app(settings=settings, database_factory=lambda: self.db)
            with TestClient(app) as client:
                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Gestor Web", "password": "senha-segura"},
                )
                csrf = client.cookies.get("gestor_csrf")
                created = client.post(
                    "/api/v1/reports/schedules",
                    headers={"X-CSRF-Token": csrf},
                    json={
                        "name": "Fechamento diário",
                        "report_type": "completo",
                        "frequency": "diario",
                        "run_time": "06:30:00",
                        "timezone": "America/Sao_Paulo",
                    },
                )
                self.assertEqual(created.status_code, 201, created.text)
                listed = client.get("/api/v1/reports/schedules")
                self.assertEqual(len(listed.json()["items"]), 1)
                schedule_id = created.json()["id"]
                removed = client.delete(
                    f"/api/v1/reports/schedules/{schedule_id}",
                    headers={"X-CSRF-Token": csrf},
                )
                self.assertEqual(removed.status_code, 204)

    def test_envio_telegram_do_artifact_exige_csrf_e_destino_do_gestor(self):
        calls = []

        class ProviderStub:
            def __init__(self, **_kwargs):
                pass

            async def send_document(self, **values):
                calls.append(values)
                return {"message_id": 11, "provider": "telegram"}

        manager_id = self.db.users[0]["id"]
        self.db.messaging_destinations.append({
            "id": 1,
            "user_id": manager_id,
            "provider": "telegram",
            "label": "Gestão",
            "destination_ref": "-1000000000000",
            "enabled": True,
            "created_at": datetime(2026, 8, 26, 8, 0),
        })
        with tempfile.TemporaryDirectory() as artifact_dir, patch(
            "backend.api.dependencies.messaging.TelegramProvider", ProviderStub
        ):
            settings = replace(
                _settings(),
                report_artifact_dir=artifact_dir,
                telegram_enabled=True,
                telegram_bot_token="token-apenas-do-teste",
            )
            app = create_app(settings=settings, database_factory=lambda: self.db)
            with TestClient(app) as client:
                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Gestor Web", "password": "senha-segura"},
                )
                csrf = client.cookies.get("gestor_csrf")
                generated = client.post(
                    "/api/v1/reports/generate",
                    headers={"X-CSRF-Token": csrf},
                    json={"report_type": "ops", "period_kind": "hoje"},
                )
                self.assertEqual(generated.status_code, 201, generated.text)
                artifact = generated.json()
                self.assertTrue(artifact["telegram_available"])
                self.assertEqual(artifact["telegram_destination_id"], 1)
                denied = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    json={"destination_id": 1},
                )
                self.assertEqual(denied.status_code, 403)
                sent = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers={"X-CSRF-Token": csrf},
                    json={"destination_id": 1},
                )
                self.assertEqual(sent.status_code, 200, sent.text)
                self.assertEqual(sent.json()["status"], "enviado")
                self.assertEqual(len(calls), 1)
                self.assertTrue(calls[0]["path"].is_file())
                self.assertNotIn("destination_ref", sent.json())

    def test_envio_telegram_http_homologa_permissao_posse_falha_e_idempotencia(self):
        """Ponta a ponta HTTP do envio: API -> autenticacao -> autorizacao ->
        artifact -> destino -> ReportMessagingService -> provider fake -> delivery."""

        chamadas = []
        falhar = {"ativo": False}

        class ProviderFake:
            def __init__(self, **_kwargs):
                pass

            async def send_document(self, **values):
                chamadas.append(values)
                if falhar["ativo"]:
                    raise RuntimeError("provider indisponivel no teste")
                return {"message_id": 99, "provider": "telegram"}

        gestor_id = self.db.users[0]["id"]
        outro_id = self.db.criar_usuario("Outro Gestor", "senha-outro", "gestor")
        self.db.messaging_destinations.append({
            "id": 1,
            "user_id": gestor_id,
            "provider": "telegram",
            "label": "Gestão",
            "destination_ref": "-1000000000000",
            "enabled": True,
            "created_at": datetime(2026, 8, 26, 8, 0),
        })
        self.db.messaging_destinations.append({
            "id": 2,
            "user_id": outro_id,
            "provider": "telegram",
            "label": "Outro destino",
            "destination_ref": "-2000000000000",
            "enabled": True,
            "created_at": datetime(2026, 8, 26, 8, 0),
        })
        inexistente = "00000000-0000-0000-0000-0000000000ff"

        with tempfile.TemporaryDirectory() as artifact_dir, patch(
            "backend.api.dependencies.messaging.TelegramProvider", ProviderFake
        ):
            base = replace(_settings(), report_artifact_dir=artifact_dir)
            ligado = replace(
                base, telegram_enabled=True, telegram_bot_token="token-apenas-do-teste"
            )

            # Telegram desligado: nem o artifact nem o provider sao tocados.
            app_off = create_app(settings=base, database_factory=lambda: self.db)
            with TestClient(app_off) as client:
                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Gestor Web", "password": "senha-segura"},
                )
                csrf = client.cookies.get("gestor_csrf")
                gerado = client.post(
                    "/api/v1/reports/generate",
                    headers={"X-CSRF-Token": csrf},
                    json={"report_type": "ops", "period_kind": "hoje"},
                )
                self.assertEqual(gerado.status_code, 201, gerado.text)
                artifact = gerado.json()
                self.assertNotIn("telegram_available", artifact)
                self.assertNotIn("telegram_destination_id", artifact)
                desligado = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers={"X-CSRF-Token": csrf},
                    json={"destination_id": 1},
                )
                self.assertEqual(desligado.status_code, 409)
                self.assertEqual(desligado.json()["code"], "messaging_disabled")
            self.assertEqual(chamadas, [])
            self.assertEqual(self.db.report_deliveries, [])

            app = create_app(settings=ligado, database_factory=lambda: self.db)

            # Operador nao alcanca o endpoint gerencial de envio.
            with TestClient(app) as operador:
                operador.post(
                    "/api/v1/auth/login",
                    json={"username": "Operador Web", "password": "senha-operador"},
                )
                negado = operador.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers={"X-CSRF-Token": operador.cookies.get("gestor_csrf")},
                    json={"destination_id": 1},
                )
                self.assertEqual(negado.status_code, 403)

            # Gestor diferente nao alcanca o artifact alheio (IDOR).
            with TestClient(app) as outro:
                outro.post(
                    "/api/v1/auth/login",
                    json={"username": "Outro Gestor", "password": "senha-outro"},
                )
                alheio = outro.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers={"X-CSRF-Token": outro.cookies.get("gestor_csrf")},
                    json={"destination_id": 2},
                )
                self.assertEqual(alheio.status_code, 404)
                self.assertEqual(alheio.json()["code"], "report_not_found")

            self.assertEqual(chamadas, [])

            with TestClient(app) as client:
                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Gestor Web", "password": "senha-segura"},
                )
                csrf = client.cookies.get("gestor_csrf")
                cabecalho = {"X-CSRF-Token": csrf}

                # Artifact inexistente.
                ausente = client.post(
                    f"/api/v1/reports/artifacts/{inexistente}/send",
                    headers=cabecalho,
                    json={"destination_id": 1},
                )
                self.assertEqual(ausente.status_code, 404)
                self.assertEqual(ausente.json()["code"], "report_not_found")

                # Destino pertencente a outro gestor.
                destino_alheio = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers=cabecalho,
                    json={"destination_id": 2},
                )
                self.assertEqual(destino_alheio.status_code, 404)
                self.assertEqual(
                    destino_alheio.json()["code"], "messaging_destination_not_found"
                )
                self.assertEqual(chamadas, [])

                # Provider falhando: erro controlado, delivery registrado como
                # falhou e o relatorio continua disponivel para download.
                falhar["ativo"] = True
                falha = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers=cabecalho,
                    json={"destination_id": 1},
                )
                self.assertEqual(falha.status_code, 502)
                self.assertEqual(falha.json()["code"], "messaging_provider_error")
                self.assertNotIn("-1000000000000", falha.text)
                self.assertEqual(len(chamadas), 1)
                registrada = self.db.report_deliveries[-1]
                self.assertEqual(registrada["status"], "falhou")
                self.assertEqual(registrada["attempt"], 1)
                self.assertEqual(registrada["error_code"], "messaging_provider_error")
                preservado = client.get(artifact["download_url"])
                self.assertEqual(preservado.status_code, 200)
                self.assertTrue(preservado.content.startswith(b"PK"))

                # Retry bem-sucedido reaproveita a mesma entrega.
                falhar["ativo"] = False
                enviado = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers=cabecalho,
                    json={"destination_id": 1},
                )
                self.assertEqual(enviado.status_code, 200, enviado.text)
                corpo = enviado.json()
                self.assertEqual(corpo["status"], "enviado")
                self.assertEqual(corpo["attempt"], 2)
                self.assertEqual(corpo["destination_id"], 1)
                self.assertEqual(corpo["report_id"], artifact["id"])
                self.assertNotIn("destination_ref", corpo)
                self.assertEqual(len(chamadas), 2)
                self.assertEqual(len(self.db.report_deliveries), 1)

                # Reenvio nao duplica arquivo, entrega nem chamada ao provider.
                repetido = client.post(
                    f"/api/v1/reports/artifacts/{artifact['id']}/send",
                    headers=cabecalho,
                    json={"destination_id": 1},
                )
                self.assertEqual(repetido.status_code, 200, repetido.text)
                self.assertEqual(repetido.json()["id"], corpo["id"])
                self.assertEqual(repetido.json()["attempt"], 2)
                self.assertEqual(len(chamadas), 2)
                self.assertEqual(len(self.db.report_deliveries), 1)
                self.assertEqual(len(self.db.generated_reports), 1)

    def test_openapi_contem_rotas_por_dominio(self):
        response = self.client.get("/api/openapi.json")
        self.assertEqual(response.status_code, 200)
        paths = response.json()["paths"]
        for path in (
            "/api/v1/auth/login",
            "/api/v1/andon",
            "/api/v1/welding",
            "/api/v1/management/overview",
            "/api/v1/management/insights",
            "/api/v1/management/kpis/{kpi_key}/explanation",
            "/api/v1/operator/context",
            "/api/v1/operator/actions",
            "/api/v1/cutting/queue",
            "/api/v1/highlight/tasks/{task_code}",
            "/api/v1/operations/resources",
            "/api/v1/orders",
            "/api/v1/analytics/{analysis_type}",
            "/api/v1/audit",
            "/api/v1/traceability/orders/{op}",
        ):
            self.assertIn(path, paths)

    def test_solda_gerencial_e_leitura_compartilhada_com_o_perfil_de_tv(self):
        """Wave 6D — a visão gerencial da Solda usa a autorização do Andon.

        A mesma projeção serve a tela de gestão e o ciclo da TV, sem segundo
        backend e sem autenticação própria. O posto não acessa: a tela é de
        acompanhamento, não de execução.
        """

        for suffix in ("A", "B"):
            self.db.pcp_ops.append({
                "codigo_op": f"OP-SOLDA-{suffix}",
                "produto_codigo": f"PECA-SOLDA-{suffix}",
                "produto_descricao": "Peça de solda Web",
                "produto_modelo": None,
                "quantidade": 1,
                "unidade": "UN",
                "ativo": True,
                "fim_planejado": datetime(2026, 12, 31, 17, 30),
            })

        self.login_andon()
        resposta = self.client.get("/api/v1/welding")
        self.assertEqual(resposta.status_code, 200, resposta.text)
        leitura = resposta.json()
        self.assertEqual(leitura["setor"], "Solda")
        linhas = [linha for grupo in leitura["estacoes"] for linha in grupo["ops"]]
        self.assertEqual({linha["op"] for linha in linhas}, {"OP-SOLDA-A", "OP-SOLDA-B"})
        # Sem execução observada, a estação continua ausente e declarada.
        self.assertTrue(all(linha["estacao"]["value"] is None for linha in linhas))
        self.assertTrue(all(linha["modelo"]["value"] is None for linha in linhas))

        self.client.cookies.clear()
        self.login_operator()
        self.assertEqual(self.client.get("/api/v1/welding").status_code, 403)

    def test_andon_entrega_snapshot_geral_agrupado_e_somente_leitura(self):
        now = datetime(2026, 8, 24, 14, 30)
        self.db.listar_recursos_pcfactory = lambda *_args, **_kwargs: [
            {"codigo": "DOBRA3", "nome": "Dobradeira 1303", "tipo_setor": "Dobra", "habilitado": True},
            {"codigo": "DOBRA4", "nome": "Dobradeira 2204", "tipo_setor": "Dobra", "habilitado": True},
            {"codigo": "SERRA1", "nome": "Serra 01", "tipo_setor": "Serra", "habilitado": True},
        ]
        self.db.listar_estados_recurso_atuais = lambda **_kwargs: [
            {"id": 1, "recurso": "DOBRA3", "tipo_setor": "Dobra", "categoria": "producao", "data_inicio": now, "op": "OP-WEB", "numero_operacao": "20", "produto_codigo": "PECA-WEB"},
            {"id": 2, "recurso": "DOBRA4", "tipo_setor": "Dobra", "categoria": "parada", "data_inicio": now, "motivo": "Falta de material"},
            {"id": 3, "recurso": "SERRA1", "tipo_setor": "Serra", "categoria": "fila", "data_inicio": now},
        ]
        self.db.listar_fatos_operacionais_periodo = lambda *_args, **_kwargs: [{
            "id": 11,
            "op": "OP-WEB",
            "tipo_setor": "Dobra",
            "maquina": "DOBRA3",
            "status": "Em processo",
            "data_inicio": now,
            "data_fim": None,
            "operador_inicio": "Operador Web",
            "quantidade": 12,
            "quantidade_planejada_pcp": 12,
            "quantidade_boa": 5,
            "quantidade_refugo": 0,
            "quantidade_retrabalho": 0,
            "produto_codigo": "PECA-WEB",
            "produto_descricao": "Peça do Andon",
            "numero_operacao": "20",
            "descricao_operacao": "DOBRA",
            "eventos": [],
        }]

        self.login_manager()
        response = self.client.get(
            "/api/v1/andon",
            params={"inicio": "2026-08-24T00:00:00", "fim": "2026-08-24T23:59:59"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        payload = response.json()
        self.assertEqual(payload["resource_count"], 2)
        self.assertEqual(
            [item["name"] for item in payload["sectors"]],
            ["Corte", "Caldeiraria", "Solda", "Pintura"],
        )
        resources = {
            item["code"]: item
            for sector in payload["sectors"]
            for item in sector["resources"]
        }
        self.assertEqual(resources["DOBRA3"]["state"]["category"], "producao")
        self.assertEqual(resources["DOBRA3"]["operation"]["op"], "OP-WEB")
        self.assertEqual(resources["DOBRA3"]["operation"]["good_quantity"], 5)
        self.assertEqual(resources["DOBRA4"]["state"]["reason"], "Falta de material")
        self.assertNotIn("SERRA1", resources)
        self.assertNotIn("actions", payload)

    def test_operador_nao_acessa_snapshot_geral_do_andon(self):
        self.login_operator()
        response = self.client.get("/api/v1/andon")
        self.assertEqual(response.status_code, 403)
        self.assertEqual(response.json()["code"], "andon_access_denied")

    def test_perfil_andon_acessa_somente_o_andon(self):
        login = self.login_andon()
        self.assertEqual(login.json()["role"], "andon")
        self.assertTrue(login.json()["andon_access"])
        self.assertFalse(login.json()["management_access"])
        self.assertFalse(login.json()["operator_access"])
        self.assertEqual(self.client.get("/api/v1/andon").status_code, 200)
        self.assertEqual(self.client.get("/api/v1/management/overview").status_code, 403)
        self.assertEqual(self.client.get("/api/v1/operator/context").status_code, 403)

    def test_todas_as_telas_gerenciais_possuem_endpoint_funcional(self):
        self.login_manager()
        paths = (
            "/api/v1/management/overview",
            "/api/v1/andon",
            "/api/v1/management/sectors",
            "/api/v1/management/alerts",
            "/api/v1/operations/overview",
            "/api/v1/operations/resources",
            "/api/v1/operations/orders",
            "/api/v1/operations/time",
            "/api/v1/orders",
            "/api/v1/orders/production",
            "/api/v1/orders/planned-vs-actual",
            "/api/v1/analytics/oee",
            "/api/v1/analytics/hours-utilization",
            "/api/v1/analytics/downtimes",
            "/api/v1/analytics/setups",
            "/api/v1/analytics/quality",
            "/api/v1/analytics/standard-vs-actual",
            "/api/v1/analytics/chronoanalysis",
            "/api/v1/analytics/capacity",
            "/api/v1/audit/appointments",
            "/api/v1/audit",
            "/api/v1/audit/reliability",
            "/api/v1/reports/gerencial",
            "/api/v1/reports/producao",
            "/api/v1/reports/perdas",
            "/api/v1/reports/indicadores",
            "/api/v1/reports/dados_analiticos",
            "/api/v1/traceability/orders/OP-TESTE",
            "/api/v1/traceability/nestings",
        )
        for path in paths:
            with self.subTest(path=path):
                response = self.client.get(path)
                self.assertEqual(response.status_code, 200, response.text)

    def test_alias_gestor_peca_e_apenas_configuracao_publica(self):
        self.assertEqual(self.app.state.settings.public_host, "gestor-peca")

    def test_build_web_pode_ser_entregue_pelo_fastapi_com_fallback_spa(self):
        response = self.client.get("/rastreabilidade/linha-do-tempo")
        if response.status_code == 404:
            self.skipTest("Build React ainda não foi gerado neste checkout.")
        self.assertEqual(response.status_code, 200)
        self.assertIn("Gestor de Peças", response.text)
        self.assertIn("frame-ancestors 'none'", response.headers["content-security-policy"])

    def test_respostas_grandes_sao_comprimidas_sem_alterar_o_contrato(self):
        self.login_manager()
        response = self.client.get(
            "/api/v1/management/overview",
            headers={"Accept-Encoding": "gzip"},
        )
        self.assertEqual(response.status_code, 200, response.text)
        self.assertEqual(response.headers.get("content-encoding"), "gzip")
        self.assertEqual(
            response.json()["periodo"]["inicio"],
            datetime.combine(date.today(), time.min).isoformat(),
        )


if __name__ == "__main__":
    unittest.main()


class QualityWebApiTests(unittest.TestCase):
    """Contrato HTTP da Qualidade: visibilidade, permissão e ausência de pull."""

    def setUp(self):
        self.db = ApiFakeDatabase()
        self.db.criar_usuario("Operador Dobra", "senha-dobra", "operador_dobra")
        self.db.criar_usuario("Operador Corte", "senha-corte", "operador_corte")
        self.db.criar_usuario("Supervisor Web", "senha-supervisor", "supervisor")
        self.db.cadastrar_operador_apontamento("77", "Inspetor")
        task_id = self.db.inserir_tarefa("T-QUAL-API")
        self.db.inserir_op_na_tarefa(task_id, "OP-QUAL-API", "PROD-API", "Dobra", 2)
        self.db.catalog_operations.append({
            "id": 701,
            "codigo_op": "OP-QUAL-API",
            "numero_operacao": "10",
            "codigo_recurso": "DOBRA3",
            "descricao_operacao": "DOBRA",
            "tipo_setor": "Dobra",
            "produto_codigo": "PROD-API",
            "produto_descricao": "Peça inspecionada",
            "quantidade": 2,
            "ordem": 1,
            "ativo": True,
        })
        self.db.catalog_operations.append({
            "id": 702,
            "codigo_op": "OP-QUAL-API",
            "numero_operacao": "20",
            "codigo_recurso": "INSPEC",
            "descricao_operacao": "INSPECAO",
            "tipo_setor": None,
            "produto_codigo": "PROD-API",
            "produto_descricao": "Peça inspecionada",
            "quantidade": 2,
            "ordem": 2,
            "ativo": False,
            "marco_terminal": False,
            "inspecao_qualidade": True,
        })
        self.app = create_app(settings=_settings(), database_factory=lambda: self.db)
        self.client = TestClient(self.app)

    def tearDown(self):
        self.client.close()

    def _login(self, username, password):
        response = self.client.post(
            "/api/v1/auth/login", json={"username": username, "password": password}
        )
        self.assertEqual(response.status_code, 200, response.text)

    def _csrf(self):
        return {"X-CSRF-Token": self.client.cookies.get("gestor_csrf")}

    def _liberar_inspecao(self):
        flow = OperatorFlowService(self.db, "OPERADOR DOBRA")
        operacao = self.db.listar_operacoes_para_op("OP-QUAL-API", "Dobra")[0]
        liberar_primeira_peca(
            self.db,
            "OPERADOR DOBRA",
            op="OP-QUAL-API",
            setor="Dobra",
            recurso="1303",
            operacao=operacao,
        )
        flow.executar(
            "Início", op="OP-QUAL-API", setor="Dobra", recurso="1303", operacao=operacao
        )
        flow.executar(
            "Finalizado",
            op="OP-QUAL-API",
            setor="Dobra",
            recurso="1303",
            operacao=operacao,
            pecas_boas=2,
            operadores_cracha=["77"],
        )

    def test_aba_qualidade_aparece_para_caldeiraria_e_nao_para_corte(self):
        self._login("Operador Dobra", "senha-dobra")
        contexto = self.client.get("/api/v1/quality/context")
        self.assertEqual(contexto.status_code, 200, contexto.text)
        self.assertTrue(contexto.json()["available"])
        self.assertEqual(contexto.json()["sector"], "Dobra")
        self.assertFalse(contexto.json()["can_edit_template"])
        self.client.post("/api/v1/auth/logout", headers=self._csrf())

        self._login("Operador Corte", "senha-corte")
        contexto = self.client.get("/api/v1/quality/context")
        self.assertEqual(contexto.status_code, 200, contexto.text)
        self.assertFalse(contexto.json()["available"])
        # Sem a capacidade, nenhum caso de uso da Qualidade é acessível.
        for path in ("/api/v1/quality/queue", "/api/v1/quality/history"):
            with self.subTest(path=path):
                negado = self.client.get(path)
                self.assertEqual(negado.status_code, 403, negado.text)
                self.assertEqual(negado.json()["code"], "quality_sector_unavailable")

    def test_fila_e_pesquisa_sao_locais_e_nao_acionam_o_erp(self):
        chamadas = []
        self.app.state.order_provisioning_factory = lambda database: chamadas.append(
            database
        )
        self._login("Operador Dobra", "senha-dobra")

        vazia = self.client.get("/api/v1/quality/queue")
        self.assertEqual(vazia.status_code, 200, vazia.text)
        self.assertEqual(vazia.json()["items"], [])

        self._liberar_inspecao()
        fila = self.client.get("/api/v1/quality/queue")
        self.assertEqual(
            [item["op"] for item in fila.json()["items"]], ["OP-QUAL-API"]
        )

        ausente = self.client.get("/api/v1/quality/queue?search=OP-QUE-NAO-EXISTE")
        self.assertEqual(ausente.json()["items"], [])
        self.assertEqual(
            ausente.json()["busca"]["message"], "OP não disponível para inspeção."
        )
        # A tela da Qualidade jamais provisiona OP: o ERP não é tocado.
        self.assertEqual(chamadas, [])

    def test_operador_cadastra_o_primeiro_template_e_nao_altera_depois(self):
        self._liberar_inspecao()
        self._login("Operador Dobra", "senha-dobra")
        abertura = self.client.post(
            "/api/v1/quality/inspections",
            headers=self._csrf(),
            json={"op": "OP-QUAL-API"},
        )
        self.assertEqual(abertura.status_code, 200, abertura.text)
        self.assertIsNone(abertura.json()["data"]["template"])

        criado = self.client.post(
            "/api/v1/quality/templates",
            headers=self._csrf(),
            json={
                "produto": "PROD-API",
                "cotas": [{"sequencia": 1, "padrao": "125,0 ± 0,5", "unidade": "mm"}],
            },
        )
        self.assertEqual(criado.status_code, 200, criado.text)

        bloqueado = self.client.post(
            "/api/v1/quality/templates",
            headers=self._csrf(),
            json={
                "produto": "PROD-API",
                "cotas": [{"sequencia": 1, "padrao": "999,0 ± 9,9"}],
            },
        )
        self.assertEqual(bloqueado.status_code, 409, bloqueado.text)
        self.assertEqual(bloqueado.json()["code"], "qualidade_template_protegido")
        self.client.post("/api/v1/auth/logout", headers=self._csrf())

        # Supervisor altera o mesmo template pelo mesmo contrato.
        self._login("Supervisor Web", "senha-supervisor")
        alterado = self.client.post(
            "/api/v1/quality/templates",
            headers=self._csrf(),
            json={
                "produto": "PROD-API",
                "cotas": [{"sequencia": 1, "padrao": "130,0 ± 0,5", "unidade": "mm"}],
            },
        )
        self.assertEqual(alterado.status_code, 200, alterado.text)
        self.assertEqual(alterado.json()["data"]["revisao"], 2)

    def test_peca_exige_csrf_e_respeita_a_sequencia(self):
        self._liberar_inspecao()
        self._login("Operador Dobra", "senha-dobra")
        inspecao = self.client.post(
            "/api/v1/quality/inspections",
            headers=self._csrf(),
            json={"op": "OP-QUAL-API"},
        ).json()["data"]
        self.client.post(
            "/api/v1/quality/templates",
            headers=self._csrf(),
            json={
                "produto": "PROD-API",
                "cotas": [{"sequencia": 1, "padrao": "125,0 ± 0,5", "unidade": "mm"}],
            },
        )
        corpo = {
            "numero_peca": 1,
            "resultado": "APROVADA",
            "medidas": [{"sequencia": 1, "medida": "124,9", "status": "CONFORME"}],
            "badges": [],
        }
        sem_csrf = self.client.post(
            f"/api/v1/quality/inspections/{inspecao['id']}/pieces", json=corpo
        )
        self.assertEqual(sem_csrf.status_code, 403)

        salva = self.client.post(
            f"/api/v1/quality/inspections/{inspecao['id']}/pieces",
            headers=self._csrf(),
            json=corpo,
        )
        self.assertEqual(salva.status_code, 200, salva.text)
        self.assertEqual(salva.json()["data"]["peca_atual"], 2)

        repetida = self.client.post(
            f"/api/v1/quality/inspections/{inspecao['id']}/pieces",
            headers=self._csrf(),
            json=corpo,
        )
        self.assertEqual(repetida.status_code, 409, repetida.text)
        self.assertEqual(repetida.json()["code"], "qualidade_sequencia_peca")

    def test_desenho_ausente_responde_404_e_upload_e_restrito(self):
        self._login("Operador Dobra", "senha-dobra")
        ausente = self.client.get("/api/v1/quality/drawings/PROD-API/file")
        self.assertEqual(ausente.status_code, 404, ausente.text)
        self.assertEqual(ausente.json()["code"], "quality_drawing_not_found")

        recusado = self.client.post(
            "/api/v1/quality/drawings",
            headers=self._csrf(),
            json={
                "produto": "PROD-API",
                "filename": "desenho.pdf",
                "conteudo_base64": base64.b64encode(b"%PDF-1.4 teste").decode(),
            },
        )
        self.assertEqual(recusado.status_code, 403, recusado.text)
        self.assertEqual(recusado.json()["code"], "quality_drawing_forbidden")

    def test_supervisor_publica_desenho_e_operador_visualiza_na_propria_tela(self):
        with tempfile.TemporaryDirectory() as directory:
            settings = replace(_settings(), quality_drawing_dir=directory)
            app = create_app(settings=settings, database_factory=lambda: self.db)
            with TestClient(app) as client:
                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Supervisor Web", "password": "senha-supervisor"},
                )
                csrf = {"X-CSRF-Token": client.cookies.get("gestor_csrf")}
                enviado = client.post(
                    "/api/v1/quality/drawings",
                    headers=csrf,
                    json={
                        "produto": "PROD-API",
                        "filename": "chapa.pdf",
                        "conteudo_base64": base64.b64encode(
                            b"%PDF-1.4 conteudo de teste"
                        ).decode(),
                    },
                )
                self.assertEqual(enviado.status_code, 200, enviado.text)
                self.assertEqual(enviado.json()["desenho"]["versao"], 1)

                invalido = client.post(
                    "/api/v1/quality/drawings",
                    headers=csrf,
                    json={
                        "produto": "PROD-API",
                        "filename": "nao-e-pdf.pdf",
                        "conteudo_base64": base64.b64encode(b"apenas texto").decode(),
                    },
                )
                self.assertEqual(invalido.status_code, 400, invalido.text)
                client.post("/api/v1/auth/logout", headers=csrf)

                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Operador Dobra", "password": "senha-dobra"},
                )
                arquivo = client.get("/api/v1/quality/drawings/PROD-API/file")
                self.assertEqual(arquivo.status_code, 200, arquivo.text)
                self.assertEqual(arquivo.headers["content-type"], "application/pdf")
                self.assertIn("inline", arquivo.headers["content-disposition"])
                self.assertTrue(arquivo.content.startswith(b"%PDF-"))

    def test_desenho_do_operador_com_nome_acentuado_nao_quebra_o_header(self):
        """Nome fora do latin-1 derrubava o header (500); agora vira filename* (RFC 5987)."""

        self.db.buscar_op_planejada = lambda codigo: (
            {"produto_codigo": "PROD-API"} if codigo == "OP-QUAL-API" else None
        )
        with tempfile.TemporaryDirectory() as directory:
            nome = "PROD-API_revisão_Ω.pdf"
            (Path(directory) / nome).write_bytes(b"%PDF-1.4 conteudo de teste")
            settings = replace(_settings(), operator_drawing_roots=directory)
            app = create_app(settings=settings, database_factory=lambda: self.db)
            with TestClient(app) as client:
                client.post(
                    "/api/v1/auth/login",
                    json={"username": "Operador Dobra", "password": "senha-dobra"},
                )
                arquivo = client.get(
                    "/api/v1/operator/drawings/file", params={"op": "OP-QUAL-API"}
                )
                self.assertEqual(arquivo.status_code, 200, arquivo.text)
                disposicao = arquivo.headers["content-disposition"]
                self.assertTrue(disposicao.startswith("inline;"), disposicao)
                self.assertIn("filename*=utf-8''", disposicao)
                self.assertEqual(arquivo.headers["cache-control"], "private, max-age=60")
                self.assertTrue(arquivo.content.startswith(b"%PDF-"))


class LoginThrottleTests(unittest.TestCase):
    """Atraso progressivo do login principal (achado A2 da auditoria de

    segurança de 2026-09-14): nunca bloqueia de vez — um operador de chão de
    fábrica não pode ficar impedido de apontar produção por errar a senha —
    mas cada falha seguida no mesmo IP+usuário atrasa a próxima tentativa.
    """

    def setUp(self):
        from backend.api.routers import auth as auth_router

        self.auth_router = auth_router
        self.key = "10.0.0.1:operador teste"
        self.auth_router._clear_login_failures(self.key)

    def tearDown(self):
        self.auth_router._clear_login_failures(self.key)

    def test_sem_falha_recente_nao_ha_atraso(self):
        self.assertEqual(self.auth_router._login_delay_seconds(self.key, now=1_000.0), 0.0)

    def test_atraso_cresce_e_satura_no_teto(self):
        router = self.auth_router
        now = 1_000.0
        router._register_login_failure(self.key, now=now)
        self.assertEqual(router._login_delay_seconds(self.key, now=now), router.LOGIN_DELAY_BASE_SECONDS)
        for _ in range(8):
            router._register_login_failure(self.key, now=now)
        self.assertEqual(router._login_delay_seconds(self.key, now=now), router.LOGIN_DELAY_MAX_SECONDS)

    def test_sucesso_zera_o_atraso(self):
        router = self.auth_router
        now = 1_000.0
        router._register_login_failure(self.key, now=now)
        self.assertGreater(router._login_delay_seconds(self.key, now=now), 0.0)
        router._clear_login_failures(self.key)
        self.assertEqual(router._login_delay_seconds(self.key, now=now), 0.0)

    def test_janela_vencida_reseta_o_contador(self):
        router = self.auth_router
        router._register_login_failure(self.key, now=1_000.0)
        depois = 1_000.0 + router.LOGIN_DELAY_WINDOW_SECONDS + 1
        self.assertEqual(router._login_delay_seconds(self.key, now=depois), 0.0)
