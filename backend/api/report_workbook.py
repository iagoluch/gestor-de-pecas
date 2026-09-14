"""Exportação XLSX dos relatórios Web sem recalcular a verdade do backend."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
import json
import re
from typing import Any, Iterable

from fastapi.encoders import jsonable_encoder
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from mes.contracts import AnalyticsFilter


NAVY = "0B2E63"
PRIMARY = "0866F5"
GREEN = "16A34A"
RED = "EF3340"
YELLOW = "F5B400"
PURPLE = "7047E8"
TEAL = "20A7A1"
TEXT = "081630"
MUTED = "71809D"
LIGHT = "F5F7FA"
BORDER = "C7D3E2"
WHITE = "FFFFFF"

REPORT_TITLES = {
    "completo": "Relatório Industrial Completo",
    "gerencial": "Relatório Gerencial",
    "producao": "Relatório de Produção",
    "perdas": "Relatório de Perdas",
    "indicadores": "Relatório de Indicadores",
    "dados_analiticos": "Dados Analíticos",
    "ops": "Ordens de Produção",
    "paradas": "Paradas",
    "setup": "Setup",
    "qualidade": "Qualidade",
    "recursos": "Recursos",
    "setores": "Setores",
    "nestings": "Nestings",
    "excecoes": "Exceções",
    "rastreabilidade": "Rastreabilidade",
    "auditoria": "Auditoria",
}


@dataclass(frozen=True)
class MetricSpec:
    label: str
    path: str
    color: str
    kind: str = "number"


REPORT_METRICS = {
    "gerencial": (
        MetricSpec("Produção boa", "overview.production.good", GREEN),
        MetricSpec("Refugo", "overview.production.scrap", RED),
        MetricSpec("Retrabalho", "overview.production.rework", YELLOW),
        MetricSpec("OEE", "overview.kpis.oee.value", PURPLE, "percent"),
    ),
    "producao": (
        MetricSpec("Produção boa", "production.good", GREEN),
        MetricSpec("Refugo", "production.scrap", RED),
        MetricSpec("Retrabalho", "production.rework", YELLOW),
        MetricSpec("OPs", "production.ops", PRIMARY),
    ),
    "perdas": (
        MetricSpec("Tempo de parada", "paradas.total_seconds", RED, "duration"),
        MetricSpec("Ocorrências", "paradas.count", YELLOW),
        MetricSpec("Refugo", "qualidade.totals.refugo", RED),
        MetricSpec("Retrabalho", "qualidade.totals.retrabalho", YELLOW),
    ),
    "indicadores": (
        MetricSpec("OEE", "analytics.oee.value", PURPLE, "percent"),
        MetricSpec("FTT", "analytics.qualidade.ftt.value", TEAL, "percent"),
        MetricSpec("Tempo físico", "analytics.tempos.physical_seconds", GREEN, "duration"),
        MetricSpec("Utilizacao", "analytics.capacidade.utilizacao_percentual", YELLOW, "percent"),
        MetricSpec("MTBF", "analytics.confiabilidade.mtbf_segundos", PRIMARY, "duration"),
        MetricSpec("MTTR", "analytics.confiabilidade.mttr_segundos", RED, "duration"),
    ),
    "dados_analiticos": (
        MetricSpec("Ordens", "orders.page.total", PRIMARY),
        MetricSpec("Inconsistências", "audit.count", RED),
        MetricSpec("Nestings", "nestings.count", PURPLE),
        MetricSpec("Fonte", "orders.source", TEAL, "text"),
    ),
}

_INTELLIGENCE_METRICS = (
    MetricSpec("OEE", "sections.Resumo Executivo.kpis.oee.value", PURPLE, "percent"),
    MetricSpec("Disponibilidade", "sections.Resumo Executivo.kpis.availability.value", PRIMARY, "percent"),
    MetricSpec("Performance", "sections.Resumo Executivo.kpis.performance.value", GREEN, "percent"),
    MetricSpec("FTT", "sections.Resumo Executivo.kpis.ftt.value", TEAL, "percent"),
)
for _report_key in (
    "completo", "gerencial", "ops", "paradas", "setup", "qualidade",
    "recursos", "setores", "nestings", "excecoes", "rastreabilidade", "auditoria",
):
    REPORT_METRICS[_report_key] = _INTELLIGENCE_METRICS

FRIENDLY_NAMES = {
    "op": "OP",
    "oee": "OEE",
    "ftt": "FTT",
    "id": "ID",
    "inicio": "Início",
    "fim": "Fim",
    "producao": "Produção",
    "operacao": "Operação",
    "descricao": "Descrição",
    "quantidade": "Quantidade",
    "recurso": "Recurso",
    "setor": "Setor",
    "periodo": "Período",
    "production": "Produção",
    "quality": "Qualidade",
    "good": "Produção boa",
    "scrap": "Refugo",
    "rework": "Retrabalho",
    "resources": "Recursos",
    "totals": "Totais",
    "count": "Quantidade",
    "ops": "OPs",
    "refugo": "Refugo",
    "retrabalho": "Retrabalho",
    "motivo": "Motivo",
    "segundos": "Segundos",
    "availability": "Disponibilidade",
    "reason": "Justificativa",
    "source": "Fonte",
    "value": "Valor",
    "unit": "Unidade",
}


def _humanize(value: str) -> str:
    parts = str(value or "").replace("-", "_").split("_")
    return " ".join(FRIENDLY_NAMES.get(part.casefold(), part.capitalize()) for part in parts if part)


# Rótulos de coluna resolvidos pela chave inteira. Existem para traduzir chaves
# de API que apareceriam em inglês no relatório e para desfazer a ambiguidade
# entre a coluna "Valor" das linhas escalares e o valor de cada ponto de série.
COLUMN_LABELS = {
    "at": "Momento",
    "period_start": "Início do período",
    "period_end": "Fim do período",
    "value": "Valor do período",
    "components": "Componentes",
    "unit": "Unidade",
}


def _column_label(key: str) -> str:
    return COLUMN_LABELS.get(str(key), _humanize(str(key)))


def _get_path(payload: dict[str, Any], path: str):
    value: Any = payload
    for part in path.split("."):
        if not isinstance(value, dict):
            return None
        value = value.get(part)
    return value


def _safe_text(value: str) -> str:
    text = str(value)
    if text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _excel_value(value: Any):
    if value is None or isinstance(value, (bool, int, float, date, datetime)):
        return value
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)
    text = str(value).strip()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?)?", text):
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
            return parsed.replace(tzinfo=None)
        except ValueError:
            pass
    return _safe_text(text)


def _metric_value(raw: Any, kind: str):
    if raw is None:
        return "Dados insuficientes", "@"
    if kind == "percent" and isinstance(raw, (int, float)):
        return float(raw) / 100.0, "0.0%"
    if kind == "duration" and isinstance(raw, (int, float)):
        return float(raw) / 86400.0, '[h]"h "mm"m "ss"s"'
    if kind == "number" and isinstance(raw, (int, float)):
        return raw, "#,##0.00" if isinstance(raw, float) and not raw.is_integer() else "#,##0"
    return _excel_value(raw), "@"


def _flatten_scalars(value: Any, path: tuple[str, ...] = ()) -> Iterable[tuple[str, Any]]:
    if isinstance(value, dict):
        for key, child in value.items():
            if key == "type":
                continue
            yield from _flatten_scalars(child, (*path, str(key)))
        return
    if isinstance(value, list):
        return
    yield " › ".join(_humanize(part) for part in path), value


def _table_sections(value: Any, path: tuple[str, ...] = ()) -> Iterable[tuple[str, list[dict[str, Any]]]]:
    if not isinstance(value, dict):
        return
    for key, child in value.items():
        current_path = (*path, str(key))
        if isinstance(child, list) and child and all(isinstance(item, dict) for item in child):
            yield ".".join(current_path), child
        elif isinstance(child, dict):
            yield from _table_sections(child, current_path)


def _sheet_title(path: str, used: set[str]) -> str:
    aliases = {
        "overview.sectors": "Setores",
        "sectors": "Produção por setor",
        "by_sector": "Resumo por setor",
        "production.by_sector": "Resumo por setor",
        "quality.by_sector": "Resumo por setor",
        "by_product": "Produção por produto",
        "production.by_product": "Produção por produto",
        "quality.by_product": "Produção por produto",
        "scrap_reasons": "Motivos de refugo",
        "production.scrap_reasons": "Motivos de refugo",
        "quality.scrap_reasons": "Motivos de refugo",
        "rework_reasons": "Motivos de retrabalho",
        "production.rework_reasons": "Motivos de retrabalho",
        "quality.rework_reasons": "Motivos de retrabalho",
        "paradas.by_reason": "Paradas por motivo",
        "orders.items": "Ordens",
        "audit.items": "Auditoria",
        "nestings.items": "Nestings",
    }
    base = aliases.get(path, _humanize(path.split(".")[-1])) or "Dados"
    base = re.sub(r"[\\/*?:\[\]]", " ", base).strip()[:31] or "Dados"
    candidate = base
    index = 2
    while candidate.casefold() in used:
        suffix = f" {index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def _apply_print_setup(sheet, *, repeat_rows: str | None = None) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A5" if sheet.title == "Resumo" else "A5"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4
    if repeat_rows:
        sheet.print_title_rows = repeat_rows


def _style_summary(sheet, report_type: str, payload: dict[str, Any], filters: AnalyticsFilter) -> None:
    title = REPORT_TITLES.get(report_type, _humanize(report_type))
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H2")
    sheet["A1"] = f"GESTOR DE PEÇAS — {title.upper()}"
    sheet["A1"].font = Font(name="Arial", size=22, bold=True, color=WHITE)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY)
    for row in sheet["A1:H2"]:
        for cell in row:
            cell.fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells("A3:H3")
    sheet["A3"] = "Dados produzidos pelo backend canônico; o Excel não recalcula indicadores industriais."
    sheet["A3"].font = Font(name="Arial", size=10, color=WHITE)
    sheet["A3"].fill = PatternFill("solid", fgColor=NAVY)

    period = f"Período: {filters.inicio:%d/%m/%Y %H:%M} a {filters.fim:%d/%m/%Y %H:%M}"
    applied = [f"{_humanize(key)}: {value}" for key, value in filters.to_dict().items() if key not in {"inicio", "fim"} and value]
    sheet.merge_cells("A4:H4")
    sheet["A4"] = period + ("  |  " + "  |  ".join(applied) if applied else "")
    sheet["A4"].font = Font(name="Arial", size=10, color=MUTED)
    sheet["A4"].alignment = Alignment(vertical="center")
    sheet["A4"].fill = PatternFill("solid", fgColor=LIGHT)

    thin = Side(style="thin", color=BORDER)
    metrics = REPORT_METRICS.get(report_type, ())
    for index, spec in enumerate(metrics):
        start = 1 + index * 2
        end = start + 1
        title_range = f"{sheet.cell(6, start).coordinate}:{sheet.cell(6, end).coordinate}"
        value_range = f"{sheet.cell(7, start).coordinate}:{sheet.cell(9, end).coordinate}"
        sheet.merge_cells(title_range)
        sheet.merge_cells(value_range)
        title_cell = sheet.cell(6, start)
        value_cell = sheet.cell(7, start)
        title_cell.value = spec.label
        title_cell.fill = PatternFill("solid", fgColor=spec.color)
        title_cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        value_cell.value, value_cell.number_format = _metric_value(_get_path(payload, spec.path), spec.kind)
        value_cell.fill = PatternFill("solid", fgColor=WHITE)
        value_cell.font = Font(name="Arial", size=20, bold=True, color=TEXT)
        value_cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in sheet.iter_rows(min_row=6, max_row=9, min_col=start, max_col=end):
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

    sheet.merge_cells("A11:H11")
    sheet["A11"] = "DETALHES DO RELATÓRIO"
    sheet["A11"].fill = PatternFill("solid", fgColor=NAVY)
    sheet["A11"].font = Font(name="Arial", size=12, bold=True, color=WHITE)
    sheet["A11"].alignment = Alignment(vertical="center")
    sheet["A12"] = "Campo"
    sheet["E12"] = "Valor"
    sheet.merge_cells("A12:D12")
    sheet.merge_cells("E12:H12")
    for cell in (sheet["A12"], sheet["E12"]):
        cell.fill = PatternFill("solid", fgColor=PRIMARY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center")

    details = list(_flatten_scalars(payload))[:120]
    for row_index, (label, value) in enumerate(details, start=13):
        sheet.merge_cells(start_row=row_index, start_column=1, end_row=row_index, end_column=4)
        sheet.merge_cells(start_row=row_index, start_column=5, end_row=row_index, end_column=8)
        label_cell = sheet.cell(row_index, 1)
        value_cell = sheet.cell(row_index, 5)
        label_cell.value = label
        value_cell.value = _excel_value(value)
        fill = LIGHT if row_index % 2 == 0 else WHITE
        for cell in (label_cell, value_cell):
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.font = Font(name="Arial", size=10, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
        if isinstance(value_cell.value, datetime):
            value_cell.number_format = "dd/mm/yyyy hh:mm:ss"
        elif isinstance(value_cell.value, date):
            value_cell.number_format = "dd/mm/yyyy"

    for column in range(1, 9):
        sheet.column_dimensions[chr(64 + column)].width = 17
    sheet.row_dimensions[1].height = 29
    sheet.row_dimensions[2].height = 25
    sheet.row_dimensions[6].height = 24
    for row in range(7, 10):
        sheet.row_dimensions[row].height = 20
    _apply_print_setup(sheet)
    sheet.print_area = f"A1:H{max(15, 12 + len(details))}"


def _write_table_sheet(sheet, title: str, path: str, rows: list[dict[str, Any]]) -> None:
    columns: list[str] = []
    for row in rows:
        for key in row:
            if key not in columns:
                columns.append(key)
    header_span = max(8, len(columns))
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=header_span)
    sheet.cell(1, 1, f"GESTOR DE PEÇAS — {title.upper()}")
    sheet.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    sheet.cell(1, 1).font = Font(name="Arial", size=18, bold=True, color=WHITE)
    sheet.cell(1, 1).alignment = Alignment(vertical="center")
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=header_span)
    path_label = {
        "sectors": "Produção por setor",
        "quality.by_sector": "Qualidade › Resumo por setor",
        "quality.by_product": "Qualidade › Produção por produto",
        "quality.scrap_reasons": "Qualidade › Motivos de refugo",
        "quality.rework_reasons": "Qualidade › Motivos de retrabalho",
    }.get(path, " › ".join(_humanize(part) for part in path.split(".")))
    sheet.cell(2, 1, f"Origem dos dados: {path_label}")
    sheet.cell(2, 1).font = Font(name="Arial", size=9, color=MUTED)
    sheet.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT)

    header_row = 4
    for column_index, key in enumerate(columns, start=1):
        cell = sheet.cell(header_row, column_index, _column_label(key))
        cell.fill = PatternFill("solid", fgColor=PRIMARY)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index, row in enumerate(rows, start=header_row + 1):
        for column_index, key in enumerate(columns, start=1):
            cell = sheet.cell(row_index, column_index, _excel_value(row.get(key)))
            cell.font = Font(name="Arial", size=9, color=TEXT)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if isinstance(cell.value, datetime):
                cell.number_format = "dd/mm/yyyy hh:mm:ss"
            elif isinstance(cell.value, date):
                cell.number_format = "dd/mm/yyyy"
            elif isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, int):
                cell.number_format = "#,##0"

    if columns and rows:
        table_ref = f"A{header_row}:{sheet.cell(header_row + len(rows), len(columns)).coordinate}"
        table_name = "tbl" + re.sub(r"[^A-Za-z0-9]", "", sheet.title.title())
        table = Table(displayName=table_name[:240], ref=table_ref)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)

    for column_index, key in enumerate(columns, start=1):
        values = [_column_label(key), *(str(row.get(key, "")) for row in rows[:200])]
        width = min(42, max(11, max(len(value) for value in values) + 2))
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    for column_index in range(len(columns) + 1, header_span + 1):
        sheet.column_dimensions[get_column_letter(column_index)].width = 12
    sheet.row_dimensions[1].height = 30
    sheet.row_dimensions[4].height = 30
    _apply_print_setup(sheet, repeat_rows="4:4")
    if columns:
        sheet.print_area = f"A1:{sheet.cell(max(5, header_row + len(rows)), len(columns)).coordinate}"


def _section_rows(value: Any) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []

    def visit(current: Any, path: tuple[str, ...] = ()):
        if isinstance(current, dict):
            scalar_found = False
            for key, child in current.items():
                if isinstance(child, (dict, list)):
                    visit(child, (*path, str(key)))
                else:
                    scalar_found = True
                    rows.append({
                        "Grupo": " › ".join(_humanize(part) for part in path) or "Resumo",
                        "Campo": _humanize(str(key)),
                        "Valor": child,
                    })
            if not current and not scalar_found:
                return
            return
        if isinstance(current, list):
            group = " › ".join(_humanize(part) for part in path) or "Itens"
            for item in current:
                if isinstance(item, dict):
                    rows.append({"Grupo": group, **item})
                else:
                    rows.append({"Grupo": group, "Valor": item})
            return
        rows.append({
            "Grupo": " › ".join(_humanize(part) for part in path) or "Resumo",
            "Valor": current,
        })

    visit(value)
    return rows


def build_report_workbook(report_type: str, payload: Any, filters: AnalyticsFilter) -> bytes:
    """Gera um arquivo Excel apresentacional a partir do payload já calculado."""

    safe_type = str(report_type or "").strip().casefold().replace("-", "_")
    encoded = jsonable_encoder(payload)
    if not isinstance(encoded, dict):
        encoded = {"dados": encoded}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumo Executivo"
    sections = encoded.get("sections") if isinstance(encoded.get("sections"), dict) else None
    if sections:
        summary_payload = {
            **encoded,
            "sections": {"Resumo Executivo": sections.get("Resumo Executivo") or {}},
        }
        _style_summary(summary, safe_type, summary_payload, filters)
    else:
        _style_summary(summary, safe_type, encoded, filters)

    used = {"resumo executivo"}
    if sections:
        for section_name, section_payload in sections.items():
            if section_name == "Resumo Executivo":
                continue
            rows = _section_rows(section_payload)
            if not rows:
                continue
            title = _sheet_title(str(section_name), used)
            sheet = workbook.create_sheet(title)
            _write_table_sheet(sheet, title, str(section_name), rows)
    else:
        for path, rows in _table_sections(encoded):
            title = _sheet_title(path, used)
            sheet = workbook.create_sheet(title)
            _write_table_sheet(sheet, title, path, rows)

    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    output = BytesIO()
    try:
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
