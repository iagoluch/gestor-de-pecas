"""Relatórios MES executivos (Gerencial, Produção, Perdas, Indicadores, Dados Analíticos).

Cada função recebe o payload já produzido pelos serviços canônicos e decide
apenas a apresentação. Nenhum indicador industrial é recalculado aqui: quando o
backend não publica o dado, a aba diz isso com a justificativa oficial em vez de
mostrar zero.
"""

from __future__ import annotations

from datetime import datetime
from typing import Any, Sequence

from openpyxl.utils import get_column_letter

from .kit import (
    AMBER,
    AVAILABILITY_LABELS,
    BLUE,
    Card,
    Column,
    FORMATS,
    GREEN,
    MISSING,
    MUTED,
    RED,
    ReportSheet,
    TableRange,
    duration_text,
    grouped_chart,
    metric,
    number_text,
    pareto_chart,
    ranking_chart,
    technical_sheet,
    trend_chart,
)


# Um gráfico com categorias longas (motivo de parada, recurso) só fica legível
# ocupando a largura inteira do painel.
WIDE_CHART = 28.0

OVERVIEW_SHEET = "Visão Geral"
QUALITY_SHEET = "Qualidade dos Dados"
SERIES_SHEET = "Séries dos Gráficos"

FOOTER = (
    "Indicadores calculados pelos serviços canônicos do Gestor de Peças; esta planilha apresenta o "
    "resultado e não recalcula regra industrial. “—” indica dado ausente — nunca zero."
)

ISSUE_LABELS = {
    "sobreposicao_estados_incompativeis": "Estados físicos incompatíveis no mesmo intervalo",
    "recurso_em_setores_conflitantes": "Recurso em setores conflitantes",
    "estado_fisico_desconhecido": "Estado físico desconhecido",
    "recurso_produzindo_sem_op": "Produção sem OP associada",
    "fim_turno_sem_interrupcao_programada": "Fim de turno sem interrupção programada",
    "rateio_nao_conserva_tempo_fisico": "Rateio não conserva o tempo físico",
    "recurso_em_turno_sem_status": "Turno disponível sem estado registrado (recurso ocioso)",
}

SOURCE_LABELS = {
    "eventos_estado_recurso": "estado físico do recurso",
    "timeline_apontamentos_fallback": "timeline de apontamentos (fallback histórico)",
    "eventos_quantidade_producao": "eventos de quantidade de produção",
    "apontamentos_operacionais_fallback": "apontamentos operacionais (fallback)",
    "calendario_x_estado_recurso": "calendário confrontado com o estado do recurso",
}

SEVERITY_LABELS = {
    "info": "Informativo",
    "aviso": "Aviso",
    "erro": "Erro",
    "critico": "Crítico",
}

SEVERITY_TONES = (
    ("Crítico", RED),
    ("Erro", RED),
    ("Aviso", AMBER),
    ("Informativo", MUTED),
)

STATUS_TONES = (
    ("Finalizado", GREEN),
    ("Em processo", BLUE),
    ("Setup", AMBER),
    ("Retrabalho", AMBER),
    ("Parada", RED),
    ("Aguardando", MUTED),
)

QUALITY_COLUMNS = (
    Column("assunto", "Assunto", "text", 34),
    Column("situacao", "Situação", "text", 32),
    Column("detalhe", "Detalhe", "text", 74, wrap=True),
)


def _estado(availability: Any) -> str:
    """Traduz o código de disponibilidade do backend para linguagem de gestão."""

    codigo = str(availability or "").strip()
    return AVAILABILITY_LABELS.get(codigo, codigo)


def _fonte(source: Any) -> str:
    """Nome de tabela vira descrição legível fora da aba técnica."""

    codigo = str(source or "").strip()
    return SOURCE_LABELS.get(codigo, codigo or MISSING)


def _rows(value: Any) -> list[dict[str, Any]]:
    if not isinstance(value, list):
        return []
    return [item for item in value if isinstance(item, dict)]


def _dict(value: Any) -> dict[str, Any]:
    return value if isinstance(value, dict) else {}


def _number(value: Any):
    return value if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _sorted(rows: Sequence[dict[str, Any]], key: str) -> list[dict[str, Any]]:
    """Ordena do maior para o menor mantendo dados ausentes no fim."""

    return sorted(rows, key=lambda row: (_number(row.get(key)) is None, -(_number(row.get(key)) or 0.0)))


def _sum(rows: Sequence[dict[str, Any]], key: str):
    """Soma apenas totalizações de coluna já publicadas pelo backend."""

    valores = [_number(row.get(key)) for row in rows]
    presentes = [valor for valor in valores if valor is not None]
    return sum(presentes) if presentes else None


def _missing_notice(sheet: ReportSheet, metricas: Sequence[tuple[str, Any]]) -> list[dict[str, Any]]:
    """Agrupa indicadores indisponíveis em um aviso único e discreto."""

    faltantes = [(label, item) for label, item in metricas if item.missing]
    if not faltantes:
        return []
    razoes: dict[str, list[str]] = {}
    for label, item in faltantes:
        chave = str(item.reason or item.state_label).strip().rstrip(".")
        razoes.setdefault(chave, []).append(label)
    if len(faltantes) == 1:
        label, item = faltantes[0]
        motivo = str(item.reason or item.state_label).strip().rstrip(".")
        sheet.notice(f"{label}: {motivo}.", tone="muted")
    else:
        nomes = ", ".join(label for label, _ in faltantes)
        sheet.notice(
            f"Alguns indicadores não puderam ser calculados no período ({nomes}). "
            f"Consulte a aba “{QUALITY_SHEET}” para o motivo oficial de cada um.",
            tone="muted",
        )
    return [
        {"assunto": ", ".join(labels), "situacao": "Indicador indisponível", "detalhe": razao}
        for razao, labels in razoes.items()
    ]


def _quality_sheet(workbook, filters, generated_at: datetime, entradas: Sequence[dict[str, Any]]) -> None:
    if not entradas:
        return
    sheet = ReportSheet(workbook.create_sheet(QUALITY_SHEET), width=8, default_column_width=18)
    sheet.header(
        "Qualidade dos dados do período",
        filters,
        generated_at=generated_at,
        subtitle="O que limita a leitura deste relatório, em linguagem de gestão.",
    )
    sheet.section("Cobertura e limitações")
    sheet.table(QUALITY_COLUMNS, list(entradas), autofilter=False)
    sheet.print_setup()


def _series_sheet(workbook):
    sheet = ReportSheet(workbook.create_sheet(SERIES_SHEET), width=6, default_column_width=22)
    return sheet


def _hide(workbook, title: str) -> None:
    """Oculta a aba auxiliar; remove quando nenhum gráfico precisou dela."""

    if title not in workbook.sheetnames:
        return
    sheet = workbook[title]
    if sheet.max_row <= 1 and sheet.max_column <= 1:
        workbook.remove(sheet)
        return
    sheet.sheet_state = "hidden"


def _hide_column(sheet: ReportSheet, table: TableRange, key: str) -> None:
    """Mantém o código técnico rastreável sem ocupar a leitura do usuário."""

    letra = get_column_letter(table.column_index(key))
    sheet.ws.column_dimensions[letra].hidden = True


def _finish(sheet: ReportSheet, *, footer: bool = True) -> None:
    if footer:
        sheet.footer(FOOTER)
    sheet.print_setup()


# --------------------------------------------------------------------------
# 1. Relatório Gerencial
# --------------------------------------------------------------------------

SECTOR_COLUMNS = (
    Column("setor", "Setor", "text", 22),
    Column("producao_boa", "Produção boa", "int"),
    Column("refugo", "Refugo", "int"),
    Column("retrabalho", "Retrabalho", "int"),
    Column("ops", "OPs", "int"),
    Column("tempo_producao_segundos", "Tempo de produção", "duration"),
    Column("tempo_setup_segundos", "Tempo de setup", "duration"),
    Column("tempo_parada_segundos", "Tempo de parada", "duration"),
    Column("tempo_retrabalho_segundos", "Tempo de retrabalho", "duration"),
    Column("tempo_produtivo_segundos", "Tempo produtivo", "duration"),
)

RESOURCE_COLUMNS = (
    Column("recurso", "Recurso", "text", 20),
    Column("setor", "Setor", "text", 18),
    Column("oee", "OEE", "percent"),
    Column("disponibilidade", "Disponibilidade", "percent"),
    Column("performance", "Performance", "percent"),
    Column("ftt", "FTT", "percent"),
    Column("tempo_disponivel", "Tempo disponível", "duration"),
    Column("tempo_operacional", "Tempo operacional", "duration"),
    Column("tempo_produtivo", "Tempo produtivo", "duration"),
)


def _resource_rows(resource_kpis: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    """Explode o contrato por recurso em colunas — nunca JSON dentro da célula."""

    linhas = []
    for item in resource_kpis:
        metricas = _dict(item.get("metrics"))
        bases = _dict(item.get("kpi_time_bases"))
        linhas.append({
            "recurso": item.get("resource"),
            "setor": item.get("sector"),
            "oee": metric(metricas.get("oee")).value,
            "disponibilidade": metric(metricas.get("availability")).value,
            "performance": metric(metricas.get("performance")).value,
            "ftt": metric(metricas.get("ftt")).value,
            "tempo_disponivel": bases.get("available_seconds"),
            "tempo_operacional": bases.get("operational_seconds"),
            "tempo_produtivo": bases.get("productive_gross_seconds"),
        })
    return _sorted(linhas, "oee")


def build_gerencial(workbook, payload: dict[str, Any], filters, generated_at: datetime) -> None:
    overview = _dict(payload.get("overview"))
    production = _dict(overview.get("production"))
    kpis = _dict(overview.get("kpis"))
    estendidos = _dict(overview.get("kpis_estendidos"))
    insights = _dict(overview.get("insights"))
    setores = _rows(overview.get("sectors"))
    composicao = _dict(overview.get("time_composition"))
    perdas_oee = _dict(overview.get("kpi_losses_breakdown"))

    painel = ReportSheet(workbook.create_sheet(OVERVIEW_SHEET))
    aba_setores = ReportSheet(workbook.create_sheet("Setores"), width=10)
    aba_recursos = ReportSheet(workbook.create_sheet("Recursos"), width=9)
    aba_perdas = ReportSheet(workbook.create_sheet("Perdas"), width=8)
    series = _series_sheet(workbook)

    # --- Setores -----------------------------------------------------------
    aba_setores.header("Desempenho por setor", filters, generated_at=generated_at)
    linhas_setor = _sorted(setores, "producao_boa")
    tabela_setores = None
    if linhas_setor:
        totais = {campo.key: _sum(linhas_setor, campo.key) for campo in SECTOR_COLUMNS[1:]}
        totais[SECTOR_COLUMNS[0].key] = "Total"
        aba_setores.section("Produção e tempos por setor")
        tabela_setores = aba_setores.table(SECTOR_COLUMNS, linhas_setor, totals=totais)
        aba_setores.notice(
            "FTT e OEE por setor não são publicados pelo backend canônico; eles existem por "
            "recurso e para a fábrica. Consulte a aba “Recursos”.",
            tone="muted",
        )
    else:
        aba_setores.empty("Nenhum setor com apontamento no período filtrado.")
    _finish(aba_setores)

    # --- Recursos ----------------------------------------------------------
    aba_recursos.header("Indicadores por recurso", filters, generated_at=generated_at)
    linhas_recurso = _resource_rows(_rows(overview.get("resource_kpis")))
    tabela_recursos = None
    if linhas_recurso:
        aba_recursos.section(
            "Recursos com indicador canônico no período",
            hint="Ordenados por OEE. Cada indicador vem do mesmo cálculo usado nas telas gerenciais.",
        )
        tabela_recursos = aba_recursos.table(RESOURCE_COLUMNS, linhas_recurso)
    else:
        aba_recursos.empty("Nenhum recurso possui indicador calculável no período filtrado.")
    _finish(aba_recursos)

    # --- Perdas ------------------------------------------------------------
    aba_perdas.header("Perdas do período", filters, generated_at=generated_at)
    perdas_linhas = [
        {"componente": "Fora de turno", "segundos": perdas_oee.get("fora_de_turno_segundos")},
        {"componente": "Parada planejada", "segundos": perdas_oee.get("parada_planejada_segundos")},
        {"componente": "Parada não planejada", "segundos": perdas_oee.get("parada_nao_planejada_segundos")},
        {"componente": "Ritmo (performance)", "segundos": perdas_oee.get("ritmo_segundos")},
    ]
    aba_perdas.section(
        "Composição canônica das perdas de OEE",
        hint=(
            "Mesma decomposição usada pelo cálculo oficial de OEE. Em “Ritmo”, o backend publica a "
            "diferença entre o numerador e o denominador da performance: valor positivo indica "
            "performance acima de 100%."
        ),
    )
    aba_perdas.table(
        (Column("componente", "Componente", "text", 30), Column("segundos", "Tempo", "duration")),
        perdas_linhas,
        autofilter=False,
    )
    quantidade_perdas = [
        {"componente": "Refugo", "quantidade": perdas_oee.get("refugo_quantidade")},
        {"componente": "Retrabalho", "quantidade": perdas_oee.get("retrabalho_quantidade")},
    ]
    aba_perdas.section("Perdas de quantidade")
    aba_perdas.table(
        (Column("componente", "Componente", "text", 30), Column("quantidade", "Peças", "int")),
        quantidade_perdas,
        autofilter=False,
        freeze=False,
    )

    causas = _rows(insights.get("losses"))
    tabela_causas = None
    # Denominador canônico: tempo de parada já consolidado pela visão gerencial.
    total_parada = _number(_dict(overview.get("hours")).get("downtime_seconds"))
    if not total_parada:
        total_parada = _sum(causas, "impact_value")
    aba_perdas.section("Principais causas de parada")
    if causas:
        acumulado = 0.0
        linhas_causa = []
        for item in causas[:20]:
            segundos = _number(item.get("impact_value")) or 0.0
            acumulado += segundos
            linhas_causa.append({
                "motivo": item.get("cause") or "Não informado",
                "segundos": segundos,
                "participacao": (segundos / total_parada * 100.0) if total_parada else None,
                "acumulado": (acumulado / total_parada * 100.0) if total_parada else None,
            })
        tabela_causas = aba_perdas.table(
            (
                Column("motivo", "Motivo", "text", 40),
                Column("segundos", "Tempo perdido", "duration"),
                Column("participacao", "% do tempo de parada", "percent"),
                Column("acumulado", "% acumulado", "percent"),
            ),
            linhas_causa,
            freeze=False,
        )
    else:
        aba_perdas.empty("Nenhuma parada registrada no período.")

    criticos = _rows(insights.get("critical_resources"))
    if criticos:
        aba_perdas.section(
            "Recursos com maior tempo de parada",
            hint="Critério publicado pelo backend: tempo físico de parada no período, sem escore composto.",
        )
        aba_perdas.table(
            (
                Column("resource", "Recurso", "text", 24),
                Column("impact_value", "Tempo parado", "duration"),
                Column("criterion", "Critério", "text", 40),
            ),
            criticos[:15],
            autofilter=False,
            freeze=False,
        )
    _finish(aba_perdas)

    # --- Séries auxiliares --------------------------------------------------
    itens_composicao = [
        item for item in _rows(composicao.get("items"))
        if (_number(item.get("seconds")) or 0.0) > 0
    ]
    tabela_composicao = None
    if itens_composicao:
        tabela_composicao = series.table(
            (
                Column("label", "Composição do tempo", "text", 26),
                Column("seconds", "Tempo", "duration"),
                Column("percentage", "Participação", "percent"),
            ),
            sorted(itens_composicao, key=lambda item: -(_number(item.get("seconds")) or 0.0)),
            autofilter=False,
            freeze=False,
            zebra=False,
        )

    # --- Painel -------------------------------------------------------------
    painel.header("Relatório Gerencial", filters, generated_at=generated_at)
    oee = metric(kpis.get("oee"))
    disponibilidade = metric(kpis.get("availability"))
    performance = metric(kpis.get("performance"))
    ftt = metric(kpis.get("ftt"))
    utilizacao = metric(estendidos.get("utilization"))
    sem_registros = str(production.get("availability")) == "sem_registros"

    painel.section("Indicadores do período")
    painel.kpis([
        Card("Produção boa", None if sem_registros else production.get("good"), "int", "peças boas", "good",
             production.get("reason") or "Sem registros de produção no período"),
        Card("OPs executadas", production.get("ops"), "int", "ordens com apontamento", "neutral"),
        Card("Recursos utilizados", production.get("resources"), "int", "recursos com registro", "neutral"),
        Card("OEE", oee.value, "percent", "cálculo canônico do Gestor", "info", oee.reason or oee.state_label),
        Card("Disponibilidade", disponibilidade.value, "percent", "tempo disponível de calendário", "info", disponibilidade.reason or disponibilidade.state_label),
        Card("Performance", performance.value, "percent", "ritmo sobre o tempo padrão", "info", performance.reason or performance.state_label),
        Card("FTT", ftt.value, "percent", "qualidade de primeira passagem", "good", ftt.reason or ftt.state_label),
        Card("Utilização", utilizacao.value, "percent", "carga sobre a capacidade", "info", utilizacao.reason or utilizacao.state_label),
    ])
    limitacoes = _missing_notice(painel, [
        ("OEE", oee), ("Disponibilidade", disponibilidade), ("Performance", performance),
        ("FTT", ftt), ("Utilização", utilizacao),
    ])

    destaques = []
    for item in _rows(overview.get("sector_highlights")):
        valor = _number(item.get("value"))
        unidade = str(item.get("unit") or "")
        # "Maior refugo: 0 peças" não é destaque: é ausência de ocorrência.
        if str(item.get("key") or "").startswith("maior_") and not valor:
            continue
        texto = duration_text(valor) if unidade == "s" else f"{number_text(valor, 0)} {unidade}".strip()
        destaques.append((str(item.get("label") or "Destaque"), f"{item.get('sector') or MISSING} — {texto}"))
    if causas:
        maior = causas[0]
        destaques.append(("Maior causa de parada", f"{maior.get('cause') or 'Não informado'} — {duration_text(maior.get('impact_value'))}"))
    if criticos:
        destaques.append(("Recurso com maior parada", f"{criticos[0].get('resource')} — {duration_text(criticos[0].get('impact_value'))}"))
    excecoes = _number(insights.get("exception_count"))
    if excecoes:
        destaques.append(("Exceções abertas", f"{number_text(excecoes, 0)} registradas pelo motor de exceções"))
    if destaques:
        painel.section("Destaques do período", hint="Fatos objetivos extraídos dos dados; sem julgamento de desempenho.")
        painel.facts(destaques)

    painel.section("Gráficos")
    ancora = painel.row
    if tabela_setores is not None:
        ranking_chart(
            painel.ws, tabela_setores,
            title="Produção boa por setor (peças)",
            label_key="setor", value_key="producao_boa",
            anchor=f"A{ancora}", color=GREEN, number_format=FORMATS["int"],
        )
    if tabela_composicao is not None:
        ranking_chart(
            painel.ws, tabela_composicao,
            title="Composição do tempo físico medido",
            label_key="label", value_key="seconds",
            anchor=f"G{ancora}", color=BLUE, number_format=FORMATS["duration"],
        )
    ancora += 17
    if tabela_recursos is not None and any(linha.get("oee") is not None for linha in linhas_recurso):
        grouped_chart(
            painel.ws, tabela_recursos,
            title="Disponibilidade, performance e FTT por recurso",
            label_key="recurso", value_keys=("disponibilidade", "performance", "ftt"),
            anchor=f"A{ancora}", colors=(BLUE, AMBER, GREEN),
            number_format=FORMATS["percent"], max_rows=12,
            width=WIDE_CHART,
        )
        ancora += 17
    if tabela_causas is not None:
        pareto_chart(
            painel.ws, tabela_causas,
            title="Pareto das causas de parada",
            label_key="motivo", value_key="segundos", cumulative_key="acumulado",
            anchor=f"A{ancora}", number_format=FORMATS["duration"],
            width=WIDE_CHART,
        )
        ancora += 17
    else:
        painel.row = ancora
        painel.empty("Nenhuma parada registrada no período — o Pareto de causas não é gerado.")
        ancora = painel.row
    painel.row = max(painel.row, ancora)
    if itens_composicao:
        painel.notice(
            "“Fila / espera” é o tempo físico em que o recurso permaneceu disponível sem execução de OP. "
            "Ele domina a composição em períodos com poucos apontamentos e não representa parada registrada.",
            tone="muted",
        )
    _finish(painel)

    # --- Qualidade dos dados e anexos técnicos ------------------------------
    qualidade = list(limitacoes)
    data_quality = _dict(overview.get("data_quality"))
    qualidade.append({
        "assunto": "Cobertura de calendário produtivo",
        "situacao": "Configurado" if data_quality.get("calendar_configured") else "Não configurado",
        "detalhe": (
            "O calendário produtivo define o tempo disponível usado por disponibilidade, OEE e utilização."
        ),
    })
    qualidade.append({
        "assunto": "Base de fatos analisada",
        "situacao": f"{number_text(data_quality.get('facts'), 0)} apontamentos • {number_text(data_quality.get('physical_state_rows'), 0)} estados físicos",
        "detalhe": (
            f"Fonte física: {_fonte(data_quality.get('physical_state_source'))}. "
            f"Cobertura declarada pelo backend: {_estado(data_quality.get('status'))}."
        ),
    })
    plano = _dict(overview.get("production_plan"))
    if plano.get("value") is None and plano.get("reason"):
        qualidade.append({
            "assunto": "Planejado x realizado",
            "situacao": "Indisponível",
            "detalhe": str(plano.get("reason")),
        })
    for item in _rows(insights.get("limitations")):
        codigo = str(item.get("code") or "").strip()
        qualidade.append({
            "assunto": "Limitação declarada pelo backend",
            "situacao": "Informativo",
            "detalhe": str(item.get("message") or "") + (f" (código: {codigo})" if codigo else ""),
        })
    rateio = _dict(overview.get("rateio"))
    if rateio.get("reason"):
        qualidade.append({
            "assunto": "Rateio de tempo por OP",
            "situacao": _estado(rateio.get("availability")),
            "detalhe": str(rateio.get("reason")),
        })
    _quality_sheet(workbook, filters, generated_at, qualidade)
    technical_sheet(workbook, payload, filters, generated_at=generated_at)
    _hide(workbook, SERIES_SHEET)


# --------------------------------------------------------------------------
# 2. Relatório de Produção
# --------------------------------------------------------------------------

PRODUCTION_SECTOR_COLUMNS = (
    Column("setor", "Setor", "text", 22),
    Column("producao_boa", "Produção boa", "int"),
    Column("refugo", "Refugo", "int"),
    Column("retrabalho", "Retrabalho", "int"),
    Column("ops", "OPs", "int"),
    Column("tempo_producao_segundos", "Tempo de produção", "duration"),
    Column("tempo_setup_segundos", "Tempo de setup", "duration"),
    Column("tempo_parada_segundos", "Tempo de parada", "duration"),
)

PRODUCT_COLUMNS = (
    Column("produto", "Produto", "text", 20),
    Column("descricao", "Descrição", "text", 44, wrap=True),
    Column("boa", "Quantidade boa", "int"),
    Column("refugo", "Refugo", "int"),
    Column("retrabalho", "Retrabalho", "int"),
)

ORDER_COLUMNS = (
    Column("op", "OP", "text", 16),
    Column("produto", "Produto", "text", 16),
    Column("descricao", "Descrição", "text", 40),
    Column("operacao_atual", "Operação", "text", 12),
    Column("descricao_operacao", "Descrição da operação", "text", 24),
    Column("setor", "Setor", "text", 16),
    Column("recurso_real", "Recurso", "text", 16),
    Column("quantidade_planejada", "Qtd. planejada", "int"),
    Column("quantidade_boa", "Qtd. boa", "int"),
    Column("refugo", "Refugo", "int"),
    Column("saldo_quantidade", "Saldo", "int"),
    Column("progresso_percentual", "Progresso", "percent"),
    Column("inicio_real", "Início real", "datetime"),
    Column("fim_real", "Fim real", "datetime"),
    Column("status", "Status", "text", 16, tones=STATUS_TONES),
)


def build_producao(workbook, payload: dict[str, Any], filters, generated_at: datetime) -> None:
    production = _dict(payload.get("production"))
    qualidade = _dict(payload.get("quality"))
    setores = _rows(payload.get("sectors"))
    ordens = _dict(payload.get("orders"))
    itens_ordens = _rows(ordens.get("items"))

    painel = ReportSheet(workbook.create_sheet(OVERVIEW_SHEET))
    aba_setores = ReportSheet(workbook.create_sheet("Setores"), width=8)
    aba_produtos = ReportSheet(workbook.create_sheet("Produtos"), width=6)
    aba_ordens = ReportSheet(workbook.create_sheet("Ordens"), width=15)

    # --- Setores -----------------------------------------------------------
    aba_setores.header("Produção por setor", filters, generated_at=generated_at)
    linhas_setor = _sorted(setores, "producao_boa")
    tabela_setores = None
    if linhas_setor:
        totais = {campo.key: _sum(linhas_setor, campo.key) for campo in PRODUCTION_SECTOR_COLUMNS[1:]}
        totais[PRODUCTION_SECTOR_COLUMNS[0].key] = "Total"
        aba_setores.section("Quantidades e tempos por setor")
        tabela_setores = aba_setores.table(PRODUCTION_SECTOR_COLUMNS, linhas_setor, totals=totais)
    else:
        aba_setores.empty("Nenhum setor com produção registrada no período filtrado.")
    _finish(aba_setores)

    # --- Produtos ----------------------------------------------------------
    descricoes = {
        str(item.get("produto")): item.get("descricao")
        for item in itens_ordens
        if item.get("produto") and item.get("descricao")
    }
    linhas_produto = [
        {**item, "descricao": descricoes.get(str(item.get("produto")))}
        for item in _rows(qualidade.get("by_product"))
    ]
    linhas_produto = _sorted(linhas_produto, "boa")
    aba_produtos.header("Produção por produto", filters, generated_at=generated_at)
    tabela_produtos = None
    if linhas_produto:
        totais = {campo.key: _sum(linhas_produto, campo.key) for campo in PRODUCT_COLUMNS[2:]}
        totais[PRODUCT_COLUMNS[0].key] = "Total"
        aba_produtos.section("Ordenado pela quantidade boa produzida")
        tabela_produtos = aba_produtos.table(PRODUCT_COLUMNS, linhas_produto, totals=totais)
    else:
        aba_produtos.empty("Nenhum evento de quantidade por produto no período filtrado.")
    _finish(aba_produtos)

    # --- Ordens ------------------------------------------------------------
    aba_ordens.header("Ordens executadas no período", filters, generated_at=generated_at)
    if itens_ordens:
        aba_ordens.section(
            "Visão operacional das OPs",
            hint="IDs internos, operadores e fontes de dado permanecem na aba técnica.",
        )
        aba_ordens.table(ORDER_COLUMNS, itens_ordens)
    else:
        aba_ordens.empty("Nenhuma ordem com apontamento no período filtrado.")
    _finish(aba_ordens)

    # --- Painel ------------------------------------------------------------
    painel.header("Relatório de Produção", filters, generated_at=generated_at)
    ftt = metric(qualidade.get("ftt"))
    sem_registros = str(production.get("availability")) == "sem_registros"
    painel.section("Produção do período")
    painel.kpis([
        Card("Produção boa", None if sem_registros else production.get("good"), "int", "peças boas", "good",
             production.get("reason") or "Sem registros de produção no período"),
        Card("Refugo", None if sem_registros else production.get("scrap"), "int", "peças refugadas", "bad",
             production.get("reason") or "Sem registros de produção no período"),
        Card("Retrabalho", None if sem_registros else production.get("rework"), "int", "peças retrabalhadas", "warn",
             production.get("reason") or "Sem registros de produção no período"),
        Card("FTT", ftt.value, "percent", "qualidade de primeira passagem", "good", ftt.reason or ftt.state_label),
        Card("OPs executadas", production.get("ops"), "int", "ordens com apontamento", "neutral"),
        Card("Recursos utilizados", production.get("resources"), "int", "recursos com registro", "neutral"),
    ])
    limitacoes = _missing_notice(painel, [("FTT", ftt)])

    destaques = []
    if linhas_setor:
        topo = linhas_setor[0]
        destaques.append(("Maior produção", f"{topo.get('setor')} — {number_text(topo.get('producao_boa'), 0)} peças"))
    if linhas_produto:
        topo = linhas_produto[0]
        destaques.append(("Produto mais produzido", f"{topo.get('produto')} — {number_text(topo.get('boa'), 0)} peças"))
    if itens_ordens:
        concluidas = [item for item in itens_ordens if str(item.get("status") or "").casefold() == "finalizado"]
        destaques.append(("Apontamentos finalizados", f"{len(concluidas)} de {len(itens_ordens)} registros do período"))
    if destaques:
        painel.section("Destaques do período")
        painel.facts(destaques)

    painel.section("Gráficos")
    ancora = painel.row
    if tabela_setores is not None:
        ranking_chart(
            painel.ws, tabela_setores,
            title="Produção boa por setor (peças)",
            label_key="setor", value_key="producao_boa",
            anchor=f"A{ancora}", color=GREEN, number_format=FORMATS["int"],
        )
    if tabela_produtos is not None:
        ranking_chart(
            painel.ws, tabela_produtos,
            title="Top produtos por quantidade boa",
            label_key="produto", value_key="boa",
            anchor=f"G{ancora}", color=BLUE, number_format=FORMATS["int"], max_rows=10,
        )
    ancora += 17
    refugo_total = _sum(linhas_setor, "refugo") or 0
    retrabalho_total = _sum(linhas_setor, "retrabalho") or 0
    if tabela_setores is not None and (refugo_total or retrabalho_total):
        grouped_chart(
            painel.ws, tabela_setores,
            title="Boa, refugo e retrabalho por setor",
            label_key="setor", value_keys=("producao_boa", "refugo", "retrabalho"),
            anchor=f"A{ancora}", colors=(GREEN, RED, AMBER),
            number_format=FORMATS["int"], stacked=True,
        )
        painel.row = ancora + 17
    else:
        painel.row = ancora
        painel.empty("Nenhum refugo ou retrabalho registrado no período — o gráfico de composição não é gerado.")
    _finish(painel)

    # --- Qualidade dos dados ------------------------------------------------
    qualidade_linhas = list(limitacoes)
    if qualidade.get("reason"):
        qualidade_linhas.append({
            "assunto": "Fonte das quantidades",
            "situacao": _estado(qualidade.get("availability")),
            "detalhe": str(qualidade.get("reason")),
        })
    if production.get("reason"):
        qualidade_linhas.append({
            "assunto": "Produção do período",
            "situacao": _estado(production.get("availability")),
            "detalhe": str(production.get("reason")),
        })
    qualidade_linhas.append({
        "assunto": "Planejado x realizado",
        "situacao": "Disponível apenas por operação",
        "detalhe": (
            "O total planejado da fábrica depende do contrato corporativo Protheus/TOTVS e não é "
            "inferido a partir de apontamentos por operação. A comparação existe linha a linha na aba “Ordens”."
        ),
    })
    qualidade_linhas.append({
        "assunto": "Produção ao longo do período",
        "situacao": "Série não publicada",
        "detalhe": "O backend não publica série temporal de produção; nenhuma curva é estimada nesta planilha.",
    })
    _quality_sheet(workbook, filters, generated_at, qualidade_linhas)
    technical_sheet(workbook, payload, filters, generated_at=generated_at)


# --------------------------------------------------------------------------
# 3. Relatório de Perdas
# --------------------------------------------------------------------------

STOP_COLUMNS = (
    Column("recurso", "Recurso", "text", 18),
    Column("setor", "Setor", "text", 16),
    Column("motivo", "Motivo", "text", 32),
    Column("inicio", "Início", "datetime"),
    Column("fim", "Fim", "datetime"),
    Column("segundos", "Duração", "duration"),
    Column("programada", "Programada", "text", 14),
    Column("op", "OP", "text", 16),
)


def build_perdas(workbook, payload: dict[str, Any], filters, generated_at: datetime) -> None:
    paradas = _dict(payload.get("paradas"))
    setups = _dict(payload.get("setup"))
    qualidade = _dict(payload.get("qualidade"))
    tempos = _dict(payload.get("tempos"))
    perdas = _dict(payload.get("losses"))
    totais_qualidade = _dict(qualidade.get("totals"))

    painel = ReportSheet(workbook.create_sheet(OVERVIEW_SHEET))
    aba_motivos = ReportSheet(workbook.create_sheet("Motivos de Parada"), width=6)
    aba_setores = ReportSheet(workbook.create_sheet("Perdas por Setor"), width=8)
    aba_qualidade = ReportSheet(workbook.create_sheet("Refugo e Retrabalho"), width=6)

    total_parada = _number(paradas.get("total_seconds"))
    registros_por_motivo: dict[str, int] = {}
    for item in _rows(paradas.get("items")):
        chave = str(item.get("motivo") or "Não informado")
        registros_por_motivo[chave] = registros_por_motivo.get(chave, 0) + 1

    # --- Motivos -----------------------------------------------------------
    aba_motivos.header("Causas de parada", filters, generated_at=generated_at)
    motivos = _rows(paradas.get("by_reason"))
    tabela_motivos = None
    if motivos:
        acumulado = 0.0
        linhas = []
        for item in motivos:
            segundos = _number(item.get("segundos")) or 0.0
            acumulado += segundos
            linhas.append({
                "motivo": item.get("motivo") or "Não informado",
                "segundos": segundos,
                "registros": registros_por_motivo.get(str(item.get("motivo") or "Não informado")),
                "participacao": (segundos / total_parada * 100.0) if total_parada else None,
                "acumulado": (acumulado / total_parada * 100.0) if total_parada else None,
            })
        aba_motivos.section(
            "Ordenado pelo tempo perdido",
            hint=(
                "Tempo consolidado sem sobreposição pelo backend. “Registros” é a contagem de estados "
                "físicos com o mesmo motivo e serve para dimensionar a recorrência."
            ),
        )
        tabela_motivos = aba_motivos.table(
            (
                Column("motivo", "Motivo", "text", 40),
                Column("segundos", "Tempo perdido", "duration"),
                Column("registros", "Registros", "int"),
                Column("participacao", "% do tempo de parada", "percent"),
                Column("acumulado", "% acumulado", "percent"),
            ),
            linhas,
        )
    else:
        aba_motivos.empty("Nenhuma parada registrada no período.")
    _finish(aba_motivos)

    # --- Perdas por setor e recurso ----------------------------------------
    aba_setores.header("Perdas por setor e por recurso", filters, generated_at=generated_at)
    linhas_setor = _sorted(_rows(tempos.get("by_sector")), "parada")
    tabela_setores = None
    if linhas_setor:
        aba_setores.section("Tempo não produtivo por setor")
        tabela_setores = aba_setores.table(
            (
                Column("setor", "Setor", "text", 22),
                Column("parada", "Parada", "duration"),
                Column("setup", "Setup", "duration"),
                Column("retrabalho", "Retrabalho", "duration"),
                Column("fila", "Fila / espera", "duration"),
                Column("fora_turno", "Fora de turno", "duration"),
                Column("producao", "Produção", "duration"),
            ),
            linhas_setor,
        )
    else:
        aba_setores.empty("Sem composição de tempo por setor no período filtrado.")

    linhas_recurso = _sorted(_rows(paradas.get("by_resource")), "segundos")
    tabela_recursos = None
    if linhas_recurso:
        aba_setores.section("Recursos com maior tempo de parada")
        tabela_recursos = aba_setores.table(
            (Column("recurso", "Recurso", "text", 24), Column("segundos", "Tempo parado", "duration")),
            linhas_recurso,
            autofilter=False,
            freeze=False,
        )
    _finish(aba_setores)

    # --- Refugo e retrabalho ------------------------------------------------
    aba_qualidade.header("Refugo e retrabalho", filters, generated_at=generated_at)
    refugo = _number(totais_qualidade.get("refugo"))
    retrabalho = _number(totais_qualidade.get("retrabalho"))
    if (refugo or 0) or (retrabalho or 0):
        por_setor = _sorted(_rows(qualidade.get("by_sector")), "refugo")
        aba_qualidade.section("Por setor")
        aba_qualidade.table(
            (
                Column("setor", "Setor", "text", 22),
                Column("boa", "Produção boa", "int"),
                Column("refugo", "Refugo", "int"),
                Column("retrabalho", "Retrabalho", "int"),
            ),
            por_setor,
        )
        motivos_refugo = _rows(qualidade.get("scrap_reasons"))
        if motivos_refugo:
            aba_qualidade.section("Motivos de refugo")
            aba_qualidade.table(
                (Column("motivo", "Motivo", "text", 40), Column("quantidade", "Peças", "int")),
                motivos_refugo,
                autofilter=False,
                freeze=False,
            )
        motivos_retrabalho = _rows(qualidade.get("rework_reasons"))
        if motivos_retrabalho:
            aba_qualidade.section("Motivos de retrabalho")
            aba_qualidade.table(
                (Column("motivo", "Motivo", "text", 40), Column("quantidade", "Peças", "int")),
                motivos_retrabalho,
                autofilter=False,
                freeze=False,
            )
    else:
        aba_qualidade.empty(
            "Nenhum refugo ou retrabalho registrado no período — há eventos de quantidade, "
            "e todos foram classificados como produção boa."
        )
    _finish(aba_qualidade)

    # --- Detalhe das paradas ------------------------------------------------
    itens_parada = _rows(paradas.get("items"))
    if itens_parada:
        aba_detalhe = ReportSheet(workbook.create_sheet("Paradas"), width=8)
        aba_detalhe.header("Paradas registradas", filters, generated_at=generated_at)
        aba_detalhe.section("Registro físico de cada parada")
        aba_detalhe.table(STOP_COLUMNS, itens_parada)
        _finish(aba_detalhe)

    # --- Painel -------------------------------------------------------------
    painel.header("Relatório de Perdas", filters, generated_at=generated_at)
    painel.section("Perdas do período")
    painel.kpis([
        Card("Tempo de parada", paradas.get("total_seconds"), "duration_card", "tempo físico consolidado", "bad"),
        Card("Ocorrências", paradas.get("count"), "int", "registros de parada", "bad"),
        Card("Parada planejada", perdas.get("planned_downtime_seconds"), "duration_card", "não afeta o OEE", "warn"),
        Card("Parada não planejada", perdas.get("unplanned_downtime_seconds"), "duration_card", "afeta o OEE", "bad"),
        Card("Setup", setups.get("total_seconds"), "duration_card", f"{number_text(setups.get('count'), 0)} ocorrências", "warn"),
        Card("Retrabalho", perdas.get("rework_seconds"), "duration_card", "tempo em retrabalho", "warn"),
        Card("Refugo", totais_qualidade.get("refugo"), "int", "peças refugadas", "bad"),
        Card("Peças retrabalhadas", totais_qualidade.get("retrabalho"), "int", "peças em retrabalho", "warn"),
    ])

    perdas_oee = _dict(perdas.get("oee_losses"))
    if perdas_oee:
        painel.section(
            "Composição canônica das perdas de OEE",
            hint=(
                "Em “Ritmo”, o backend publica a diferença entre o numerador e o denominador da "
                "performance: valor positivo indica performance acima de 100%."
            ),
        )
        painel.facts([
            ("Fora de turno", duration_text(perdas_oee.get("fora_de_turno_segundos"))),
            ("Parada planejada", duration_text(perdas_oee.get("parada_planejada_segundos"))),
            ("Parada não planejada", duration_text(perdas_oee.get("parada_nao_planejada_segundos"))),
            ("Ritmo (performance)", duration_text(perdas_oee.get("ritmo_segundos"))),
            ("Refugo", f"{number_text(perdas_oee.get('refugo_quantidade'), 0)} peças"),
            ("Retrabalho", f"{number_text(perdas_oee.get('retrabalho_quantidade'), 0)} peças"),
        ])

    destaques = []
    if motivos:
        destaques.append(("Maior causa de parada", f"{motivos[0].get('motivo')} — {duration_text(motivos[0].get('segundos'))}"))
    if linhas_recurso:
        destaques.append(("Recurso com maior parada", f"{linhas_recurso[0].get('recurso')} — {duration_text(linhas_recurso[0].get('segundos'))}"))
    if linhas_setor and (_number(linhas_setor[0].get("parada")) or 0) > 0:
        destaques.append(("Setor com maior parada", f"{linhas_setor[0].get('setor')} — {duration_text(linhas_setor[0].get('parada'))}"))
    if destaques:
        painel.section("Destaques do período")
        painel.facts(destaques)

    painel.section("Gráficos")
    ancora = painel.row
    if tabela_motivos is not None:
        pareto_chart(
            painel.ws, tabela_motivos,
            title="Pareto das causas de parada",
            label_key="motivo", value_key="segundos", cumulative_key="acumulado",
            anchor=f"A{ancora}", number_format=FORMATS["duration"],
            width=WIDE_CHART,
        )
        ancora += 17
    if tabela_recursos is not None:
        ranking_chart(
            painel.ws, tabela_recursos,
            title="Tempo parado por recurso",
            label_key="recurso", value_key="segundos",
            anchor=f"A{ancora}", color=RED, number_format=FORMATS["duration"],
        )
    if tabela_setores is not None and (_sum(linhas_setor, "parada") or 0) > 0:
        grouped_chart(
            painel.ws, tabela_setores,
            title="Tempo não produtivo por setor",
            label_key="setor", value_keys=("parada", "setup", "retrabalho"),
            anchor=f"G{ancora}", colors=(RED, AMBER, BLUE),
            number_format=FORMATS["duration"], stacked=True,
        )
    if tabela_motivos is None and tabela_recursos is None and tabela_setores is None:
        painel.row = ancora
        painel.empty("Nenhuma parada registrada no período — nenhum gráfico de perdas é gerado.")
    else:
        painel.row = max(painel.row, ancora + 17)
    _finish(painel)

    # --- Qualidade dos dados ------------------------------------------------
    qualidade_linhas = []
    if paradas.get("overlap_removed_seconds"):
        qualidade_linhas.append({
            "assunto": "Sobreposição de paradas",
            "situacao": duration_text(paradas.get("overlap_removed_seconds")),
            "detalhe": "Tempo removido pela consolidação física para que paradas simultâneas não sejam contadas duas vezes.",
        })
    if qualidade.get("reason"):
        qualidade_linhas.append({
            "assunto": "Fonte de refugo e retrabalho",
            "situacao": _estado(qualidade.get("availability")),
            "detalhe": str(qualidade.get("reason")),
        })
    if tempos.get("reason"):
        qualidade_linhas.append({
            "assunto": "Fonte da composição de tempo",
            "situacao": _estado(tempos.get("availability")),
            "detalhe": str(tempos.get("reason")),
        })
    qualidade_linhas.append({
        "assunto": "Tendência de perdas no período",
        "situacao": "Série não publicada",
        "detalhe": "O backend não publica série temporal de paradas; nenhuma curva de tendência é estimada aqui.",
    })
    _quality_sheet(workbook, filters, generated_at, qualidade_linhas)
    technical_sheet(workbook, payload, filters, generated_at=generated_at)


# --------------------------------------------------------------------------
# 4. Relatório de Indicadores
# --------------------------------------------------------------------------

CAPACITY_COLUMNS = (
    Column("codigo", "Recurso", "text", 18),
    Column("nome", "Descrição", "text", 30),
    Column("tipo_setor", "Setor", "text", 16),
    Column("capacidade_segundos", "Capacidade", "duration"),
    Column("carga_segundos", "Carga", "duration"),
    Column("capacidade_restante_segundos", "Disponível", "duration"),
    Column("utilizacao_percentual", "Utilização", "percent"),
    Column("fila", "Fila", "int"),
)


def build_indicadores(workbook, payload: dict[str, Any], filters, generated_at: datetime) -> None:
    analytics = _dict(payload.get("analytics"))
    oee = _dict(analytics.get("oee"))
    kpis = _dict(analytics.get("kpis"))
    estendidos = _dict(analytics.get("kpis_estendidos"))
    capacidade = _dict(analytics.get("capacidade"))
    confiabilidade = _dict(analytics.get("confiabilidade"))
    qualidade = _dict(analytics.get("qualidade"))
    evolucao = _dict(oee.get("evolution"))

    painel = ReportSheet(workbook.create_sheet(OVERVIEW_SHEET))
    aba_recursos = ReportSheet(workbook.create_sheet("Recursos"), width=9)
    aba_capacidade = ReportSheet(workbook.create_sheet("Utilização"), width=8)
    aba_confiabilidade = ReportSheet(workbook.create_sheet("Confiabilidade"), width=6)
    series = _series_sheet(workbook)

    # --- Indicadores por recurso -------------------------------------------
    linhas_recurso = _resource_rows(_rows(analytics.get("recursos")))
    aba_recursos.header("Indicadores por recurso", filters, generated_at=generated_at)
    tabela_recursos = None
    if linhas_recurso:
        aba_recursos.section("Ordenados por OEE", hint="Mesma instância de cálculo usada pelas telas gerenciais.")
        tabela_recursos = aba_recursos.table(RESOURCE_COLUMNS, linhas_recurso)
    else:
        aba_recursos.empty("Nenhum recurso possui indicador calculável no período filtrado.")
    _finish(aba_recursos)

    # --- Utilização / capacidade -------------------------------------------
    itens_capacidade = [
        item for item in _rows(capacidade.get("items"))
        if item.get("utilizacao_percentual") is not None
    ]
    itens_capacidade = _sorted(itens_capacidade, "utilizacao_percentual")
    aba_capacidade.header("Utilização da capacidade", filters, generated_at=generated_at)
    tabela_capacidade = None
    if itens_capacidade:
        aba_capacidade.section(
            "Recursos com calendário produtivo configurado",
            hint="Sem calendário não existe tempo disponível para comparar; esses recursos ficam fora da tabela.",
        )
        tabela_capacidade = aba_capacidade.table(CAPACITY_COLUMNS, itens_capacidade)
    else:
        aba_capacidade.empty("Nenhum recurso do filtro possui calendário produtivo configurado.")
    sem_calendario = _number(capacidade.get("recursos_sem_calendario"))
    if sem_calendario:
        aba_capacidade.notice(
            f"{number_text(sem_calendario, 0)} de {number_text(capacidade.get('recursos'), 0)} recursos "
            "ainda não possuem calendário/turno produtivo cadastrado.",
            tone="warn",
        )
    _finish(aba_capacidade)

    # --- Confiabilidade -----------------------------------------------------
    aba_confiabilidade.header("Confiabilidade de equipamento", filters, generated_at=generated_at)
    por_recurso = _rows(confiabilidade.get("por_recurso"))
    if por_recurso:
        aba_confiabilidade.section("Falhas por recurso")
        aba_confiabilidade.table(
            (
                Column("recurso", "Recurso", "text", 24),
                Column("falhas", "Falhas", "int"),
                Column("segundos", "Tempo em reparo", "duration"),
            ),
            por_recurso,
        )
        motivos_falha = _rows(confiabilidade.get("motivos"))
        if motivos_falha:
            aba_confiabilidade.section("Motivos classificados como falha")
            aba_confiabilidade.table(
                (Column("motivo", "Motivo", "text", 40), Column("quantidade", "Ocorrências", "int")),
                motivos_falha,
                autofilter=False,
                freeze=False,
            )
    else:
        aba_confiabilidade.empty(str(confiabilidade.get("reason") or "Nenhuma falha de equipamento registrada no período."))
    aba_confiabilidade.notice(str(confiabilidade.get("taxonomia") or ""), tone="muted")
    _finish(aba_confiabilidade)

    # --- Evolução do OEE ----------------------------------------------------
    pontos = [
        {
            "periodo": item.get("period_start") or item.get("at"),
            "oee": item.get("value"),
        }
        for item in _rows(evolucao.get("points"))
        if item.get("value") is not None
    ]
    tabela_evolucao = None
    if len(pontos) >= 2:
        aba_evolucao = ReportSheet(workbook.create_sheet("Evolução do OEE"), width=4)
        aba_evolucao.header("Evolução do OEE", filters, generated_at=generated_at)
        aba_evolucao.section(
            "Pontos calculáveis no período",
            hint=f"Agrupamento de {number_text(evolucao.get('bucket_days'), 0)} dia(s) definido pelo backend.",
        )
        tabela_evolucao = aba_evolucao.table(
            (Column("periodo", "Período", "date", 16), Column("oee", "OEE", "percent")),
            pontos,
        )
        _finish(aba_evolucao)

    # --- Comparação de indicadores (série auxiliar) -------------------------
    oee_metric = metric(oee)
    disponibilidade = metric(kpis.get("availability"))
    performance = metric(kpis.get("performance"))
    ftt = metric(kpis.get("ftt") or qualidade.get("ftt"))
    utilizacao = metric(estendidos.get("utilization"))
    produtividade = metric(estendidos.get("productivity"))

    comparacao = [
        {"indicador": label, "valor": item.value}
        for label, item in (
            ("OEE", oee_metric), ("Disponibilidade", disponibilidade),
            ("Performance", performance), ("FTT", ftt),
        )
        if not item.missing
    ]
    tabela_comparacao = None
    if len(comparacao) >= 2:
        tabela_comparacao = series.table(
            (Column("indicador", "Indicador", "text", 22), Column("valor", "Valor", "percent")),
            comparacao,
            autofilter=False,
            freeze=False,
            zebra=False,
        )

    # --- Painel -------------------------------------------------------------
    painel.header("Relatório de Indicadores", filters, generated_at=generated_at)
    mtbf = confiabilidade.get("mtbf_segundos")
    mttr = confiabilidade.get("mttr_segundos")
    painel.section("Indicadores do período")
    painel.kpis([
        Card("OEE", oee_metric.value, "percent", "cálculo canônico do Gestor", "info", oee_metric.reason or oee_metric.state_label),
        Card("Disponibilidade", disponibilidade.value, "percent", "tempo disponível de calendário", "info", disponibilidade.reason or disponibilidade.state_label),
        Card("Performance", performance.value, "percent", "ritmo sobre o tempo padrão", "info", performance.reason or performance.state_label),
        Card("FTT", ftt.value, "percent", "qualidade de primeira passagem", "good", ftt.reason or ftt.state_label),
        Card("Utilização", utilizacao.value, "percent", "carga sobre a capacidade", "info", utilizacao.reason or utilizacao.state_label),
        Card("Produtividade", produtividade.value, "percent", "indicador estendido do Gestor", "info", produtividade.reason or produtividade.state_label),
        Card("MTBF", mtbf, "duration_card", "tempo médio entre falhas", "info", str(confiabilidade.get("reason") or "Sem base de falhas no período")),
        Card("MTTR", mttr, "duration_card", "tempo médio de reparo", "info", str(confiabilidade.get("reason") or "Sem base de reparos no período")),
    ])
    limitacoes = _missing_notice(painel, [
        ("OEE", oee_metric), ("Disponibilidade", disponibilidade), ("Performance", performance),
        ("FTT", ftt), ("Utilização", utilizacao), ("Produtividade", produtividade),
        ("MTBF", metric({"value": mtbf, "availability": confiabilidade.get("availability"), "reason": confiabilidade.get("reason")})),
        ("MTTR", metric({"value": mttr, "availability": confiabilidade.get("availability"), "reason": confiabilidade.get("reason")})),
    ])

    destaques = []
    if capacidade.get("bottleneck_resource"):
        destaques.append((
            "Maior utilização registrada",
            f"{capacidade.get('bottleneck_resource')} — "
            f"{number_text(capacidade.get('bottleneck_utilizacao_percentual'), 1)}%",
        ))
    if linhas_recurso and linhas_recurso[0].get("oee") is not None:
        destaques.append(("Maior OEE por recurso", f"{linhas_recurso[0].get('recurso')} — {number_text(linhas_recurso[0].get('oee'), 1)}%"))
    if _number(confiabilidade.get("falhas")):
        destaques.append(("Falhas de equipamento", f"{number_text(confiabilidade.get('falhas'), 0)} no período"))
    if destaques:
        painel.section("Destaques do período", hint="Sem semáforo: não existe meta gerencial configurada no sistema.")
        painel.facts(destaques)

    painel.section("Gráficos")
    ancora = painel.row
    if tabela_evolucao is not None:
        trend_chart(
            painel.ws, tabela_evolucao,
            title="Evolução do OEE",
            label_key="periodo", value_key="oee",
            anchor=f"A{ancora}", color=BLUE, number_format=FORMATS["percent"],
        )
    elif evolucao.get("reason"):
        painel.row = ancora
        painel.notice(f"Evolução do OEE: {evolucao.get('reason')}", tone="muted")
        ancora = painel.row
    if tabela_comparacao is not None:
        ranking_chart(
            painel.ws, tabela_comparacao,
            title="Disponibilidade, performance, FTT e OEE",
            label_key="indicador", value_key="valor",
            anchor=f"G{ancora}", color=BLUE, number_format=FORMATS["percent"], horizontal=True,
        )
    ancora += 17
    if tabela_recursos is not None and any(linha.get("oee") is not None for linha in linhas_recurso):
        ranking_chart(
            painel.ws, tabela_recursos,
            title="OEE por recurso",
            label_key="recurso", value_key="oee",
            anchor=f"A{ancora}", color=BLUE, number_format=FORMATS["percent"], max_rows=12,
        )
    if tabela_capacidade is not None:
        ranking_chart(
            painel.ws, tabela_capacidade,
            title="Utilização por recurso",
            label_key="codigo", value_key="utilizacao_percentual",
            anchor=f"G{ancora}", color=AMBER, number_format=FORMATS["percent"], max_rows=12,
        )
    painel.row = max(painel.row, ancora + 17)
    _finish(painel)

    # --- Qualidade dos dados ------------------------------------------------
    qualidade_linhas = list(limitacoes)
    if capacidade.get("reason"):
        qualidade_linhas.append({
            "assunto": "Cobertura de calendário produtivo",
            "situacao": (
                f"{number_text(capacidade.get('recursos_com_calendario'), 0)} de "
                f"{number_text(capacidade.get('recursos'), 0)} recursos configurados"
            ),
            "detalhe": str(capacidade.get("reason")),
        })
    if evolucao.get("reason"):
        qualidade_linhas.append({
            "assunto": "Evolução do OEE",
            "situacao": _estado(evolucao.get("availability")),
            "detalhe": str(evolucao.get("reason")),
        })
    qualidade_linhas.append({
        "assunto": "Metas gerenciais",
        "situacao": "Não configuradas",
        "detalhe": (
            "Nenhuma fonte oficial de metas está configurada; os indicadores são apresentados sem "
            "classificação de bom, ruim ou crítico."
        ),
    })
    qualidade_linhas.append({
        "assunto": "OEE por setor",
        "situacao": "Não publicado",
        "detalhe": "O backend publica OEE por recurso e para a fábrica; nenhum OEE setorial é derivado aqui.",
    })
    _quality_sheet(workbook, filters, generated_at, qualidade_linhas)
    technical_sheet(workbook, payload, filters, generated_at=generated_at)
    _hide(workbook, SERIES_SHEET)


# --------------------------------------------------------------------------
# 5. Dados Analíticos
# --------------------------------------------------------------------------

ISSUE_COLUMNS = (
    Column("severidade", "Severidade", "text", 14, tones=SEVERITY_TONES),
    Column("tipo", "Tipo", "text", 40),
    Column("mensagem", "Mensagem", "text", 58, wrap=True),
    Column("op", "OP", "text", 16),
    Column("operacao", "Operação", "text", 12),
    Column("recurso", "Recurso", "text", 16),
    Column("setor", "Setor", "text", 16),
    Column("inicio", "Início", "datetime"),
    Column("fim", "Fim", "datetime"),
    Column("segundos", "Duração", "duration"),
    Column("codigo", "Código técnico", "text", 38),
)

NESTING_COLUMNS = (
    Column("nesting", "Nesting", "text", 16),
    Column("maquina", "Recurso", "text", 16),
    Column("programa", "Programa", "text", 20),
    Column("real_segundos", "Tempo real", "duration"),
    Column("real_periodo_segundos", "Tempo no período", "duration"),
)


def _issue_rows(issues: Sequence[dict[str, Any]]) -> list[dict[str, Any]]:
    return [
        {
            "severidade": SEVERITY_LABELS.get(str(item.get("severity")), str(item.get("severity") or "")),
            "tipo": ISSUE_LABELS.get(str(item.get("type")), str(item.get("type") or "")),
            "mensagem": item.get("message"),
            "op": item.get("op"),
            "operacao": item.get("operation"),
            "recurso": item.get("resource"),
            "setor": item.get("sector"),
            "inicio": item.get("start"),
            "fim": item.get("end"),
            "segundos": item.get("seconds"),
            "codigo": item.get("type"),
        }
        for item in issues
    ]


def build_dados_analiticos(workbook, payload: dict[str, Any], filters, generated_at: datetime) -> None:
    ordens = _dict(payload.get("orders"))
    auditoria = _dict(payload.get("audit"))
    nestings = _dict(payload.get("nestings"))
    diagnosticos = _dict(auditoria.get("diagnostics"))
    por_severidade = _dict(auditoria.get("by_severity"))
    itens_ordens = _rows(ordens.get("items"))
    inconsistencias = _rows(auditoria.get("issues"))
    lacunas = _rows(diagnosticos.get("calendar_gaps"))

    painel = ReportSheet(workbook.create_sheet("Resumo da Base"), width=12)
    aba_ordens = ReportSheet(workbook.create_sheet("Ordens"), width=15)
    aba_inconsistencias = ReportSheet(workbook.create_sheet("Inconsistências"), width=11)
    aba_diagnosticos = ReportSheet(workbook.create_sheet("Diagnósticos"), width=8)

    # --- Ordens -------------------------------------------------------------
    aba_ordens.header("Ordens analisadas", filters, generated_at=generated_at)
    if itens_ordens:
        aba_ordens.section("Base operacional do período")
        aba_ordens.table(ORDER_COLUMNS, itens_ordens)
    else:
        aba_ordens.empty("Nenhuma ordem com apontamento no período filtrado.")
    _finish(aba_ordens)

    # --- Inconsistências ----------------------------------------------------
    aba_inconsistencias.header("Inconsistências encontradas", filters, generated_at=generated_at)
    if inconsistencias:
        aba_inconsistencias.section(
            "Somente inconsistências reais",
            hint="Recurso ocioso dentro do turno não entra aqui: ele é diagnóstico, não inconsistência.",
        )
        tabela_incon = aba_inconsistencias.table(ISSUE_COLUMNS, _issue_rows(inconsistencias))
        _hide_column(aba_inconsistencias, tabela_incon, "codigo")
    else:
        aba_inconsistencias.empty("Nenhuma inconsistência encontrada no período filtrado.")
    _finish(aba_inconsistencias)

    # --- Diagnósticos -------------------------------------------------------
    aba_diagnosticos.header("Diagnósticos de cobertura", filters, generated_at=generated_at)
    recursos_afetados = sorted({str(item.get("resource") or "") for item in lacunas if item.get("resource")})
    setores_afetados = sorted({str(item.get("sector") or "") for item in lacunas if item.get("sector")})
    aba_diagnosticos.section("Resumo")
    aba_diagnosticos.table(
        (
            Column("indicador", "Diagnóstico", "text", 36),
            Column("intervalos", "Intervalos", "int"),
            Column("tempo", "Tempo sem estado", "duration"),
        ),
        [{
            "indicador": "Tempo de turno sem estado físico registrado",
            "intervalos": diagnosticos.get("calendar_gap_count"),
            "tempo": diagnosticos.get("calendar_gap_seconds"),
        }],
        autofilter=False,
    )
    aba_diagnosticos.facts([
        ("Recursos afetados", f"{number_text(len(recursos_afetados), 0)} recurso(s)"),
        ("Setores afetados", f"{number_text(len(setores_afetados), 0)} setor(es)"),
    ])
    if diagnosticos.get("reason"):
        aba_diagnosticos.notice(str(diagnosticos.get("reason")), tone="muted")
    if recursos_afetados:
        por_recurso: dict[str, dict[str, Any]] = {}
        for item in lacunas:
            chave = str(item.get("resource") or "Não informado")
            atual = por_recurso.setdefault(chave, {"recurso": chave, "setor": item.get("sector"), "intervalos": 0, "segundos": 0.0})
            atual["intervalos"] += 1
            atual["segundos"] += _number(item.get("seconds")) or 0.0
        aba_diagnosticos.section("Recursos com tempo de turno sem estado registrado")
        aba_diagnosticos.table(
            (
                Column("recurso", "Recurso", "text", 20),
                Column("setor", "Setor", "text", 18),
                Column("intervalos", "Intervalos", "int"),
                Column("segundos", "Tempo sem estado", "duration"),
            ),
            _sorted(list(por_recurso.values()), "segundos"),
            freeze=False,
        )
    _finish(aba_diagnosticos)

    if lacunas:
        detalhe = ReportSheet(workbook.create_sheet("Diagnósticos (detalhe)"), width=11)
        detalhe.header("Intervalos sem estado físico", filters, generated_at=generated_at)
        detalhe.section("Detalhe completo", hint="Aba oculta por padrão: volume alto e uso de conferência.")
        tabela_detalhe = detalhe.table(ISSUE_COLUMNS, _issue_rows(lacunas))
        _hide_column(detalhe, tabela_detalhe, "codigo")
        _finish(detalhe)
        detalhe.ws.sheet_state = "hidden"

    # --- Nestings -----------------------------------------------------------
    itens_nesting = _rows(nestings.get("items"))
    if itens_nesting:
        aba_nestings = ReportSheet(workbook.create_sheet("Nestings"), width=5)
        aba_nestings.header("Nestings do Corte", filters, generated_at=generated_at)
        aba_nestings.section(
            "Tempo por nesting",
            hint="“Tempo real” é o total do nesting; “tempo no período” é a parcela dentro do filtro.",
        )
        aba_nestings.table(NESTING_COLUMNS, itens_nesting)
        _finish(aba_nestings)

    # --- Resumo da base -----------------------------------------------------
    painel.header("Dados Analíticos", filters, generated_at=generated_at)
    confiabilidade_base = metric({
        "value": auditoria.get("reliability_percentage"),
        "availability": "dados_insuficientes" if auditoria.get("reliability_percentage") is None else "disponivel",
        "reason": auditoria.get("reliability_reason"),
    })
    painel.section("Base analisada")
    painel.kpis([
        Card("Ordens analisadas", ordens.get("count"), "int", "apontamentos no período", "neutral"),
        Card("Estados físicos", auditoria.get("resource_states_analyzed"), "int", "registros de estado do recurso", "neutral"),
        Card("Inconsistências", auditoria.get("count"), "int", "ocorrências reais", "bad"),
        Card("Nestings", nestings.get("count"), "int", "planos de corte", "neutral"),
        Card("Intervalos sem estado", diagnosticos.get("calendar_gap_count"), "int", "diagnóstico de cobertura", "warn"),
        Card("Confiabilidade da base", confiabilidade_base.value, "percent", "", "info", confiabilidade_base.reason or confiabilidade_base.state_label),
    ])
    _missing_notice(painel, [("Confiabilidade da base", confiabilidade_base)])

    if any(_number(valor) for valor in por_severidade.values()):
        painel.section("Inconsistências por severidade")
        painel.table(
            (Column("severidade", "Severidade", "text", 20), Column("quantidade", "Ocorrências", "int")),
            [
                {"severidade": SEVERITY_LABELS.get(chave, chave), "quantidade": valor}
                for chave, valor in por_severidade.items()
            ],
            autofilter=False,
            freeze=False,
        )
    if inconsistencias:
        tipos: dict[str, int] = {}
        for item in inconsistencias:
            chave = ISSUE_LABELS.get(str(item.get("type")), str(item.get("type") or ""))
            tipos[chave] = tipos.get(chave, 0) + 1
        painel.section("Principais tipos encontrados")
        painel.facts([
            (tipo, f"{number_text(quantidade, 0)} ocorrência(s)")
            for tipo, quantidade in sorted(tipos.items(), key=lambda item: -item[1])[:6]
        ])
    else:
        painel.empty("Nenhuma inconsistência encontrada no período filtrado.")
    _finish(painel)

    qualidade_linhas = []
    if auditoria.get("reliability_reason"):
        qualidade_linhas.append({
            "assunto": "Percentual de confiabilidade da base",
            "situacao": "Não publicado",
            "detalhe": str(auditoria.get("reliability_reason")),
        })
    if diagnosticos.get("reason"):
        qualidade_linhas.append({
            "assunto": "Intervalos sem estado físico",
            "situacao": f"{number_text(diagnosticos.get('calendar_gap_count'), 0)} intervalos",
            "detalhe": str(diagnosticos.get("reason")),
        })
    if ordens.get("availability"):
        qualidade_linhas.append({
            "assunto": "Cobertura das ordens",
            "situacao": _estado(ordens.get("availability")),
            "detalhe": f"Fonte física declarada pelo backend: {_fonte(auditoria.get('physical_state_source'))}.",
        })
    _quality_sheet(workbook, filters, generated_at, qualidade_linhas)
    technical_sheet(workbook, payload, filters, generated_at=generated_at)


BUILDERS = {
    "gerencial": build_gerencial,
    "producao": build_producao,
    "perdas": build_perdas,
    "indicadores": build_indicadores,
    "dados_analiticos": build_dados_analiticos,
}


__all__ = ["BUILDERS", "FOOTER", "ISSUE_LABELS", "OVERVIEW_SHEET", "QUALITY_SHEET", "SEVERITY_LABELS"]
