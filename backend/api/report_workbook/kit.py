"""Identidade visual e primitivas reutilizáveis das exportações XLSX.

Este módulo é a única fonte de estilo dos relatórios. Ele não conhece
indicadores industriais: recebe valores já calculados pelo backend canônico e
apenas decide como apresentá-los. Zero e dado ausente são tratados como coisas
diferentes em todas as primitivas.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from functools import lru_cache
import json
import re
from typing import Any, Sequence

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# --- Paleta -----------------------------------------------------------------
# Restrita e com significado. Nenhuma cor existe para decorar um KPI.
NAVY = "10243F"      # institucional: títulos e cabeçalhos
BLUE = "2A6099"      # informação neutra
GREEN = "1B7F4B"     # produção conforme / resultado bom
AMBER = "B5751A"     # atenção
RED = "B3261E"       # perda, refugo, parada, situação crítica
TEXT = "1F2933"      # texto padrão
MUTED = "6B7A8F"     # texto auxiliar
SURFACE = "F4F6F9"   # fundo de tabelas e áreas auxiliares
LINE = "DCE3EB"      # linhas discretas
WHITE = "FFFFFF"

FONT = "Calibri"
MISSING = "—"

TONE_COLORS = {
    "neutral": NAVY,
    "info": BLUE,
    "good": GREEN,
    "warn": AMBER,
    "bad": RED,
    "muted": MUTED,
}

# --- Formatos ---------------------------------------------------------------
# Convenções únicas para todo o pacote: percentual executivo com uma casa,
# duração em [h]:mm:ss nas tabelas e em "0h 00min" nos cartões executivos.
FORMATS = {
    "text": "@",
    "int": "#,##0",
    "number": "#,##0.0",
    # Uma casa decimal na leitura executiva; abaixo de 0,1% o valor ganha
    # precisão em vez de virar "0,0%" e se confundir com zero.
    "percent": "[<0.001]0.000%;0.0%",
    "percent_exact": "0.00%",
    "duration": "[h]:mm:ss",
    "duration_card": '[h]"h "mm"min"',
    "date": "dd/mm/yyyy",
    "datetime": "dd/mm/yyyy hh:mm",
}

DURATION_KINDS = frozenset({"duration", "duration_card"})

# Eixo de gráfico não usa o formato condicional de percentual: em um eixo,
# "0,000%" no zero só polui a leitura.
AXIS_PERCENT = "0.0%"


def _chart_format(number_format: str | None) -> str | None:
    return AXIS_PERCENT if number_format == FORMATS["percent"] else number_format


_ISO = re.compile(r"\d{4}-\d{2}-\d{2}(?:[T ]\d{2}:\d{2}:\d{2}(?:\.\d+)?(?:Z|[+-]\d{2}:?\d{2})?)?")
_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")

THIN = Side(style="thin", color=LINE)


# Estilos memoizados: uma tabela de auditoria pode ter milhares de células e
# criar um objeto de fonte por célula dominava o tempo de geração.
@lru_cache(maxsize=256)
def font(*, size: int = 10, bold: bool = False, color: str = TEXT, italic: bool = False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


@lru_cache(maxsize=256)
def align(*, horizontal: str | None = None, vertical: str = "center", wrap: bool = False) -> Alignment:
    return Alignment(horizontal=horizontal, vertical=vertical, wrap_text=wrap)


@lru_cache(maxsize=64)
def fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


@lru_cache(maxsize=64)
def bottom_border(color: str = LINE, style: str = "thin") -> Border:
    return Border(bottom=Side(style=style, color=color))


def safe_text(value: Any) -> str:
    """Neutraliza fórmula injetada por dado de origem."""

    text = str(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def parse_datetime(value: Any):
    if isinstance(value, (datetime, date)):
        return value
    text = str(value or "").strip()
    if not text or not _ISO.fullmatch(text):
        return None
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except ValueError:
        return None


def coerce(value: Any, kind: str):
    """Converte o valor canônico para o tipo nativo do Excel.

    Retorna ``None`` quando o dado está ausente — nunca zero.
    """

    if value is None:
        return None
    if kind in {"percent", "percent_exact"}:
        return float(value) / 100.0 if isinstance(value, (int, float)) else None
    if kind in {"duration", "duration_card"}:
        return float(value) / 86400.0 if isinstance(value, (int, float)) else None
    if kind == "int":
        return int(value) if isinstance(value, (int, float)) else None
    if kind == "number":
        return float(value) if isinstance(value, (int, float)) else None
    if kind in {"date", "datetime"}:
        return parse_datetime(value)
    if isinstance(value, bool):
        return "Sim" if value else "Não"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, (dict, list, tuple)):
        # Estrutura só chega aqui na aba técnica/genérica; JSON compacto é mais
        # legível e auditável do que o repr do Python.
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"), default=str)
    parsed = parse_datetime(value)
    return parsed if parsed is not None else safe_text(value)


def _card_reason(card: "Card", card_width: int) -> str:
    """Motivo curto cabe no cartão; o texto completo fica no aviso e na aba de qualidade."""

    texto = str(card.reason or "").strip()
    return texto if texto and len(texto) <= card_width * 11 else "Dados insuficientes"


def duration_text(seconds: Any) -> str:
    """Duração legível para frases (os números ficam em células de duração)."""

    if not isinstance(seconds, (int, float)):
        return MISSING
    total = int(round(float(seconds)))
    horas, resto = divmod(abs(total), 3600)
    minutos, segundos = divmod(resto, 60)
    # Sinal tipográfico (U+2212): o Excel não exibe duração negativa como número
    # e o hífen inicial ainda dispararia a proteção contra fórmula injetada.
    sinal = "−" if total < 0 else ""
    if horas:
        return f"{sinal}{horas}h {minutos:02d}min"
    if minutos:
        return f"{sinal}{minutos}min {segundos:02d}s"
    return f"{sinal}{segundos}s"


def number_text(value: Any, casas: int = 1) -> str:
    if not isinstance(value, (int, float)):
        return MISSING
    if float(value).is_integer() and casas == 0:
        return f"{int(value):,}".replace(",", ".")
    # Valor pequeno e não nulo ganha precisão em vez de virar "0,0".
    if value and abs(float(value)) < 0.1 and casas:
        casas = 3
    inteiro, _, decimal = f"{float(value):,.{casas}f}".partition(".")
    inteiro = inteiro.replace(",", ".")
    return f"{inteiro},{decimal}" if decimal else inteiro


def metric(payload: Any) -> "Metric":
    """Lê o contrato canônico ``{value, availability, unit, reason}``."""

    data = payload if isinstance(payload, dict) else {}
    return Metric(
        value=data.get("value"),
        availability=str(data.get("availability") or ("disponivel" if data.get("value") is not None else "dados_insuficientes")),
        unit=data.get("unit"),
        reason=data.get("reason"),
    )


AVAILABILITY_LABELS = {
    "disponivel": "Disponível",
    "parcial": "Parcial",
    "sem_registros": "Sem registros no período",
    "dados_insuficientes": "Dados insuficientes",
    "nao_configurado": "Não configurado",
}


@dataclass(frozen=True)
class Metric:
    value: Any
    availability: str = "disponivel"
    unit: str | None = None
    reason: str | None = None

    @property
    def missing(self) -> bool:
        return self.value is None

    @property
    def state_label(self) -> str:
        return AVAILABILITY_LABELS.get(self.availability, "Dados insuficientes")


@dataclass(frozen=True)
class Card:
    label: str
    value: Any
    kind: str = "int"
    context: str = ""
    tone: str = "neutral"
    reason: str | None = None


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    kind: str = "text"
    width: int | None = None
    wrap: bool = False
    # Cor por valor (severidade, status). Discreta e sem julgar desempenho.
    tones: tuple[tuple[str, str], ...] = ()

    def tone_for(self, value: Any) -> str | None:
        if not self.tones:
            return None
        texto = str(value or "").strip().casefold()
        for chave, cor in self.tones:
            if chave.casefold() == texto:
                return cor
        return None


def write_value(cell, raw: Any, kind: str, *, size: int = 10, bold: bool = False, color: str = TEXT, wrap: bool = False) -> bool:
    """Escreve um valor canônico respeitando ausência e duração negativa.

    O Excel não consegue exibir um tempo negativo como número de duração
    (mostraria ``########``); nesses casos o valor vira texto legível com o
    sinal explícito. Retorna ``False`` quando o dado está ausente.
    """

    if raw is None:
        cell.value = MISSING
        cell.number_format = FORMATS["text"]
        cell.font = font(size=size, bold=bold, color=MUTED)
        cell.alignment = align(horizontal="right" if kind != "text" else "left")
        return False
    if kind in DURATION_KINDS and isinstance(raw, (int, float)) and raw < 0:
        cell.value = duration_text(raw)
        cell.number_format = FORMATS["text"]
        cell.font = font(size=size, bold=bold, color=color)
        return True
    valor = coerce(raw, kind)
    if valor is None:
        cell.value = MISSING
        cell.number_format = FORMATS["text"]
        cell.font = font(size=size, bold=bold, color=MUTED)
        return False
    cell.value = valor
    if kind == "text" and isinstance(valor, (datetime, date)):
        # Data reconhecida dentro de um campo textual continua sendo data.
        cell.number_format = FORMATS["datetime"] if isinstance(valor, datetime) else FORMATS["date"]
    else:
        cell.number_format = FORMATS.get(kind, FORMATS["text"])
    cell.font = font(size=size, bold=bold, color=color)
    if wrap:
        cell.alignment = align(vertical="top", wrap=True)
    return True


@dataclass
class TableRange:
    """Coordenadas de uma tabela escrita, para ancorar gráficos."""

    sheet_title: str
    header_row: int
    first_row: int
    last_row: int
    first_col: int
    columns: list[Column] = field(default_factory=list)

    def column_index(self, key: str) -> int:
        for offset, column in enumerate(self.columns):
            if column.key == key:
                return self.first_col + offset
        raise KeyError(key)

    @property
    def rows(self) -> int:
        return max(0, self.last_row - self.first_row + 1)


def sheet_title(name: str, used: set[str]) -> str:
    base = _INVALID_SHEET_CHARS.sub(" ", str(name or "Dados")).strip()[:31] or "Dados"
    candidate = base
    index = 2
    while candidate.casefold() in used:
        suffix = f" {index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def filters_label(filters) -> str:
    """Resume os filtros aplicados em linguagem de chão de fábrica."""

    data = filters.to_dict() if hasattr(filters, "to_dict") else dict(filters or {})
    rotulos = (
        ("setor", "Setor", "Todos os setores"),
        ("recurso", "Recurso", "Todos os recursos"),
        ("turno", "Turno", None),
        ("op", "OP", None),
        ("operacao", "Operação", None),
        ("produto", "Produto", None),
        ("operador", "Operador", None),
    )
    partes = []
    for key, label, default in rotulos:
        value = data.get(key)
        if value not in (None, "", []):
            partes.append(f"{label}: {value}")
        elif default:
            partes.append(default)
    return " • ".join(partes)


def period_label(filters) -> str:
    data = filters.to_dict() if hasattr(filters, "to_dict") else dict(filters or {})
    inicio = parse_datetime(data.get("inicio"))
    fim = parse_datetime(data.get("fim"))
    if inicio is None or fim is None:
        return "Período não informado"
    return f"{inicio:%d/%m/%Y %H:%M} — {fim:%d/%m/%Y %H:%M}"


class ReportSheet:
    """Escreve uma aba seguindo o padrão visual único das exportações."""

    HEADER_HEIGHT = 5

    def __init__(self, worksheet, *, width: int = 12, default_column_width: float = 12.0):
        self.ws = worksheet
        self.width = width
        self.row = 1
        self.ws.sheet_view.showGridLines = False
        for index in range(1, width + 1):
            self.ws.column_dimensions[get_column_letter(index)].width = default_column_width

    # -- blocos estruturais --------------------------------------------------
    def header(self, title: str, filters, *, generated_at: datetime, subtitle: str | None = None) -> None:
        ws = self.ws
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=self.width)
        marca = ws.cell(1, 1, "GESTOR DE PEÇAS")
        marca.font = Font(name=FONT, size=9, bold=True, color=MUTED)

        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=self.width)
        titulo = ws.cell(2, 1, title)
        titulo.font = Font(name=FONT, size=18, bold=True, color=NAVY)
        titulo.alignment = Alignment(vertical="center")
        ws.row_dimensions[2].height = 24

        linhas = [f"Período: {period_label(filters)}"]
        filtros = filters_label(filters)
        if filtros:
            linhas.append(f"Filtros: {filtros}")
        linhas.append(f"Gerado em: {generated_at:%d/%m/%Y %H:%M}")
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=self.width)
        contexto = ws.cell(3, 1, "   •   ".join(linhas))
        contexto.font = Font(name=FONT, size=9, color=MUTED)

        if subtitle:
            ws.merge_cells(start_row=4, start_column=1, end_row=4, end_column=self.width)
            legenda = ws.cell(4, 1, subtitle)
            legenda.font = Font(name=FONT, size=9, italic=True, color=MUTED)

        regua = 5
        ws.row_dimensions[regua].height = 4
        for index in range(1, self.width + 1):
            ws.cell(regua, index).fill = PatternFill("solid", fgColor=NAVY)
        self.row = regua + 2

    def section(self, title: str, hint: str | None = None) -> None:
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = ws.cell(self.row, 1, title.upper())
        cell.font = Font(name=FONT, size=11, bold=True, color=NAVY)
        cell.alignment = Alignment(vertical="center")
        cell.border = Border(bottom=Side(style="thin", color=NAVY))
        ws.row_dimensions[self.row].height = 20
        self.row += 1
        if hint:
            ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
            nota = ws.cell(self.row, 1, hint)
            nota.font = Font(name=FONT, size=9, color=MUTED)
            self.row += 1
        self.row += 1

    def spacer(self, rows: int = 1) -> None:
        self.row += rows

    def kpis(self, cards: Sequence[Card], *, per_row: int = 4, card_width: int = 3) -> None:
        """Cartões executivos com espaço suficiente para o texto completo."""

        ws = self.ws
        for start in range(0, len(cards), per_row):
            bloco = cards[start:start + per_row]
            label_row = self.row
            value_row = self.row + 1
            context_row = self.row + 3
            for index, card in enumerate(bloco):
                first = 1 + index * card_width
                last = first + card_width - 1
                if last > self.width:
                    break
                ws.merge_cells(start_row=label_row, start_column=first, end_row=label_row, end_column=last)
                ws.merge_cells(start_row=value_row, start_column=first, end_row=value_row + 1, end_column=last)
                ws.merge_cells(start_row=context_row, start_column=first, end_row=context_row, end_column=last)

                rotulo = ws.cell(label_row, first, card.label.upper())
                rotulo.font = Font(name=FONT, size=9, bold=True, color=MUTED)
                rotulo.alignment = Alignment(horizontal="left", vertical="center")

                valor = ws.cell(value_row, first)
                presente = write_value(
                    valor, card.value, card.kind,
                    size=20, bold=True, color=TONE_COLORS.get(card.tone, NAVY),
                )
                valor.alignment = Alignment(horizontal="left", vertical="center")

                contexto = ws.cell(context_row, first)
                contexto.value = card.context if presente else _card_reason(card, card_width)
                contexto.font = Font(name=FONT, size=9, color=MUTED)
                contexto.alignment = Alignment(horizontal="left", vertical="top")

                for column in range(first, last + 1):
                    for linha in (label_row, value_row, value_row + 1, context_row):
                        ws.cell(linha, column).fill = PatternFill("solid", fgColor=SURFACE)
                    ws.cell(context_row, column).border = Border(bottom=Side(
                        style="medium",
                        color=TONE_COLORS.get(card.tone, NAVY) if presente else LINE,
                    ))
            ws.row_dimensions[label_row].height = 16
            ws.row_dimensions[value_row].height = 20
            ws.row_dimensions[value_row + 1].height = 12
            ws.row_dimensions[context_row].height = 15
            self.row = context_row + 2

    def notice(self, text: str, *, tone: str = "info") -> None:
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = ws.cell(self.row, 1, text)
        cell.font = Font(name=FONT, size=9, color=TONE_COLORS.get(tone, BLUE))
        cell.alignment = Alignment(vertical="center", wrap_text=False)
        cell.fill = PatternFill("solid", fgColor=SURFACE)
        cell.border = Border(left=Side(style="medium", color=TONE_COLORS.get(tone, BLUE)))
        for column in range(2, self.width + 1):
            ws.cell(self.row, column).fill = PatternFill("solid", fgColor=SURFACE)
        ws.row_dimensions[self.row].height = 18
        self.row += 2

    def empty(self, message: str) -> None:
        self.notice(message, tone="muted")

    def facts(self, linhas: Sequence[tuple[str, str]]) -> None:
        """Destaques objetivos: rótulo à esquerda, fato à direita."""

        ws = self.ws
        for rotulo, fato in linhas:
            ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=3)
            ws.merge_cells(start_row=self.row, start_column=4, end_row=self.row, end_column=self.width)
            chave = ws.cell(self.row, 1, rotulo)
            chave.font = Font(name=FONT, size=10, color=MUTED)
            valor = ws.cell(self.row, 4, safe_text(fato))
            valor.font = Font(name=FONT, size=10, bold=True, color=TEXT)
            for column in range(1, self.width + 1):
                ws.cell(self.row, column).border = Border(bottom=THIN)
            ws.row_dimensions[self.row].height = 16
            self.row += 1
        self.row += 1

    def table(
        self,
        columns: Sequence[Column],
        rows: Sequence[dict[str, Any]],
        *,
        start_col: int = 1,
        totals: dict[str, Any] | None = None,
        autofilter: bool = True,
        freeze: bool = True,
        zebra: bool = True,
    ) -> TableRange:
        ws = self.ws
        header_row = self.row
        for offset, column in enumerate(columns):
            cell = ws.cell(header_row, start_col + offset, column.label)
            cell.fill = PatternFill("solid", fgColor=NAVY)
            cell.font = Font(name=FONT, size=10, bold=True, color=WHITE)
            # Cabeçalho numérico centralizado: alinhado à direita ele fica
            # escondido atrás do botão de filtro do Excel.
            cell.alignment = align(
                horizontal="center" if column.kind not in {"text", "date", "datetime"} else "left",
                wrap=True,
            )
        ws.row_dimensions[header_row].height = 28

        first_row = header_row + 1
        for index, row in enumerate(rows):
            linha = first_row + index
            fundo = SURFACE if zebra and index % 2 else None
            for offset, column in enumerate(columns):
                cell = ws.cell(linha, start_col + offset)
                write_value(
                    cell, row.get(column.key), column.kind,
                    wrap=column.wrap,
                    color=column.tone_for(row.get(column.key)) or TEXT,
                )
                cell.alignment = align(
                    horizontal="left" if column.kind in {"text", "date", "datetime"} else "right",
                    vertical="top" if column.wrap else "center",
                    wrap=column.wrap,
                )
                if fundo:
                    cell.fill = fill(fundo)
                cell.border = bottom_border()
        last_row = first_row + len(rows) - 1

        if totals:
            linha = last_row + 1
            for offset, column in enumerate(columns):
                cell = ws.cell(linha, start_col + offset)
                if offset == 0:
                    cell.value = totals.get(column.key, "Total")
                    cell.number_format = FORMATS["text"]
                elif column.key in totals:
                    write_value(cell, totals.get(column.key), column.kind, bold=True, color=NAVY)
                cell.font = Font(name=FONT, size=10, bold=True, color=NAVY)
                cell.alignment = Alignment(
                    horizontal="left" if column.kind in {"text", "date", "datetime"} else "right",
                    vertical="center",
                )
                cell.border = Border(top=Side(style="medium", color=NAVY))
            self.row = linha + 2
        else:
            self.row = last_row + 2

        for offset, column in enumerate(columns):
            letra = get_column_letter(start_col + offset)
            largura = column.width or self._auto_width(column, rows)
            atual = ws.column_dimensions[letra].width or 0
            ws.column_dimensions[letra].width = max(atual, largura)

        if rows and autofilter and start_col == 1:
            ws.auto_filter.ref = (
                f"{get_column_letter(start_col)}{header_row}:"
                f"{get_column_letter(start_col + len(columns) - 1)}{last_row}"
            )
        if freeze and start_col == 1:
            ws.freeze_panes = ws.cell(first_row, 1)
            ws.print_title_rows = f"{header_row}:{header_row}"

        return TableRange(
            sheet_title=ws.title,
            header_row=header_row,
            first_row=first_row,
            last_row=last_row,
            first_col=start_col,
            columns=list(columns),
        )

    @staticmethod
    def _auto_width(column: Column, rows: Sequence[dict[str, Any]]) -> float:
        if column.kind in {"percent", "percent_exact"}:
            base = 12
        elif column.kind == "duration":
            base = 14
        elif column.kind == "datetime":
            base = 17
        elif column.kind == "date":
            base = 12
        elif column.kind in {"int", "number"}:
            base = 13
        else:
            tamanhos = [len(str(row.get(column.key) or "")) for row in rows[:200]]
            base = min(46, max(14, (max(tamanhos) if tamanhos else 0) + 2))
        return max(base, min(30, len(column.label) + 4))

    def footer(self, text: str) -> None:
        ws = self.ws
        ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=self.width)
        cell = ws.cell(self.row, 1, text)
        cell.font = Font(name=FONT, size=8, color=MUTED)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[self.row].height = 14
        self.row += 2

    def print_setup(self, *, landscape: bool = True, fit_width: int = 1) -> None:
        ws = self.ws
        ws.page_setup.orientation = "landscape" if landscape else "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = fit_width
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_options.horizontalCentered = True
        ws.page_margins.left = 0.4
        ws.page_margins.right = 0.4
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5
        ws.oddFooter.left.text = "Gestor de Peças"
        ws.oddFooter.left.size = 8
        ws.oddFooter.left.color = MUTED
        ws.oddFooter.right.text = "Página &P de &N"
        ws.oddFooter.right.size = 8
        ws.oddFooter.right.color = MUTED


# --- Gráficos ---------------------------------------------------------------
# Todo gráfico responde a uma pergunta e só existe quando há dado suficiente.

def _value_labels(number_format: str) -> DataLabelList:
    """Só o valor no rótulo: sem repetir série e categoria dentro do gráfico."""

    labels = DataLabelList()
    labels.showVal = True
    labels.showSerName = False
    labels.showCatName = False
    labels.showLegendKey = False
    labels.showPercent = False
    labels.showBubbleSize = False
    labels.numFmt = number_format
    return labels


def _series_fill(series, color: str) -> None:
    series.graphicalProperties = GraphicalProperties(solidFill=color)
    series.graphicalProperties.line = LineProperties(noFill=True)


def _base_style(chart, title: str, *, horizontal: bool = False) -> None:
    chart.title = title
    # Sem isto o Excel desenha o título por cima da área de plotagem.
    chart.title.overlay = False
    chart.style = None
    chart.roundedCorners = False
    chart.visible_cells_only = False
    if chart.x_axis is not None:
        chart.x_axis.delete = False
        chart.x_axis.majorGridlines = None
        # Categoria à esquerda na barra horizontal; embaixo nos demais tipos.
        chart.x_axis.axPos = "l" if horizontal else "b"
    if chart.y_axis is not None:
        chart.y_axis.delete = False
        chart.y_axis.axPos = "b" if horizontal else "l"
    if chart.legend is not None:
        # Legenda embaixo: à direita ela rouba largura e corta o texto.
        chart.legend.position = "b"
        chart.legend.overlay = False


def ranking_chart(
    target,
    table: TableRange,
    *,
    title: str,
    label_key: str,
    value_key: str,
    anchor: str,
    color: str = BLUE,
    max_rows: int = 12,
    number_format: str | None = None,
    horizontal: bool = True,
    width: float = 13.8,
    height: float = 8.0,
) -> bool:
    """Ranking: barras horizontais ordenadas da maior para a menor."""

    if table.rows <= 0:
        return False
    source = target.parent[table.sheet_title]
    last = min(table.last_row, table.first_row + max_rows - 1)
    chart = BarChart()
    chart.type = "bar" if horizontal else "col"
    chart.grouping = "clustered"
    _base_style(chart, title, horizontal=horizontal)
    chart.legend = None
    data = Reference(source, min_col=table.column_index(value_key), min_row=table.header_row, max_row=last)
    cats = Reference(source, min_col=table.column_index(label_key), min_row=table.first_row, max_row=last)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    _series_fill(chart.series[0], color)
    if horizontal:
        # Sem isto o Excel desenha a maior barra embaixo e o ranking fica invertido.
        chart.x_axis.scaling.orientation = "maxMin"
        chart.y_axis.crosses = "max"
    number_format = _chart_format(number_format)
    if number_format:
        chart.y_axis.numFmt = number_format
        chart.dataLabels = _value_labels(number_format)
    chart.width = width
    chart.height = height
    chart.gapWidth = 60
    target.add_chart(chart, anchor)
    return True


def composition_chart(
    target,
    table: TableRange,
    *,
    title: str,
    label_key: str,
    value_key: str,
    anchor: str,
    color: str = BLUE,
    number_format: str | None = None,
    width: float = 13.8,
    height: float = 8.0,
) -> bool:
    """Composição: barras horizontais de participação por categoria."""

    return ranking_chart(
        target,
        table,
        title=title,
        label_key=label_key,
        value_key=value_key,
        anchor=anchor,
        color=color,
        number_format=number_format,
        horizontal=True,
        width=width,
        height=height,
        max_rows=20,
    )


def grouped_chart(
    target,
    table: TableRange,
    *,
    title: str,
    label_key: str,
    value_keys: Sequence[str],
    anchor: str,
    colors: Sequence[str] = (GREEN, RED, AMBER),
    stacked: bool = False,
    number_format: str | None = None,
    max_rows: int = 15,
    width: float = 13.8,
    height: float = 8.0,
) -> bool:
    """Comparação de séries por categoria (lado a lado ou empilhada)."""

    if table.rows <= 0 or not value_keys:
        return False
    source = target.parent[table.sheet_title]
    last = min(table.last_row, table.first_row + max_rows - 1)
    chart = BarChart()
    chart.type = "col"
    chart.grouping = "stacked" if stacked else "clustered"
    if stacked:
        chart.overlap = 100
    _base_style(chart, title)
    for index, key in enumerate(value_keys):
        data = Reference(source, min_col=table.column_index(key), min_row=table.header_row, max_row=last)
        chart.add_data(data, titles_from_data=True)
    chart.set_categories(
        Reference(source, min_col=table.column_index(label_key), min_row=table.first_row, max_row=last)
    )
    for index, series in enumerate(chart.series):
        _series_fill(series, colors[index % len(colors)])
    number_format = _chart_format(number_format)
    if number_format:
        chart.y_axis.numFmt = number_format
    chart.width = width
    chart.height = height
    chart.gapWidth = 80
    target.add_chart(chart, anchor)
    return True


def trend_chart(
    target,
    table: TableRange,
    *,
    title: str,
    label_key: str,
    value_key: str,
    anchor: str,
    color: str = BLUE,
    number_format: str | None = None,
    width: float = 13.8,
    height: float = 8.0,
) -> bool:
    """Evolução temporal — exige pelo menos dois pontos para existir."""

    if table.rows < 2:
        return False
    source = target.parent[table.sheet_title]
    chart = LineChart()
    _base_style(chart, title)
    chart.legend = None
    data = Reference(source, min_col=table.column_index(value_key), min_row=table.header_row, max_row=table.last_row)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(
        Reference(source, min_col=table.column_index(label_key), min_row=table.first_row, max_row=table.last_row)
    )
    series = chart.series[0]
    series.graphicalProperties = GraphicalProperties()
    series.graphicalProperties.line = LineProperties(solidFill=color, w=22000)
    series.smooth = False
    number_format = _chart_format(number_format)
    if number_format:
        chart.y_axis.numFmt = number_format
    chart.width = width
    chart.height = height
    target.add_chart(chart, anchor)
    return True


def pareto_chart(
    target,
    table: TableRange,
    *,
    title: str,
    label_key: str,
    value_key: str,
    cumulative_key: str,
    anchor: str,
    color: str = RED,
    number_format: str | None = None,
    max_rows: int = 10,
    width: float = 13.8,
    height: float = 8.0,
) -> bool:
    """Pareto: barras ordenadas + percentual acumulado no eixo secundário."""

    if table.rows <= 0:
        return False
    source = target.parent[table.sheet_title]
    last = min(table.last_row, table.first_row + max_rows - 1)
    barras = BarChart()
    barras.type = "col"
    _base_style(barras, title)
    barras.add_data(
        Reference(source, min_col=table.column_index(value_key), min_row=table.header_row, max_row=last),
        titles_from_data=True,
    )
    barras.set_categories(
        Reference(source, min_col=table.column_index(label_key), min_row=table.first_row, max_row=last)
    )
    _series_fill(barras.series[0], color)
    number_format = _chart_format(number_format)
    if number_format:
        barras.y_axis.numFmt = number_format
    barras.gapWidth = 40

    linha = LineChart()
    linha.add_data(
        Reference(source, min_col=table.column_index(cumulative_key), min_row=table.header_row, max_row=last),
        titles_from_data=True,
    )
    linha.y_axis.axId = 200
    linha.y_axis.numFmt = "0%"
    linha.y_axis.title = "% acumulado"
    linha.y_axis.delete = False
    acumulado = linha.series[0]
    acumulado.graphicalProperties = GraphicalProperties()
    acumulado.graphicalProperties.line = LineProperties(solidFill=NAVY, w=20000)
    acumulado.smooth = False

    # Recomendação do openpyxl para eixo secundário: sem isto o "% acumulado"
    # não ganha o próprio eixo à direita.
    barras.y_axis.crosses = "max"
    barras += linha
    barras.width = width
    barras.height = height
    target.add_chart(barras, anchor)
    return True


# --- Abas técnicas ----------------------------------------------------------

FRIENDLY = {
    "op": "OP", "oee": "OEE", "ftt": "FTT", "id": "ID", "kpi": "KPI", "kpis": "KPIs",
    "mtbf": "MTBF", "mttr": "MTTR", "ops": "OPs", "inicio": "Início", "fim": "Fim",
    "producao": "Produção", "operacao": "Operação", "descricao": "Descrição",
    "quantidade": "Quantidade", "recurso": "Recurso", "setor": "Setor", "periodo": "Período",
    "production": "Produção", "quality": "Qualidade", "good": "Produção boa", "scrap": "Refugo",
    "rework": "Retrabalho", "resources": "Recursos", "totals": "Totais", "count": "Quantidade",
    "refugo": "Refugo", "retrabalho": "Retrabalho", "motivo": "Motivo", "segundos": "Segundos",
    "availability": "Disponibilidade", "reason": "Justificativa", "source": "Fonte",
    "value": "Valor", "unit": "Unidade", "seconds": "Segundos", "sectors": "Setores",
    "hours": "Tempos", "policy": "Política", "policies": "Políticas", "audit": "Auditoria",
    "orders": "Ordens", "items": "Itens", "evidence": "Evidências", "insights": "Análises",
}


def humanize(value: str) -> str:
    partes = str(value or "").replace("-", "_").split("_")
    return " ".join(FRIENDLY.get(parte.casefold(), parte.capitalize()) for parte in partes if parte)


TECHNICAL_COLUMNS = (
    Column("grupo", "Grupo", "text", 38),
    Column("campo", "Campo", "text", 32),
    Column("valor", "Valor", "text", 60, wrap=True),
    Column("caminho", "Caminho canônico", "text", 46),
)

CONVENTIONS = (
    ("Percentuais", "Uma casa decimal na leitura executiva; abaixo de 0,1% o valor ganha precisão para não se confundir com zero."),
    ("Durações", "[h]:mm:ss nas tabelas e “0h 00min” nos cartões; duração negativa aparece como texto com sinal, limitação do Excel."),
    ("Datas", "dd/mm/aaaa e dd/mm/aaaa hh:mm."),
    ("Ausência de dado", "“—” significa dado ausente e nunca é substituído por zero."),
    ("Fonte dos números", "Serviços canônicos do Gestor de Peças; a planilha não recalcula indicador industrial."),
)


def technical_sheet(
    workbook,
    payload: dict[str, Any],
    filters,
    *,
    generated_at: datetime,
    title: str = "Dados Técnicos",
    hidden: bool = True,
    limit: int = 4000,
) -> None:
    """Rastreabilidade completa fora da experiência do gestor."""

    ws = workbook.create_sheet(title)
    sheet = ReportSheet(ws, width=4, default_column_width=30)
    sheet.header(
        "Dados técnicos e auditoria",
        filters,
        generated_at=generated_at,
        subtitle=(
            "Campos canônicos do backend preservados para auditoria; esta aba não é a interface de "
            "leitura do relatório."
        ),
    )
    linhas = [
        {"grupo": "Convenções de apresentação", "campo": campo, "valor": texto, "caminho": ""}
        for campo, texto in CONVENTIONS
    ]
    for caminho, (grupo, campo, valor) in _flatten_with_path(payload):
        linhas.append({
            "grupo": grupo,
            "campo": campo,
            "valor": valor if isinstance(valor, (int, float, bool)) or valor is None else safe_text(valor),
            "caminho": caminho,
        })
        if len(linhas) >= limit:
            break
    if linhas:
        sheet.table(TECHNICAL_COLUMNS, linhas, zebra=True)
    else:
        sheet.empty("Nenhum campo escalar foi produzido pelo backend neste relatório.")
    sheet.print_setup()
    if hidden:
        ws.sheet_state = "hidden"


def _flatten_with_path(value: Any, path: tuple[str, ...] = ()):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from _flatten_with_path(child, (*path, str(key)))
        return
    if isinstance(value, list):
        return
    grupo = " › ".join(humanize(parte) for parte in path[:-1]) or "Raiz"
    campo = humanize(path[-1]) if path else "Valor"
    yield ".".join(path), (grupo, campo, value)


__all__ = [
    "AVAILABILITY_LABELS", "AMBER", "BLUE", "Card", "Column", "FORMATS", "GREEN", "LINE",
    "MISSING", "MUTED", "Metric", "NAVY", "RED", "ReportSheet", "SURFACE", "TEXT",
    "TableRange", "WHITE", "coerce", "composition_chart", "duration_text", "filters_label",
    "grouped_chart", "humanize", "metric", "number_text", "pareto_chart", "parse_datetime", "period_label",
    "ranking_chart", "safe_text", "sheet_title", "technical_sheet", "trend_chart",
]
